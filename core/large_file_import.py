#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
大文件导入处理模块。
使用 openpyxl 的只读模式流式读取，避免把整张表一次性装入内存。
"""

import json
import os
import sqlite3
import threading

from openpyxl import load_workbook

from config.config import Config


PROGRESS_FILE = os.path.join(Config.BASE_DIR, 'data', 'import_progress.json')


def _quote_identifier(identifier):
    return '"' + str(identifier).replace('"', '""') + '"'


def _normalize_headers(row_values):
    headers = []
    used = {}
    for index, cell in enumerate(row_values, start=1):
        header = str(cell).strip() if cell is not None else ''
        if not header:
            header = f'列{index}'
        suffix = used.get(header, 0) + 1
        used[header] = suffix
        if suffix > 1:
            header = f'{header}_{suffix}'
        headers.append(header)
    return headers


class LargeFileImporter:
    """大文件导入器。"""

    MAX_ROWS = 990000
    BATCH_SIZE = 5000
    READ_PROGRESS_INTERVAL = 5000

    def __init__(self, db_path, table_name, use_json_data=False):
        self.db_path = db_path
        self.table_name = table_name
        self.use_json_data = use_json_data
        self.progress = {
            'status': 'pending',
            'total_rows': 0,
            'processed_rows': 0,
            'percent': 0,
            'message': '准备中...',
            'error': None
        }
        self._lock = threading.Lock()
        self._last_save_time = 0

    def update_progress(self, **kwargs):
        with self._lock:
            self.progress.update(kwargs)
            total_rows = self.progress.get('total_rows') or 0
            processed_rows = self.progress.get('processed_rows') or 0
            if total_rows > 0:
                self.progress['percent'] = min(100, int(processed_rows / total_rows * 100))

            import time
            current_time = time.time()
            if current_time - self._last_save_time > 1 or kwargs.get('status') in {'completed', 'error'}:
                self._save_progress()
                self._last_save_time = current_time

    def _save_progress(self):
        try:
            with open(PROGRESS_FILE, 'w', encoding='utf-8') as f:
                json.dump(self.progress, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f'[ERROR] 保存进度失败: {e}')

    def get_progress(self):
        with self._lock:
            return self.progress.copy()

    def _prepare_table(self, cursor, headers):
        if self.use_json_data:
            cursor.executescript(
                f'''
                DROP TABLE IF EXISTS {self.table_name};
                CREATE TABLE {self.table_name} (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    json_data TEXT
                );
                '''
            )
            return 'INSERT INTO {table_name} (json_data) VALUES (?)'.format(
                table_name=self.table_name
            )

        column_defs = ', '.join([f'{_quote_identifier(col)} TEXT' for col in headers])
        cursor.executescript(
            f'''
            DROP TABLE IF EXISTS {self.table_name};
            CREATE TABLE {self.table_name} (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                {column_defs}
            );
            '''
        )
        col_names = ', '.join([_quote_identifier(col) for col in headers])
        placeholders = ', '.join(['?'] * len(headers))
        return f'INSERT INTO {self.table_name} ({col_names}) VALUES ({placeholders})'

    def _flush_batch(self, cursor, insert_sql, batch):
        if batch:
            cursor.executemany(insert_sql, batch)
            batch.clear()

    def import_excel_file(self, file_path, sheet_name=0, expected_headers=None, allow_extra_columns=False):
        conn = None
        wb = None
        try:
            self.update_progress(status='reading', message='正在读取 Excel 文件...')
            wb = load_workbook(file_path, read_only=True, data_only=True)

            if isinstance(sheet_name, int):
                if sheet_name < 0 or sheet_name >= len(wb.worksheets):
                    raise ValueError(f'工作表索引超出范围: {sheet_name}')
                ws = wb.worksheets[sheet_name]
            elif isinstance(sheet_name, str):
                if sheet_name not in wb.sheetnames:
                    raise ValueError(f'工作表不存在: {sheet_name}')
                ws = wb[sheet_name]
            else:
                raise ValueError(f'不支持的工作表参数类型: {type(sheet_name).__name__}')

            total_rows = max((ws.max_row or 0) - 1, 0)
            if total_rows <= 0:
                raise ValueError('文件为空或无法读取')
            if total_rows > self.MAX_ROWS:
                total_rows = self.MAX_ROWS

            row_iter = ws.iter_rows(values_only=True)
            header_row = next(row_iter, None)
            if header_row is None:
                raise ValueError('文件为空或无法读取')
            actual_headers = _normalize_headers(header_row)
            if expected_headers:
                headers = [str(item).strip() for item in expected_headers if str(item).strip()]
                if not headers:
                    raise ValueError('系统标准表头为空，无法导入')
                if allow_extra_columns:
                    if len(actual_headers) < len(headers):
                        raise ValueError('导入文件列数量不足，无法匹配系统标准表头')
                elif len(actual_headers) != len(headers):
                    raise ValueError('导入文件列数量与系统标准表头不一致')
            else:
                headers = actual_headers

            self.update_progress(
                status='importing',
                total_rows=total_rows,
                processed_rows=0,
                message=f'开始写入数据库，共 {total_rows} 行...'
            )

            conn = sqlite3.connect(self.db_path)
            cursor = conn.cursor()
            cursor.execute('PRAGMA journal_mode = WAL')
            cursor.execute('PRAGMA synchronous = OFF')
            cursor.execute('PRAGMA cache_size = -300000')
            cursor.execute('PRAGMA temp_store = MEMORY')
            cursor.execute('PRAGMA mmap_size = 536870912')
            cursor.execute('PRAGMA locking_mode = EXCLUSIVE')

            insert_sql = self._prepare_table(cursor, headers)
            conn.commit()

            processed_rows = 0
            batch = []

            for row in row_iter:
                if processed_rows >= total_rows:
                    break

                row_values = [str(cell) if cell is not None else '' for cell in row]
                if len(row_values) < len(headers):
                    row_values.extend([''] * (len(headers) - len(row_values)))
                elif len(row_values) > len(headers):
                    row_values = row_values[:len(headers)]

                if self.use_json_data:
                    row_dict = {headers[idx]: row_values[idx] for idx in range(len(headers))}
                    batch.append((json.dumps(row_dict, ensure_ascii=False),))
                else:
                    batch.append(tuple(row_values))

                processed_rows += 1

                if len(batch) >= self.BATCH_SIZE:
                    self._flush_batch(cursor, insert_sql, batch)
                    conn.commit()
                    self.update_progress(
                        status='importing',
                        total_rows=total_rows,
                        processed_rows=processed_rows,
                        message=f'正在写入数据库...{processed_rows}/{total_rows} 行'
                    )
                elif processed_rows % self.READ_PROGRESS_INTERVAL == 0:
                    self.update_progress(
                        status='importing',
                        total_rows=total_rows,
                        processed_rows=processed_rows,
                        message=f'正在处理...{processed_rows}/{total_rows} 行'
                    )

            self._flush_batch(cursor, insert_sql, batch)
            conn.commit()

            self.update_progress(
                status='completed',
                total_rows=processed_rows,
                processed_rows=processed_rows,
                message=f'导入完成！共导入 {processed_rows:,} 条记录'
            )

            return {
                'success': True,
                'total_rows': processed_rows,
                'imported_rows': processed_rows
            }
        except Exception as e:
            import traceback
            traceback.print_exc()
            self.update_progress(
                status='error',
                error=str(e),
                message=f'导入失败: {str(e)}'
            )
            return {'success': False, 'error': str(e)}
        finally:
            if wb:
                try:
                    wb.close()
                except Exception:
                    pass
            if conn:
                try:
                    conn.close()
                except Exception:
                    pass


def get_import_progress():
    """获取当前导入进度。"""
    try:
        if os.path.exists(PROGRESS_FILE):
            with open(PROGRESS_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        return {'status': 'pending', 'message': '无导入任务'}
    except Exception:
        return {'status': 'pending', 'message': '无导入任务'}


def clear_import_progress():
    """清除导入进度。"""
    try:
        if os.path.exists(PROGRESS_FILE):
            os.remove(PROGRESS_FILE)
    except Exception:
        pass
