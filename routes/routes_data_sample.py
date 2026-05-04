# -*- coding: utf-8 -*-
"""
数据样例标准管理路由
"""
from io import BytesIO

from flask import Blueprint, jsonify, request, send_file
import json
import os
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# 创建Blueprint
data_sample_bp = Blueprint('data_sample', __name__)

# 配置文件路径
CONFIG_FILE = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'config', 'data_sample_standards.json')



def load_standards():
    """加载数据样例标准配置"""
    try:
        with open(CONFIG_FILE, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception as e:
        return {
            "name": "数据分级标准样例",
            "description": "用于根据数据名称自动填充数据样例，并应用敏感数据脱敏/加密规则",
            "version": "1.0",
            "total_count": 0,
            "standards": []
        }


def save_standards(data):
    """保存数据样例标准配置"""
    try:
        # 更新元数据
        data['last_updated'] = datetime.now().strftime('%Y-%m-%d')
        data['total_count'] = len(data.get('standards', []))

        with open(CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        return True
    except Exception as e:
        print(f'保存配置失败: {e}')
        return False


@data_sample_bp.route('/api/data-sample-standards', methods=['GET'])
def get_standards():
    """获取所有数据样例标准"""
    try:
        data = load_standards()
        return jsonify({
            'success': True,
            'data': data
        })
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@data_sample_bp.route('/api/data-sample-standards', methods=['POST'])
def add_standard():
    """添加数据样例标准"""
    try:
        req_data = request.get_json()

        data = load_standards()

        # 生成新ID
        max_id = max([s.get('id', 0) for s in data.get('standards', [])], default=0)
        new_id = max_id + 1

        # 处理 samples 字段，兼容旧格式
        samples = req_data.get('samples', [])
        if not samples and req_data.get('sample'):
            samples = [req_data.get('sample')]

        new_standard = {
            'id': new_id,
            'name': req_data.get('name', ''),
            'category': req_data.get('category', ''),
            'level': req_data.get('level', ''),
            'samples': samples,
            'enabled': req_data.get('enabled', True)
        }

        data['standards'].append(new_standard)

        if save_standards(data):
            return jsonify({
                'success': True,
                'data': new_standard,
                'message': '添加成功'
            })
        else:
            return jsonify({
                'success': False,
                'error': '保存失败'
            }), 500
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@data_sample_bp.route('/api/data-sample-standards/<int:standard_id>', methods=['PUT'])
def update_standard(standard_id):
    """更新数据样例标准"""
    try:
        req_data = request.get_json()

        data = load_standards()

        # 查找并更新标准
        for standard in data.get('standards', []):
            if standard.get('id') == standard_id:
                standard['name'] = req_data.get('name', standard['name'])
                standard['category'] = req_data.get('category', standard['category'])
                standard['level'] = req_data.get('level', standard['level'])

                # 处理 samples 字段，兼容旧格式
                samples = req_data.get('samples', [])
                if not samples and req_data.get('sample'):
                    samples = [req_data.get('sample')]
                if samples:
                    standard['samples'] = samples

                standard['enabled'] = req_data.get('enabled', standard['enabled'])
                break

        if save_standards(data):
            return jsonify({
                'success': True,
                'message': '更新成功'
            })
        else:
            return jsonify({
                'success': False,
                'error': '保存失败'
            }), 500
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@data_sample_bp.route('/api/data-sample-standards/<int:standard_id>', methods=['DELETE'])
def delete_standard(standard_id):
    """删除数据样例标准"""
    try:
        data = load_standards()

        # 删除标准
        standards = data.get('standards', [])
        original_count = len(standards)
        data['standards'] = [s for s in standards if s.get('id') != standard_id]

        if len(data['standards']) < original_count:
            if save_standards(data):
                return jsonify({
                    'success': True,
                    'message': '删除成功'
                })

        return jsonify({
            'success': False,
            'error': '标准不存在或删除失败'
        }), 404
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@data_sample_bp.route('/api/data-sample-standards/batch', methods=['POST'])
def batch_import():
    """批量导入数据样例标准"""
    try:
        req_data = request.get_json()
        standards = req_data.get('standards', [])

        data = load_standards()

        # 生成新ID
        max_id = max([s.get('id', 0) for s in data.get('standards', [])], default=0)

        imported_count = 0
        for standard_data in standards:
            max_id += 1

            # 处理 samples 字段，兼容旧格式
            samples = standard_data.get('samples', [])
            if not samples and standard_data.get('sample'):
                samples = [standard_data.get('sample')]
            if not samples:
                samples = ['']

            new_standard = {
                'id': max_id,
                'name': standard_data.get('name', ''),
                'category': standard_data.get('category', ''),
                'level': standard_data.get('level', ''),
                'samples': samples,
                'enabled': standard_data.get('enabled', True)
            }
            data['standards'].append(new_standard)
            imported_count += 1

        if save_standards(data):
            return jsonify({
                'success': True,
                'imported_count': imported_count,
                'message': f'成功导入{imported_count}条标准'
            })
        else:
            return jsonify({
                'success': False,
                'error': '保存失败'
            }), 500
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@data_sample_bp.route('/api/data-sample-standards/export', methods=['GET'])
def export_standards():
    """????????? Excel ??"""
    try:
        name_col = '\u6570\u636e\u7c7b\u578b\uff08\u6570\u7ba1\u5f00\u53d1\u63d0\u4f9b\uff09'
        category_col = '\u6570\u636e\u5206\u7c7b'
        level_col = '\u6570\u636e\u5206\u7ea7'
        sample_col = '\u6570\u636e\u6837\u4f8b'
        sheet_name = '\u6570\u636e\u5206\u7ea7\u6807\u51c6\u6837\u4f8b'

        data = load_standards()
        standards = data.get('standards', [])
        df = pd.DataFrame({
            name_col: [s.get('name', '') for s in standards],
            category_col: [s.get('category', '') for s in standards],
            level_col: [s.get('level', '') for s in standards],
            sample_col: [
                '\u3001'.join(s.get('samples', [''])) if isinstance(s.get('samples'), list) else s.get('sample', '')
                for s in standards
            ]
        })

        filename = f"{sheet_name}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        output = BytesIO()

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            worksheet = writer.sheets[sheet_name]

            header_font = Font(bold=True, size=11)
            header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
            header_alignment = Alignment(horizontal='center', vertical='center')
            cell_alignment = Alignment(vertical='center')
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            for col, width in {'A': 30, 'B': 35, 'C': 18, 'D': 40}.items():
                worksheet.column_dimensions[col].width = width

            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
                for cell in row:
                    cell.alignment = cell_alignment
                    cell.border = thin_border

            worksheet.row_dimensions[1].height = 25

        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@data_sample_bp.route('/api/data-sample-standards/import', methods=['POST'])
def import_standards():
    """? Excel ??????????"""
    try:
        if 'file' not in request.files:
            return jsonify({
                'success': False,
                'error': '\u672a\u627e\u5230\u4e0a\u4f20\u6587\u4ef6'
            }), 400

        file = request.files['file']
        if file.filename == '':
            return jsonify({
                'success': False,
                'error': '\u672a\u9009\u62e9\u6587\u4ef6'
            }), 400

        if not file.filename.endswith(('.xlsx', '.xls')):
            return jsonify({
                'success': False,
                'error': '\u4ec5\u652f\u6301Excel\u6587\u4ef6\uff08.xlsx\u6216.xls\uff09'
            }), 400

        name_col = '\u6570\u636e\u7c7b\u578b\uff08\u6570\u7ba1\u5f00\u53d1\u63d0\u4f9b\uff09'
        category_col = '\u6570\u636e\u5206\u7c7b'
        level_col = '\u6570\u636e\u5206\u7ea7'
        sample_col = '\u6570\u636e\u6837\u4f8b'

        file.stream.seek(0)
        df = pd.read_excel(file.stream, engine='openpyxl')

        required_columns = [name_col, category_col, level_col, sample_col]
        missing_columns = [col for col in required_columns if col not in df.columns]
        if missing_columns:
            return jsonify({
                'success': False,
                'error': f"Excel\u6587\u4ef6\u7f3a\u5c11\u5fc5\u9700\u7684\u5217: {', '.join(missing_columns)}"
            }), 400

        data = load_standards()
        existing_standards = data.get('standards', [])
        max_id = max([s.get('id', 0) for s in existing_standards], default=0)

        imported_count = 0
        updated_count = 0
        skipped_count = 0
        errors = []

        for idx, row in df.iterrows():
            try:
                name = str(row.get(name_col, '')).strip()
                category = str(row.get(category_col, '')).strip()
                level = str(row.get(level_col, '')).strip()
                sample_value = str(row.get(sample_col, '')).strip()

                if not name or name == 'nan':
                    skipped_count += 1
                    continue

                samples = []
                if sample_value and sample_value != 'nan':
                    separators = ['\u3001', '\uff0c', ',', ';', '\uff1b', '\n']
                    samples_temp = [sample_value]
                    for sep in separators:
                        new_samples = []
                        for item in samples_temp:
                            new_samples.extend(item.split(sep))
                        samples_temp = new_samples
                    samples = [item.strip() for item in samples_temp if item.strip()]
                if not samples:
                    samples = ['']

                existing = None
                for item in existing_standards:
                    if item.get('name', '').strip() == name:
                        existing = item
                        break

                if existing:
                    existing['category'] = category
                    existing['level'] = level
                    existing['samples'] = samples
                    existing['enabled'] = True
                    updated_count += 1
                else:
                    max_id += 1
                    existing_standards.append({
                        'id': max_id,
                        'name': name,
                        'category': category,
                        'level': level,
                        'samples': samples,
                        'enabled': True
                    })
                    imported_count += 1
            except Exception as e:
                errors.append(f"\u7b2c{idx + 2}\u884c\u5904\u7406\u5931\u8d25: {str(e)}")

        data['standards'] = existing_standards
        data['total_count'] = len(existing_standards)
        data['last_updated'] = datetime.now().strftime('%Y-%m-%d')

        if save_standards(data):
            response_data = {
                'success': True,
                'imported_count': imported_count,
                'updated_count': updated_count,
                'skipped_count': skipped_count,
                'total_count': len(existing_standards),
                'message': f"\u5bfc\u5165\u6210\u529f\uff1a\u65b0\u589e{imported_count}\u6761\uff0c\u66f4\u65b0{updated_count}\u6761\uff0c\u8df3\u8fc7{skipped_count}\u6761"
            }
            if errors:
                response_data['errors'] = errors[:10]
                response_data['message'] += f"\uff0c\u53e6\u6709{len(errors)}\u6761\u9519\u8bef"
            return jsonify(response_data)

        return jsonify({
            'success': False,
            'error': '\u4fdd\u5b58\u5931\u8d25'
        }), 500
    except Exception as e:
        return jsonify({
            'success': False,
            'error': f"\u5bfc\u5165\u5931\u8d25: {str(e)}"
        }), 500


@data_sample_bp.route('/api/data-sample-standards/template', methods=['GET'])
def download_template():
    """??????"""
    try:
        name_col = '\u6570\u636e\u7c7b\u578b\uff08\u6570\u7ba1\u5f00\u53d1\u63d0\u4f9b\uff09'
        category_col = '\u6570\u636e\u5206\u7c7b'
        level_col = '\u6570\u636e\u5206\u7ea7'
        sample_col = '\u6570\u636e\u6837\u4f8b'
        sheet_name = '\u6570\u636e\u5206\u7ea7\u6807\u51c6\u6837\u4f8b'
        info_sheet_name = '\u586b\u5199\u8bf4\u660e'

        df = pd.DataFrame({
            name_col: [],
            category_col: [],
            level_col: [],
            sample_col: []
        })

        filename = '\u6570\u636e\u5206\u7ea7\u6807\u51c6\u6837\u4f8b_\u5bfc\u5165\u6a21\u677f.xlsx'
        output = BytesIO()

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name=sheet_name)
            workbook = writer.book
            worksheet = writer.sheets[sheet_name]

            header_font = Font(bold=True, size=11, color='FFFFFF')
            header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
            header_alignment = Alignment(horizontal='center', vertical='center')
            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            for col, width in {'A': 30, 'B': 35, 'C': 18, 'D': 40}.items():
                worksheet.column_dimensions[col].width = width

            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            worksheet.row_dimensions[1].height = 25

            info_sheet = workbook.create_sheet(info_sheet_name)
            instructions = [
                ['\u6570\u636e\u5206\u7ea7\u6807\u51c6\u6837\u4f8b\u5bfc\u5165\u6a21\u677f\u586b\u5199\u8bf4\u660e'],
                [''],
                ['1. \u6a21\u677f\u8bf4\u660e'],
                ['   \u672c\u6a21\u677f\u7528\u4e8e\u6279\u91cf\u5bfc\u5165\u6570\u636e\u5206\u7ea7\u6807\u51c6\u6837\u4f8b\u3002'],
                ['   \u5bfc\u5165\u540e\u4f1a\u6839\u636e\u6570\u636e\u540d\u79f0\u81ea\u52a8\u5339\u914d\uff0c\u5df2\u5b58\u5728\u7684\u8bb0\u5f55\u4f1a\u66f4\u65b0\uff0c\u4e0d\u5b58\u5728\u7684\u8bb0\u5f55\u4f1a\u65b0\u589e\u3002'],
                [''],
                ['2. \u586b\u5199\u89c4\u8303'],
                [f'   {name_col}\uff1a\u5fc5\u586b\uff0c\u586b\u5199\u6807\u51c6\u6570\u636e\u540d\u79f0\u3002'],
                [f'   {category_col}\uff1a\u5fc5\u586b\u3002'],
                [f'   {level_col}\uff1a\u5fc5\u586b\u3002'],
                [f'   {sample_col}\uff1a\u5fc5\u586b\uff0c\u591a\u4e2a\u6837\u4f8b\u53ef\u7528\u987f\u53f7\u3001\u9017\u53f7\u3001\u5206\u53f7\u6216\u6362\u884c\u5206\u9694\u3002'],
                [''],
                ['3. \u6ce8\u610f\u4e8b\u9879'],
                ['   - \u4e0d\u8981\u4fee\u6539\u8868\u5934\u5217\u540d\u3002'],
                ['   - \u540c\u540d\u6570\u636e\u7c7b\u578b\u518d\u6b21\u5bfc\u5165\u65f6\u4f1a\u8986\u76d6\u539f\u8bb0\u5f55\u3002'],
            ]

            for row_idx, row_data in enumerate(instructions, start=1):
                for col_idx, cell_value in enumerate(row_data, start=1):
                    cell = info_sheet.cell(row=row_idx, column=col_idx, value=cell_value)
                    cell.font = Font(bold=(row_idx == 1), size=14 if row_idx == 1 else 11)
                    info_sheet.column_dimensions[chr(64 + col_idx)].width = 80

        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500
