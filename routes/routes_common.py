#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
通用API路由模块
处理通用的API请求，包括首页、列配置、统计信息、文件管理等
"""
from flask import Blueprint, request, jsonify, send_file
from io import BytesIO
import copy
from difflib import SequenceMatcher
import json
import math
import os
import pandas as pd
import re
import sqlite3
from datetime import datetime
from urllib.parse import quote, urlparse

from config.config import Config, DatabaseConfig, PathConfig, PaginationConfig
from core.large_file_import import (
    LargeFileImporter,
    clear_import_progress,
    PROGRESS_FILE,
    _normalize_headers as normalize_import_headers,
)
from core.large_file_import import get_import_progress as get_import_progress_func
from core.utils import success_response, error_response

# 创建Blueprint
common_bp = Blueprint('common', __name__)
PUBLIC_CONFIG_BUSINESS_SYSTEM_NAME_MAPPING_KEY = 'business_system_name_mapping'
PUBLIC_CONFIG_ASSET_NAME_MAPPING_KEY = 'asset_name_mapping'
BUSINESS_SYSTEM_NAME_MAPPING_FILE = os.path.join(
    Config.BASE_DIR, 'config', 'business_system_name_mapping.json'
)
ASSET_NAME_MAPPING_FILE = os.path.join(
    Config.BASE_DIR, 'config', 'asset_name_mapping.json'
)
PUBLIC_CONFIG_DEFINITIONS_FILE = os.path.join(
    Config.BASE_DIR, 'config', 'public_configs.json'
)
PUBLIC_CONFIG_BINDINGS_FILE = os.path.join(
    Config.BASE_DIR, 'config', 'public_config_bindings.json'
)
DYNAMIC_CONFIG_QUERY_DEFINITIONS_FILE = os.path.join(
    Config.BASE_DIR, 'config', 'dynamic_config_query_definitions.json'
)
PUBLIC_CONFIG_REGISTRY = {
    PUBLIC_CONFIG_BUSINESS_SYSTEM_NAME_MAPPING_KEY: {
        'title': '公共配置',
        'name': '业务系统名称映射',
        'description': '用于将合并结果表中的业务系统名称转换为数据概览表中的标准名称',
        'contract_role': 'runtime',
        'affects_export': True,
        'effect_summary': '保存后会直接影响已接线报表导出',
        'config_type': 'mapping_table',
        'detail_api': '/api/public-configs/business-system-name-mappings',
        'save_api': '/api/public-configs/business-system-name-mappings/save',
        'template_api': '/api/public-configs/business-system-name-mappings/template',
        'import_api': '/api/public-configs/business-system-name-mappings/import',
        'export_api': '/api/public-configs/business-system-name-mappings/export',
        'editor_mode': 'business_system_name_mapping',
        'supports_detail': True,
        'supports_crud': True,
        'supports_import_export': True,
        'storage_file': BUSINESS_SYSTEM_NAME_MAPPING_FILE,
    },
    PUBLIC_CONFIG_ASSET_NAME_MAPPING_KEY: {
        'title': '公共配置',
        'name': '资产名称映射',
        'description': '用于将源数据中的资产名称转换为统一口径的标准资产名称',
        'contract_role': 'reserve',
        'affects_export': False,
        'effect_summary': '配置层已初始化，具体使用方式待后续按报表列决定',
        'config_type': 'mapping_table',
        'detail_api': '/api/public-configs/asset-name-mappings',
        'template_api': '/api/public-configs/asset-name-mappings/template',
        'import_api': '/api/public-configs/asset-name-mappings/import',
        'export_api': '/api/public-configs/asset-name-mappings/export',
        'editor_mode': 'asset_name_mapping',
        'supports_detail': True,
        'supports_crud': True,
        'supports_import_export': True,
        'storage_file': ASSET_NAME_MAPPING_FILE,
    }
}
EXPORT_SCOPE_CONFIG_FILE = os.path.join(Config.BASE_DIR, 'config', 'export_scope_config.json')
DEFAULT_EXPORT_SCOPE_CONFIG = {
    'assets': {
        'enabled': False,
        'mode': 'all',
        'value': 0
    },
    'merge_results': {
        'enabled': False,
        'mode': 'all',
        'value': 0
    }
}
MERGE_LEVEL_RANK_PATTERN = re.compile(r'一般级\s*-\s*第\s*([1-4])\s*小级')
MERGE_FINE_SCOPE_POLICY_FILE = os.path.join(
    Config.BASE_DIR, 'config', 'merge_fine_scope_policy.json'
)
DEFAULT_MERGE_FINE_SCOPE_POLICY = {
    'enabled': False,
    'max_target_rows': 950000,
    'tolerance_rows': 50000,
    'manual_ratio_enabled': False,
    'manual_ratio': 0.0,
    'target_mode': 'absolute',
    'target_min': 900000,
    'target_max': 950000,
    'preferred_target': 950000,
    'target_min_percent': 90.0,
    'target_max_percent': 95.0,
    'preferred_target_percent': 95.0,
    'small_source_threshold': 10000,
    'rounding': 'largest_remainder',
    'keep_at_least_one_when_nonzero': False,
    'fallback_when_small_exceeds_max': 'warn_only',
}

ADVANCED_SEARCH_FIELD_MAPPING = {
    'business_system': '业务系统',
    'data_asset_name': '数据资产名称',
    'ip_address': 'IP地址',
    'data_type': '数据类型',
    'data_level': '数据分级',
    'status': '状态',
    'responsible_person': '负责人'
}

ADVANCED_SEARCH_OPERATORS = {
    'contains',
    'equals',
    'not_contains',
    'not_equals',
    'is_empty',
    'not_empty'
}

# 导入日志记录函数（延迟导入避免循环依赖）
def get_log_operation():
    """延迟导入日志记录函数"""
    from routes.routes_logs import log_operation
    return log_operation


def resolve_safe_file_path(base_dir, filename):
    """解析指定目录下的安全文件路径，阻止路径穿越。"""
    raw_filename = str(filename or '')
    safe_filename = os.path.basename(raw_filename)
    if safe_filename != raw_filename or safe_filename in ('', '.', '..'):
        raise ValueError('无效文件名')
    return os.path.join(base_dir, safe_filename), safe_filename


def _safe_remove_file(file_path):
    """安全删除文件，失败时仅记录日志。"""
    safe_path = str(file_path or '').strip()
    if not safe_path:
        return
    try:
        if os.path.exists(safe_path):
            os.remove(safe_path)
    except Exception as cleanup_err:
        print(f'[临时文件清理失败] {cleanup_err}')


def get_datafiles_source_dir(data_type):
    """获取 DataFiles 下的源文件目录。"""
    datafiles_base_dir = os.path.join(Config.BASE_DIR, 'DataFiles')
    if data_type in ('merge', 'merge_results'):
        return os.path.join(datafiles_base_dir, 'merge_results')
    return os.path.join(datafiles_base_dir, 'assets')


def find_source_file_path(filename, file_type='assets'):
    """在兼容目录中查找源文件路径。"""
    _, safe_filename = resolve_safe_file_path('', filename)

    search_dirs = [PathConfig.UPLOAD_FOLDER]
    if file_type in ('merge', 'merge_results'):
        search_dirs.append(get_datafiles_source_dir('merge_results'))
        search_dirs.append(get_datafiles_source_dir('assets'))
    else:
        search_dirs.append(get_datafiles_source_dir('assets'))
        search_dirs.append(get_datafiles_source_dir('merge_results'))

    seen = set()
    for base_dir in search_dirs:
        if not base_dir or base_dir in seen:
            continue
        seen.add(base_dir)
        file_path, _ = resolve_safe_file_path(base_dir, safe_filename)
        if os.path.isfile(file_path):
            return file_path, safe_filename

    return None, safe_filename


def read_excel_headers(file_path):
    """读取 Excel 首行表头，用于同步列配置。"""
    from openpyxl import load_workbook

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header_row is None:
            return []
        return normalize_import_headers(header_row)
    finally:
        close = getattr(workbook, 'close', None)
        if callable(close):
            close()


def _extract_column_names_from_config(config_data):
    """从列配置中提取列名，兼容旧版根配置和按模式分组配置。"""
    if isinstance(config_data, dict) and isinstance(config_data.get('全部'), dict):
        default_config = config_data.get('全部') or {}
        candidates = default_config.get('columns') or []
    elif isinstance(config_data, dict) and isinstance(config_data.get('columns'), list):
        candidates = config_data.get('columns') or []
    elif isinstance(config_data, dict):
        candidates = []
        ordered_configs = []
        ordered_configs.extend(config_data.values())
        for item in ordered_configs:
            if isinstance(item, dict) and isinstance(item.get('columns'), list):
                candidates = item.get('columns') or []
                break
    else:
        candidates = []

    headers = []
    for column in candidates:
        if not isinstance(column, dict):
            continue
        name = str(column.get('name') or '').strip()
        if not name or name.lower() == 'id':
            continue
        headers.append(name)
    return headers


def _get_workspace_template_headers(page_type):
    """根据页面类型读取当前列配置，生成模板表头。"""
    normalized_page_type = str(page_type or '').strip().lower()
    if normalized_page_type == 'device':
        config_path = PathConfig.ASSETS_COLUMNS_JSON
    elif normalized_page_type == 'merge':
        config_path = PathConfig.MERGE_COLUMNS_JSON
    else:
        raise ValueError('不支持的页面类型')

    try:
        with open(config_path, 'r', encoding='utf-8') as f:
            config_data = json.load(f)
    except Exception as exc:
        raise ValueError('读取列配置失败') from exc

    headers = _extract_column_names_from_config(config_data)
    if not headers:
        raise ValueError('当前列配置为空，无法生成模板')
    return headers


def _get_workspace_template_headers_safe(page_type):
    """安全读取标准表头，读取失败时返回空数组，不阻断导入流程。"""
    try:
        return _get_workspace_template_headers(page_type)
    except Exception:
        return []


def _get_workspace_page_label(page_type):
    normalized_page_type = str(page_type or '').strip().lower()
    return '合并结果' if normalized_page_type == 'merge' else '数据概览'


def _validate_import_headers_against_expected(imported_headers, expected_headers, page_type):
    """校验导入表头是否与系统当前标准表头一致，允许额外列但不允许前缀错位。"""
    page_label = _get_workspace_page_label(page_type)
    normalized_expected = [str(header or '').strip() for header in (expected_headers or [])]
    normalized_imported = [str(header or '').strip() for header in (imported_headers or [])]

    expected_count = len(normalized_expected)
    imported_count = len(normalized_imported)

    if expected_count == 0:
        return False, f'{page_label}导入失败：当前系统未配置标准表头，无法校验导入文件。'

    if imported_count < expected_count:
        return False, (
            f'{page_label}导入失败：列数量不对应，系统要求至少 {expected_count} 列，'
            f'导入文件只有 {imported_count} 列。'
        )

    for index, expected_header in enumerate(normalized_expected):
        actual_header = normalized_imported[index]
        if actual_header == expected_header:
            continue

        actual_position = -1
        try:
            actual_position = normalized_imported.index(expected_header)
        except ValueError:
            actual_position = -1

        if actual_position != -1:
            return False, (
                f'{page_label}导入失败：表头顺序错误，第 {index + 1} 列应为“{expected_header}”，'
                f'当前却是“{actual_header}”。'
            )

        return False, (
            f'{page_label}导入失败：表头名不对应，第 {index + 1} 列应为“{expected_header}”，'
            f'导入文件为“{actual_header}”。'
        )

    return True, None


def _is_force_import_requested():
    """判断当前请求是否为用户确认后的强制导入。"""
    value = str(request.form.get('force_import', '') or '').strip().lower()
    return value in {'1', 'true', 'yes', 'y', 'on'}


def _compare_import_headers(imported_headers, expected_headers):
    """比较导入表头与当前标准表头，返回缺失列与新增列。"""
    normalized_expected = [str(header or '').strip() for header in (expected_headers or []) if str(header or '').strip()]
    normalized_imported = [str(header or '').strip() for header in (imported_headers or []) if str(header or '').strip()]

    if not normalized_expected:
        return [], []

    imported_set = set(normalized_imported)
    expected_set = set(normalized_expected)

    missing_columns = [header for header in normalized_expected if header not in imported_set]
    extra_columns = [header for header in normalized_imported if header not in expected_set]
    return missing_columns, extra_columns


def _build_import_schema_warning(page_type, missing_columns, extra_columns):
    """构建导入列变化提醒文案。"""
    page_label = _get_workspace_page_label(page_type)
    detail_items = []
    if missing_columns:
        detail_items.append(f'缺失列: {", ".join(missing_columns)}')
    if extra_columns:
        detail_items.append(f'新增列: {", ".join(extra_columns)}')

    detail_text = f'（{"；".join(detail_items)}）' if detail_items else ''
    return (
        f'{page_label}导入检测到列变化{detail_text}。'
        '继续导入可能导致映射数据偏差，请确认是否继续。'
    )


def _build_workspace_template_excel(headers, page_type):
    """构建仅包含表头的 Excel 模板。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '合并结果模板' if str(page_type).lower() == 'merge' else '数据概览模板'
    worksheet.append(headers)
    worksheet.freeze_panes = 'A2'
    worksheet.row_dimensions[1].height = 24

    header_fill = PatternFill(fill_type='solid', fgColor='1F6FB2')
    header_font = Font(bold=True, color='FFFFFF')
    header_alignment = Alignment(horizontal='center', vertical='center')
    header_border = Border(
        left=Side(style='thin', color='D7E0EA'),
        right=Side(style='thin', color='D7E0EA'),
        top=Side(style='thin', color='D7E0EA'),
        bottom=Side(style='thin', color='D7E0EA')
    )

    for index, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=index)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = header_border
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(str(header)) * 2 + 4, 14),
            36
        )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def sync_columns_config(config_path, columns, source_file):
    """同步列配置中的列定义与 sourceFile。"""
    try:
        with open(config_path, 'r', encoding='utf-8') as f:
            config = json.load(f)
    except Exception:
        config = {}

    existing_columns = {}
    for item in config.get('columns', []):
        if isinstance(item, dict) and item.get('name') is not None:
            existing_columns[str(item.get('name'))] = item

    normalized_columns = []
    for column_name in columns:
        key = str(column_name)
        if key in existing_columns:
            normalized_columns.append(existing_columns[key])
        else:
            normalized_columns.append({
                'name': key,
                'visible': True,
                'width': 120
            })

    config['columns'] = normalized_columns
    config['sourceFile'] = source_file
    os.makedirs(os.path.dirname(config_path), exist_ok=True)
    with open(config_path, 'w', encoding='utf-8') as f:
        json.dump(config, f, ensure_ascii=False, indent=2)

# 数据概览导入进度追踪（专用于 /api/import/confirm）
import_progress = {
    'status': 'idle',  # idle, running, completed, error
    'current': 0,
    'total': 0,
    'message': '',
    'error': ''
}


def _get_assets_expected_column_count():
    """从当前数据库表结构获取期望的数据列数（不含id）。"""
    try:
        import sqlite3
        conn = sqlite3.connect(DatabaseConfig.ASSETS_DB)
        cursor = conn.cursor()
        cursor.execute('PRAGMA table_info(assets)')
        cols = cursor.fetchall()
        conn.close()
        data_cols = [c for c in cols if c[1] != 'id']
        return len(data_cols) if data_cols else None
    except Exception:
        return None


def _validate_assets_import_columns(uploaded_columns):
    """校验导入文件列数是否与当前assets表结构一致，返回 (ok, error_msg)。"""
    expected = _get_assets_expected_column_count()
    if expected is None:
        return True, None
    actual = [c for c in uploaded_columns if c != 'id']
    if len(actual) != expected:
        return False, (
            f'列数不匹配：当前数据概览表有 {expected} 列数据，'
            f'导入文件有 {len(actual)} 列（不含id），请使用正确格式的文件'
        )
    return True, None


def _write_import_progress_file(status, total_rows, processed_rows, message='', error=''):
    """将普通导入接口的进度同步到统一进度文件，供 /api/import/progress 读取。"""
    safe_total_rows = max(int(total_rows or 0), 0)
    safe_processed_rows = max(int(processed_rows or 0), 0)
    percent = 0
    if safe_total_rows > 0:
        percent = min(100, int(safe_processed_rows / safe_total_rows * 100))
    elif status == 'completed':
        percent = 100

    payload = {
        'status': status,
        'total_rows': safe_total_rows,
        'processed_rows': safe_processed_rows,
        'percent': percent,
        'message': message,
        'error': error or None
    }

    os.makedirs(os.path.dirname(PROGRESS_FILE), exist_ok=True)
    with open(PROGRESS_FILE, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def _set_asset_import_progress(status, current, total, message='', error=''):
    """同时维护旧内存进度和统一进度文件，避免接口返回不同步。"""
    global import_progress

    import_progress = {
        'status': status,
        'current': max(int(current or 0), 0),
        'total': max(int(total or 0), 0),
        'message': message,
        'error': error
    }
    _write_import_progress_file(
        status=status,
        total_rows=import_progress['total'],
        processed_rows=import_progress['current'],
        message=message,
        error=error
    )


@common_bp.route('/api/workspace-template/<page_type>', methods=['GET'])
def download_workspace_template(page_type):
    """下载数据概览或合并结果页面的导入模板。"""
    try:
        headers = _get_workspace_template_headers(page_type)
        output = _build_workspace_template_excel(headers, page_type)
        filename = '合并结果导入模板.xlsx' if str(page_type).lower() == 'merge' else '数据概览导入模板.xlsx'
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except ValueError as exc:
        return jsonify({'error': str(exc)}), 400
    except Exception as exc:
        return jsonify({'error': f'模板生成失败: {exc}'}), 500


def get_db_connection():
    """获取数据库连接"""
    import sqlite3
    conn = sqlite3.connect(DatabaseConfig.ASSETS_DB)
    conn.row_factory = sqlite3.Row
    return conn


def normalize_advanced_search_pagination(page, page_size):
    """标准化高级搜索分页参数。"""
    try:
        page = int(page)
    except (TypeError, ValueError):
        raise ValueError('page 必须为正整数')

    try:
        page_size = int(page_size)
    except (TypeError, ValueError):
        raise ValueError('page_size 必须为正整数')

    if page <= 0:
        raise ValueError('page 必须为正整数')
    if page_size <= 0:
        raise ValueError('page_size 必须为正整数')

    return page, min(page_size, PaginationConfig.MAX_PAGE_SIZE)


def resolve_advanced_search_field(field, valid_db_columns):
    """解析并校验高级搜索字段。"""
    db_field = ADVANCED_SEARCH_FIELD_MAPPING.get(field, field)
    if not db_field or db_field not in valid_db_columns:
        raise ValueError(f'不支持的搜索字段: {field}')
    return db_field


def validate_advanced_search_operator(operator):
    """校验高级搜索操作符。"""
    if operator not in ADVANCED_SEARCH_OPERATORS:
        raise ValueError(f'不支持的搜索操作符: {operator}')
    return operator


def _normalize_export_scope_rule(raw_rule):
    """统一解析整体筛选配置项。"""
    rule = raw_rule if isinstance(raw_rule, dict) else {}
    mode = str(rule.get('mode', 'all') or 'all').strip().lower()
    if mode not in ('all', 'top_n', 'top_percent'):
        mode = 'all'

    raw_value = rule.get('value', 0)
    if mode == 'top_percent':
        try:
            value = float(raw_value)
        except (TypeError, ValueError):
            value = 0.0
    else:
        try:
            value = int(float(raw_value))
        except (TypeError, ValueError):
            value = 0

    return {
        'enabled': bool(rule.get('enabled', False)),
        'mode': mode,
        'value': value
    }


def _load_export_scope_config():
    """读取整体筛选配置，缺失时回退到默认值。"""
    config = copy.deepcopy(DEFAULT_EXPORT_SCOPE_CONFIG)

    try:
        if os.path.exists(EXPORT_SCOPE_CONFIG_FILE):
            with open(EXPORT_SCOPE_CONFIG_FILE, 'r', encoding='utf-8-sig') as f:
                raw_config = json.load(f)
            if isinstance(raw_config, dict):
                for source_name in config.keys():
                    config[source_name] = _normalize_export_scope_rule(
                        raw_config.get(source_name, config[source_name])
                    )
    except Exception:
        pass

    return config


def _save_export_scope_config(config):
    """保存整体筛选配置到配置文件。"""
    normalized_config = {}
    for source_name, default_rule in DEFAULT_EXPORT_SCOPE_CONFIG.items():
        normalized_config[source_name] = _normalize_export_scope_rule(
            (config or {}).get(source_name, default_rule)
        )

    with open(EXPORT_SCOPE_CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump(normalized_config, f, ensure_ascii=False, indent=2)

    return normalized_config


def _normalize_merge_fine_scope_policy(raw_policy):
    policy = raw_policy if isinstance(raw_policy, dict) else {}
    normalized = dict(DEFAULT_MERGE_FINE_SCOPE_POLICY)

    normalized['enabled'] = bool(policy.get('enabled', normalized['enabled']))
    try:
        normalized['max_target_rows'] = max(1, int(float(policy.get('max_target_rows', normalized['max_target_rows']))))
    except (TypeError, ValueError):
        pass
    try:
        normalized['tolerance_rows'] = max(0, int(float(policy.get('tolerance_rows', normalized['tolerance_rows']))))
    except (TypeError, ValueError):
        pass
    normalized['manual_ratio_enabled'] = bool(policy.get('manual_ratio_enabled', normalized['manual_ratio_enabled']))
    try:
        normalized['manual_ratio'] = max(0.0, min(1.0, float(policy.get('manual_ratio', normalized['manual_ratio']))))
    except (TypeError, ValueError):
        normalized['manual_ratio'] = 0.0
    target_mode = str(policy.get('target_mode', normalized['target_mode']) or '').strip().lower()
    normalized['target_mode'] = 'percent' if target_mode == 'percent' else 'absolute'

    try:
        normalized['target_min'] = max(1, int(float(policy.get('target_min', normalized['target_min']))))
    except (TypeError, ValueError):
        pass
    try:
        normalized['target_max'] = max(1, int(float(policy.get('target_max', normalized['target_max']))))
    except (TypeError, ValueError):
        pass
    if normalized['target_min'] > normalized['target_max']:
        normalized['target_min'], normalized['target_max'] = normalized['target_max'], normalized['target_min']

    try:
        preferred = int(float(policy.get('preferred_target', normalized['preferred_target'])))
        preferred = max(normalized['target_min'], min(normalized['target_max'], preferred))
        normalized['preferred_target'] = preferred
    except (TypeError, ValueError):
        normalized['preferred_target'] = normalized['target_max']

    try:
        normalized['small_source_threshold'] = max(0, int(float(
            policy.get('small_source_threshold', normalized['small_source_threshold'])
        )))
    except (TypeError, ValueError):
        pass

    try:
        normalized['target_min_percent'] = max(0.0, min(100.0, float(
            policy.get('target_min_percent', normalized['target_min_percent'])
        )))
    except (TypeError, ValueError):
        pass
    try:
        normalized['target_max_percent'] = max(0.0, min(100.0, float(
            policy.get('target_max_percent', normalized['target_max_percent'])
        )))
    except (TypeError, ValueError):
        pass
    if normalized['target_min_percent'] > normalized['target_max_percent']:
        normalized['target_min_percent'], normalized['target_max_percent'] = (
            normalized['target_max_percent'],
            normalized['target_min_percent']
        )
    try:
        preferred_percent = float(policy.get('preferred_target_percent', normalized['preferred_target_percent']))
        preferred_percent = max(
            normalized['target_min_percent'],
            min(normalized['target_max_percent'], preferred_percent)
        )
        normalized['preferred_target_percent'] = preferred_percent
    except (TypeError, ValueError):
        normalized['preferred_target_percent'] = normalized['target_max_percent']

    rounding = str(policy.get('rounding', normalized['rounding']) or '').strip().lower()
    normalized['rounding'] = 'largest_remainder' if rounding == 'largest_remainder' else 'floor'

    normalized['keep_at_least_one_when_nonzero'] = bool(
        policy.get('keep_at_least_one_when_nonzero', normalized['keep_at_least_one_when_nonzero'])
    )
    fallback = str(policy.get(
        'fallback_when_small_exceeds_max',
        normalized['fallback_when_small_exceeds_max']
    ) or '').strip().lower()
    normalized['fallback_when_small_exceeds_max'] = (
        'compress_all' if fallback == 'compress_all' else 'warn_only'
    )

    return normalized


def _load_merge_fine_scope_policy():
    policy = dict(DEFAULT_MERGE_FINE_SCOPE_POLICY)
    try:
        if os.path.exists(MERGE_FINE_SCOPE_POLICY_FILE):
            with open(MERGE_FINE_SCOPE_POLICY_FILE, 'r', encoding='utf-8-sig') as f:
                raw = json.load(f)
            policy = _normalize_merge_fine_scope_policy(raw)
    except Exception:
        policy = dict(DEFAULT_MERGE_FINE_SCOPE_POLICY)
    return policy


def _save_merge_fine_scope_policy(raw_policy):
    policy = _normalize_merge_fine_scope_policy(raw_policy)
    with open(MERGE_FINE_SCOPE_POLICY_FILE, 'w', encoding='utf-8') as f:
        json.dump(policy, f, ensure_ascii=False, indent=2)
    return policy


def _build_merge_fine_scope_preview(stats_result, policy):
    items = list(stats_result.get('items') or [])
    summary = stats_result.get('summary') or {}
    total_before = int(summary.get('total_rows') or 0)

    small_threshold = int(policy.get('small_source_threshold') or 0)
    max_target_rows = max(1, int(policy.get('max_target_rows') or 950000))
    tolerance_rows = max(0, int(policy.get('tolerance_rows') or 50000))
    manual_ratio_enabled = bool(policy.get('manual_ratio_enabled'))
    manual_ratio = max(0.0, min(1.0, float(policy.get('manual_ratio') or 0.0)))

    preview_items = []
    adjustable_groups = []
    adjustable_total = 0  # L
    fixed_total = 0       # S
    small_total = 0
    large_total = 0

    for item in items:
        source_name = str(item.get('source_name') or '')
        source_label = str(item.get('source_label') or source_name or '空值/未填写')
        source_total = int(item.get('total_count') or 0)
        level_counts = {
            1: int(item.get('level_1_count') or 0),
            2: int(item.get('level_2_count') or 0),
            3: int(item.get('level_3_count') or 0),
            4: int(item.get('level_4_count') or 0),
        }
        level_total = sum(level_counts.values())
        other_count = max(0, source_total - level_total)
        is_small = source_total <= small_threshold

        row = {
            'source_name': source_name,
            'source_label': source_label,
            'source_total': source_total,
            'is_small_source': is_small,
            'level_1_before': level_counts[1],
            'level_2_before': level_counts[2],
            'level_3_before': level_counts[3],
            'level_4_before': level_counts[4],
            'level_1_keep': level_counts[1],
            'level_2_keep': level_counts[2],
            'level_3_keep': level_counts[3],
            'level_4_keep': level_counts[4],
            'other_before': other_count,
            'other_keep': other_count,
            'keep_total': source_total,
            'cut_total': 0,
            'cut_ratio': 0.0
        }
        preview_items.append(row)

        if is_small:
            small_total += source_total
            fixed_total += source_total
            continue

        large_total += source_total
        for rank in (1, 2, 3, 4):
            count = level_counts[rank]
            if count <= 0:
                continue
            adjustable_groups.append({
                'row': row,
                'rank': rank,
                'count': count,
                'keep': 0,
                'remainder': 0.0
            })
            adjustable_total += count
            row[f'level_{rank}_keep'] = 0
        if other_count > 0:
            adjustable_groups.append({
                'row': row,
                'rank': 0,
                'count': other_count,
                'keep': 0,
                'remainder': 0.0
            })
            adjustable_total += other_count
            row['other_keep'] = 0

    target_rows = total_before
    auto_ratio = 1.0
    applied_ratio = 1.0
    status = 'no_change'
    warning_message = ''

    if bool(policy.get('enabled')) and adjustable_total > 0:
        auto_ratio = (max_target_rows - fixed_total) / adjustable_total
        auto_ratio = max(0.0, min(1.0, auto_ratio))
        applied_ratio = manual_ratio if manual_ratio_enabled else auto_ratio
        applied_ratio = max(0.0, min(1.0, applied_ratio))
        keep_adjustable_target = int(math.floor(adjustable_total * applied_ratio))
        keep_adjustable_target = max(0, min(adjustable_total, keep_adjustable_target))
        target_rows = fixed_total + keep_adjustable_target

        base_keep_sum = 0
        for group in adjustable_groups:
            raw_keep = group['count'] * applied_ratio
            base_keep = int(math.floor(raw_keep))
            if (
                policy.get('keep_at_least_one_when_nonzero')
                and group['count'] > 0
                and applied_ratio > 0
                and base_keep == 0
            ):
                base_keep = 1
            base_keep = min(group['count'], max(0, base_keep))
            group['keep'] = base_keep
            group['remainder'] = raw_keep - base_keep
            base_keep_sum += base_keep

        if base_keep_sum > keep_adjustable_target:
            adjustable_groups.sort(key=lambda g: (g['remainder'], g['count']), reverse=False)
            overflow = base_keep_sum - keep_adjustable_target
            for group in adjustable_groups:
                if overflow <= 0:
                    break
                if group['keep'] <= 0:
                    continue
                reducible = min(group['keep'], overflow)
                group['keep'] -= reducible
                overflow -= reducible

        if base_keep_sum < keep_adjustable_target:
            delta = keep_adjustable_target - base_keep_sum
            adjustable_groups.sort(key=lambda g: (g['remainder'], g['count']), reverse=True)
            for group in adjustable_groups:
                if delta <= 0:
                    break
                room = group['count'] - group['keep']
                if room <= 0:
                    continue
                add = min(room, delta)
                group['keep'] += add
                delta -= add

        for group in adjustable_groups:
            if group['rank'] in (1, 2, 3, 4):
                group['row'][f"level_{group['rank']}_keep"] = int(group['keep'])
            else:
                group['row']['other_keep'] = int(group['keep'])

    total_after = 0
    for row in preview_items:
        keep_total = (
            int(row['level_1_keep']) +
            int(row['level_2_keep']) +
            int(row['level_3_keep']) +
            int(row['level_4_keep']) +
            int(row.get('other_keep') or 0)
        )
        row['keep_total'] = keep_total
        row['cut_total'] = max(0, int(row['source_total']) - keep_total)
        row['cut_ratio'] = round((row['cut_total'] * 100.0 / row['source_total']), 2) if row['source_total'] > 0 else 0.0
        for rank in (1, 2, 3, 4):
            before = int(row[f'level_{rank}_before'])
            keep = int(row[f'level_{rank}_keep'])
            row[f'level_{rank}_cut'] = max(0, before - keep)
        row['other_cut'] = max(0, int(row.get('other_before') or 0) - int(row.get('other_keep') or 0))
        total_after += keep_total

    preview_items.sort(key=lambda row: (-int(row.get('source_total') or 0), str(row.get('source_label') or '')))
    low_bound = max(0, max_target_rows - tolerance_rows)
    high_bound = max_target_rows + tolerance_rows
    if bool(policy.get('enabled')):
        if low_bound <= total_after <= high_bound:
            status = 'in_range'
        elif total_after < low_bound:
            status = 'below_range'
            warning_message = '结果低于目标区间下限，可提高 MAX 或手动提高 C。'
        else:
            status = 'above_range'
            warning_message = '结果高于目标区间上限，可降低 MAX 或手动降低 C。'

    return {
        'policy': policy,
        'status': status,
        'warning_message': warning_message,
        'summary': {
            'total_before': total_before,
            'total_after': total_after,
            'target_mode': 'max_plus_tolerance',
            'target_min': low_bound,
            'target_max': high_bound,
            'target_preferred': max_target_rows,
            'max_target_rows': max_target_rows,
            'tolerance_rows': tolerance_rows,
            'auto_ratio': round(auto_ratio, 6),
            'applied_ratio': round(applied_ratio, 6),
            'manual_ratio_enabled': manual_ratio_enabled,
            'manual_ratio': round(manual_ratio, 6),
            'target_min_percent': float(policy.get('target_min_percent') or 0.0),
            'target_max_percent': float(policy.get('target_max_percent') or 0.0),
            'target_preferred_percent': float(policy.get('preferred_target_percent') or 0.0),
            'small_source_threshold': small_threshold,
            'small_total': small_total,
            'large_total': large_total,
            'fixed_total': fixed_total,
            'adjustable_total': adjustable_total
        },
        'items': preview_items
    }


def _get_assets_total_count():
    """返回 assets 总记录数。"""
    conn = None
    try:
        conn = sqlite3.connect(DatabaseConfig.ASSETS_DB)
        cursor = conn.cursor()
        cursor.execute('SELECT COUNT(*) FROM assets')
        return int(cursor.fetchone()[0] or 0)
    except Exception:
        return 0
    finally:
        if conn:
            conn.close()


def _get_merge_results_total_count():
    """返回 merge_results 总记录数。"""
    conn = None
    try:
        conn = sqlite3.connect(DatabaseConfig.MERGE_RESULTS_DB)
        cursor = conn.cursor()
        cursor.execute('SELECT COUNT(*) FROM merge_results')
        return int(cursor.fetchone()[0] or 0)
    except Exception:
        return 0
    finally:
        if conn:
            conn.close()


def _quote_sql_identifier(identifier):
    """对 SQLite 标识符加引号，避免特殊字符影响 SQL。"""
    return '"' + str(identifier).replace('"', '""') + '"'


def _build_sqlite_json_path(field_name):
    safe_field_name = str(field_name or '').replace('\\', '\\\\').replace('"', '\\"')
    return f'$."{safe_field_name}"'


def _resolve_merge_results_source_column(data_columns):
    columns = [str(col or '').strip() for col in (data_columns or []) if str(col or '').strip()]
    if not columns:
        return None
    if '数据源名称' in columns:
        return '数据源名称'
    for name in columns:
        if '数据源' in name and '名称' in name:
            return name
    if len(columns) >= 12:
        return columns[11]
    return None


def _resolve_merge_results_level_column(data_columns):
    columns = [str(col or '').strip() for col in (data_columns or []) if str(col or '').strip()]
    if not columns:
        return None
    if '字段数据分级' in columns:
        return '字段数据分级'
    for name in columns:
        if '字段' in name and '分级' in name:
            return name
    if len(columns) >= 11:
        return columns[10]
    return None


def _parse_merge_level_rank(level_value):
    text = '' if level_value is None else str(level_value).strip()
    if not text:
        return None
    matched = MERGE_LEVEL_RANK_PATTERN.search(text)
    if not matched:
        return None
    try:
        rank = int(matched.group(1))
    except (TypeError, ValueError):
        return None
    return rank if rank in (1, 2, 3, 4) else None


def _collect_merge_export_scope_source_level_stats():
    """汇总 merge_results 的“数据源名称 × 字段数据分级(1-4级)”统计。"""
    conn = None
    try:
        conn = sqlite3.connect(DatabaseConfig.MERGE_RESULTS_DB)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        cursor.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name='merge_results'"
        )
        if cursor.fetchone() is None:
            return {
                'summary': {
                    'total_rows': 0,
                    'source_count': 0,
                    'non_empty_source_count': 0,
                    'empty_source_rows': 0,
                    'level_totals': {
                        'level_1_count': 0,
                        'level_2_count': 0,
                        'level_3_count': 0,
                        'level_4_count': 0,
                        'other_count': 0,
                        'level_1_ratio': 0.0,
                        'level_2_ratio': 0.0,
                        'level_3_ratio': 0.0,
                        'level_4_ratio': 0.0,
                        'other_ratio': 0.0
                    }
                },
                'source_column': '数据源名称',
                'level_column': '字段数据分级',
                'items': []
            }

        cursor.execute('PRAGMA table_info(merge_results)')
        columns_info = cursor.fetchall()
        column_names = [str(row['name']) for row in columns_info]
        data_columns = [name for name in column_names if name.lower() != 'id']
        use_json_format = 'json_data' in column_names and len(column_names) <= 3

        source_column = _resolve_merge_results_source_column(data_columns)
        level_column = _resolve_merge_results_level_column(data_columns)
        if not source_column:
            raise ValueError('合并结果表缺少“数据源名称”列，无法统计。')
        if not level_column:
            raise ValueError('合并结果表缺少“字段数据分级”列，无法统计。')

        if use_json_format:
            source_expr = 'NULLIF(TRIM(COALESCE(json_extract(json_data, ?), "")), "")'
            level_expr = 'NULLIF(TRIM(COALESCE(json_extract(json_data, ?), "")), "")'
            expr_params = [
                _build_sqlite_json_path(source_column),
                _build_sqlite_json_path(level_column)
            ]
        else:
            source_expr = f'NULLIF(TRIM(COALESCE({_quote_sql_identifier(source_column)}, "")), "")'
            level_expr = f'NULLIF(TRIM(COALESCE({_quote_sql_identifier(level_column)}, "")), "")'
            expr_params = []

        cursor.execute('SELECT COUNT(*) AS total_rows FROM merge_results')
        total_rows = int(cursor.fetchone()['total_rows'] or 0)

        cursor.execute(
            f'''
            SELECT source_name, level_name, COUNT(*) AS count
            FROM (
                SELECT {source_expr} AS source_name, {level_expr} AS level_name
                FROM merge_results
            ) merged_rows
            GROUP BY source_name, level_name
            ''',
            expr_params
        )
        grouped_rows = cursor.fetchall()

        source_map = {}
        level_totals = {
            'level_1_count': 0,
            'level_2_count': 0,
            'level_3_count': 0,
            'level_4_count': 0,
            'other_count': 0
        }

        for row in grouped_rows:
            count = int(row['count'] or 0)
            if count <= 0:
                continue

            source_name = row['source_name']
            is_empty_source = source_name is None
            source_value = '' if is_empty_source else str(source_name)
            source_item = source_map.get(source_value)
            if source_item is None:
                source_item = {
                    'source_name': source_value,
                    'source_label': '空值/未填写' if is_empty_source else source_value,
                    'is_empty_source': is_empty_source,
                    'total_count': 0,
                    'level_1_count': 0,
                    'level_2_count': 0,
                    'level_3_count': 0,
                    'level_4_count': 0,
                    'other_count': 0
                }
                source_map[source_value] = source_item

            source_item['total_count'] += count
            rank = _parse_merge_level_rank(row['level_name'])
            if rank in (1, 2, 3, 4):
                key = f'level_{rank}_count'
                source_item[key] += count
                level_totals[key] += count
            else:
                source_item['other_count'] += count
                level_totals['other_count'] += count

        items = list(source_map.values())
        items.sort(key=lambda item: (-int(item.get('total_count') or 0), str(item.get('source_label') or '')))

        for item in items:
            source_total = int(item.get('total_count') or 0)
            for rank in (1, 2, 3, 4):
                count_key = f'level_{rank}_count'
                ratio_key = f'level_{rank}_ratio'
                ratio = (item[count_key] * 100.0 / source_total) if source_total > 0 else 0.0
                item[ratio_key] = round(ratio, 2)
            item['other_ratio'] = round((item['other_count'] * 100.0 / source_total), 2) if source_total > 0 else 0.0

        non_empty_source_count = sum(1 for item in items if not item.get('is_empty_source'))
        empty_source_rows = sum(int(item.get('total_count') or 0) for item in items if item.get('is_empty_source'))

        level_totals_payload = {
            **level_totals,
            'level_1_ratio': round(level_totals['level_1_count'] * 100.0 / total_rows, 2) if total_rows > 0 else 0.0,
            'level_2_ratio': round(level_totals['level_2_count'] * 100.0 / total_rows, 2) if total_rows > 0 else 0.0,
            'level_3_ratio': round(level_totals['level_3_count'] * 100.0 / total_rows, 2) if total_rows > 0 else 0.0,
            'level_4_ratio': round(level_totals['level_4_count'] * 100.0 / total_rows, 2) if total_rows > 0 else 0.0,
            'other_ratio': round(level_totals['other_count'] * 100.0 / total_rows, 2) if total_rows > 0 else 0.0
        }

        return {
            'summary': {
                'total_rows': total_rows,
                'source_count': len(items),
                'non_empty_source_count': non_empty_source_count,
                'empty_source_rows': empty_source_rows,
                'level_totals': level_totals_payload
            },
            'source_column': source_column,
            'level_column': level_column,
            'items': items
        }
    finally:
        if conn:
            conn.close()


@common_bp.route('/api/export-scope-config', methods=['GET'])
def get_export_scope_config():
    """获取填报导出整体筛选配置。"""
    try:
        return jsonify({
            'success': True,
            'config': _load_export_scope_config(),
            'source_totals': {
                'assets': _get_assets_total_count(),
                'merge_results': _get_merge_results_total_count()
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/export-scope-config', methods=['POST'])
def save_export_scope_config():
    """保存填报导出整体筛选配置。"""
    try:
        payload = request.get_json(silent=True) or {}
        normalized_config = _save_export_scope_config(payload.get('config') or {})
        return jsonify({
            'success': True,
            'config': normalized_config,
            'source_totals': {
                'assets': _get_assets_total_count(),
                'merge_results': _get_merge_results_total_count()
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def is_path_within(base_dir, target_path):
    """使用 commonpath 做真实目录边界判断，避免同前缀路径绕过。"""
    try:
        abs_base_dir = os.path.normcase(os.path.abspath(base_dir))
        abs_target_path = os.path.normcase(os.path.abspath(target_path))
        return os.path.commonpath([abs_base_dir, abs_target_path]) == abs_base_dir
    except ValueError:
        return False


def reject_cross_site_delete_request():
    """拦截明显来自外部站点的删除请求，降低交付态被浏览器跨站触发的风险。"""
    origin = str(request.headers.get('Origin') or '').strip()
    sec_fetch_site = str(request.headers.get('Sec-Fetch-Site') or '').strip().lower()

    if sec_fetch_site == 'cross-site':
        return jsonify({'error': '跨站删除请求已被拒绝'}), 403

    if not origin:
        return None

    origin_netloc = urlparse(origin).netloc.lower()
    request_netloc = urlparse(request.host_url).netloc.lower()

    if origin_netloc and origin_netloc != request_netloc:
        return jsonify({'error': '跨站删除请求已被拒绝'}), 403

    return None


def get_project_file_base_dirs(file_type):
    """返回工程文件类别允许访问的基础目录列表。"""
    export_dirs = [
        PathConfig.YEZHI_EXPORT_DIR,
        PathConfig.SMC_EXPORT_DIR,
        PathConfig.XINAN_EXPORT_DIR,
        PathConfig.EXPORT_FOLDER,
    ]

    type_dirs = {
        'templates': [PathConfig.TEMPLATES_DIR],
        'template': [PathConfig.TEMPLATES_DIR],
        'databases': [os.path.join(Config.BASE_DIR, 'data')],
        'data_process': [os.path.join(Config.BASE_DIR, 'temp_data_process')],
        'merge_results': [
            os.path.join(Config.BASE_DIR, 'DataFiles', 'merge_results'),
            PathConfig.EXPORT_FOLDER,
            PathConfig.YEZHI_EXPORT_DIR,
        ],
        'assets': [
            os.path.join(Config.BASE_DIR, 'DataFiles', 'assets'),
            PathConfig.UPLOAD_FOLDER,
            PathConfig.EXPORT_FOLDER,
        ],
        'exports': export_dirs,
        'export': export_dirs,
    }

    dirs = type_dirs.get(file_type)
    if not dirs:
        return None

    normalized = []
    for directory in dirs:
        abs_dir = os.path.abspath(directory)
        if abs_dir not in normalized:
            normalized.append(abs_dir)
    return normalized


def _normalize_business_system_name_mapping_pairs(raw_mappings):
    """将不同存储结构统一为 {source: target} 形式。"""
    normalized = {}

    if isinstance(raw_mappings, dict):
        items = raw_mappings.items()
        for source, target in items:
            source_text = str(source).strip()
            target_text = str(target).strip()
            if source_text and target_text:
                normalized[source_text] = target_text
        return normalized

    if isinstance(raw_mappings, list):
        for item in raw_mappings:
            if not isinstance(item, dict):
                continue
            if not item.get('enabled', True):
                continue
            source_text = str(item.get('source', '')).strip()
            target_text = str(item.get('target', '')).strip()
            if source_text and target_text:
                normalized[source_text] = target_text

    return normalized


def _build_business_system_name_mapping_public_entry(mappings, existing_config=None):
    entry = copy.deepcopy(existing_config) if isinstance(existing_config, dict) else {}
    entry['name'] = '公共配置'
    entry['description'] = '用于将合并结果表中的业务系统名称转换为数据概览表中的标准名称'
    entry['version'] = entry.get('version') or '1.0'
    entry['last_updated'] = datetime.now().strftime('%Y-%m-%d')
    entry['total_count'] = len(mappings)
    entry['mappings'] = copy.deepcopy(mappings)
    return entry


def _count_public_mapping_items(raw_mappings):
    """统计公共映射项数量，兼容 list/dict 两种结构。"""
    if isinstance(raw_mappings, list):
        return len([item for item in raw_mappings if isinstance(item, dict)])
    if isinstance(raw_mappings, dict):
        return len(_normalize_business_system_name_mapping_pairs(raw_mappings))
    return 0


def _build_public_mapping_registry_entry(config_key, existing_config=None):
    """根据注册表和独立存储文件构建公共映射条目。"""
    registry_entry = copy.deepcopy(_build_public_config_registry().get(config_key, {}))
    entry = copy.deepcopy(existing_config) if isinstance(existing_config, dict) else {}

    storage_file = registry_entry.get('storage_file')
    if storage_file and os.path.exists(storage_file):
        file_payload = _load_json_config_payload(storage_file, {})
        if isinstance(file_payload, dict):
            entry.update(copy.deepcopy(file_payload))

    raw_mappings = entry.get('mappings', [])
    if not isinstance(raw_mappings, (list, dict)):
        raw_mappings = []

    entry['key'] = config_key
    entry['title'] = registry_entry.get('title') or entry.get('title') or '公共配置'
    entry['name'] = registry_entry.get('name') or entry.get('name') or config_key
    entry['description'] = registry_entry.get('description') or entry.get('description') or ''
    entry['version'] = entry.get('version') or '1.0'
    entry['last_updated'] = entry.get('last_updated') or datetime.now().strftime('%Y-%m-%d')
    entry['mappings'] = copy.deepcopy(raw_mappings)
    entry['total_count'] = _count_public_mapping_items(raw_mappings)
    return entry


def _load_mapping_config_payload():
    try:
        with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8-sig') as f:
            return json.load(f)
    except Exception:
        return {'reports': {}, 'public_config': {}}


def _save_mapping_config_payload(payload):
    with open(MAPPING_CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def _load_logic_rules_payload():
    try:
        with open(LOGIC_RULES_FILE, 'r', encoding='utf-8-sig') as f:
            payload = json.load(f)
        return payload if isinstance(payload, dict) else {'rules': {}}
    except Exception:
        return {'rules': {}}


def _save_logic_rules_payload(payload):
    with open(LOGIC_RULES_FILE, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def _load_json_config_payload(file_path, default_payload):
    """读取 JSON 配置文件，失败时返回默认值副本。"""
    fallback = copy.deepcopy(default_payload)
    try:
        with open(file_path, 'r', encoding='utf-8-sig') as f:
            payload = json.load(f)
        return payload if isinstance(payload, dict) else fallback
    except Exception:
        return fallback


def _normalize_public_config_center_items(raw_items):
    """统一过滤公共配置中心条目列表中的非法项。"""
    if not isinstance(raw_items, list):
        return []

    return [
        copy.deepcopy(item)
        for item in raw_items
        if isinstance(item, dict)
    ]


def _normalize_public_config_field_definitions(raw_fields):
    """标准化公共配置字段定义列表。"""
    if not isinstance(raw_fields, list):
        return []

    normalized_fields = []
    for field in raw_fields:
        if not isinstance(field, dict):
            continue
        field_name = str(field.get('name', '') or '').strip()
        if not field_name:
            continue
        normalized_fields.append(copy.deepcopy(field))
    return normalized_fields


def _normalize_public_config_type_definition(raw_item, base_registry_entry=None):
    """将定义层配置统一收敛成通用类型注册结构。"""
    base_entry = copy.deepcopy(base_registry_entry) if isinstance(base_registry_entry, dict) else {}
    raw_entry = copy.deepcopy(raw_item) if isinstance(raw_item, dict) else {}

    key = str(raw_entry.get('key') or base_entry.get('key') or '').strip()
    if not key:
        return None

    schema = copy.deepcopy(base_entry.get('schema')) if isinstance(base_entry.get('schema'), dict) else {}
    if isinstance(raw_entry.get('schema'), dict):
        schema.update(copy.deepcopy(raw_entry.get('schema')))

    config = copy.deepcopy(base_entry.get('config')) if isinstance(base_entry.get('config'), dict) else {}
    if isinstance(raw_entry.get('config'), dict):
        config.update(copy.deepcopy(raw_entry.get('config')))

    storage = copy.deepcopy(base_entry.get('storage')) if isinstance(base_entry.get('storage'), dict) else {}
    if isinstance(raw_entry.get('storage'), dict):
        storage.update(copy.deepcopy(raw_entry.get('storage')))
    storage_type = (
        raw_entry.get('storage_type')
        or storage.get('type')
        or base_entry.get('storage_type')
        or 'json_file'
    )
    storage['type'] = storage_type
    storage_file = (
        storage.get('file')
        or raw_entry.get('storage_file')
        or config.get('storage_file')
        or base_entry.get('storage_file')
    )
    if storage_file:
        storage['file'] = storage_file

    capabilities = (
        copy.deepcopy(base_entry.get('capabilities'))
        if isinstance(base_entry.get('capabilities'), dict)
        else {}
    )
    if isinstance(raw_entry.get('capabilities'), dict):
        capabilities.update(copy.deepcopy(raw_entry.get('capabilities')))
    capabilities.setdefault(
        'detail',
        bool(base_entry.get('supports_detail', raw_entry.get('supports_detail', False)))
    )
    capabilities.setdefault(
        'crud',
        bool(base_entry.get('supports_crud', raw_entry.get('supports_crud', False)))
    )
    capabilities.setdefault(
        'import_export',
        bool(
            base_entry.get(
                'supports_import_export',
                raw_entry.get('supports_import_export', schema.get('supports_import_export', False))
            )
        )
    )
    capabilities.setdefault(
        'frontend_edit',
        bool(schema.get('supports_frontend_edit', raw_entry.get('supports_frontend_edit', False)))
    )
    capabilities.setdefault(
        'template_download',
        bool(
            raw_entry.get('template_api')
            or base_entry.get('template_api')
            or (isinstance(raw_entry.get('apis'), dict) and raw_entry.get('apis', {}).get('template_api'))
        )
    )

    apis = copy.deepcopy(base_entry.get('apis')) if isinstance(base_entry.get('apis'), dict) else {}
    if isinstance(raw_entry.get('apis'), dict):
        apis.update(copy.deepcopy(raw_entry.get('apis')))
    for api_key in ('detail_api', 'save_api', 'template_api', 'import_api', 'export_api'):
        flat_value = raw_entry.get(api_key, base_entry.get(api_key))
        if flat_value:
            apis[api_key] = flat_value

    fields = _normalize_public_config_field_definitions(
        raw_entry.get('fields', schema.get('fields', base_entry.get('fields', [])))
    )
    config_type = (
        raw_entry.get('config_type')
        or raw_entry.get('type')
        or base_entry.get('config_type')
        or base_entry.get('type')
        or 'generic'
    )

    affects_export = raw_entry.get('affects_export')
    if affects_export is None:
        affects_export = config.get('affects_export')
    if affects_export is None:
        affects_export = base_entry.get('affects_export', False)
    affects_export = bool(affects_export)

    contract_role = (
        raw_entry.get('contract_role')
        or config.get('contract_role')
        or base_entry.get('contract_role')
        or ('runtime' if affects_export else 'reserve')
    )
    effect_summary = (
        raw_entry.get('effect_summary')
        or config.get('effect_summary')
        or base_entry.get('effect_summary')
        or ('保存后会直接影响已接线报表导出' if contract_role == 'runtime' else '当前仅为预留或储备配置，不直接影响导出')
    )
    editor_mode = (
        raw_entry.get('editor_mode')
        or schema.get('editor_mode')
        or base_entry.get('editor_mode')
        or 'generic'
    )

    normalized_entry = copy.deepcopy(raw_entry)
    normalized_entry['key'] = key
    normalized_entry['title'] = raw_entry.get('title') or base_entry.get('title') or '公共配置'
    normalized_entry['name'] = raw_entry.get('name') or base_entry.get('name') or key
    normalized_entry['description'] = raw_entry.get('description') or base_entry.get('description') or ''
    normalized_entry['type'] = config_type
    normalized_entry['config_type'] = config_type
    normalized_entry['category'] = raw_entry.get('category') or base_entry.get('category') or 'general'
    normalized_entry['enabled'] = bool(raw_entry.get('enabled', base_entry.get('enabled', True)))
    normalized_entry['version'] = raw_entry.get('version') or base_entry.get('version') or '1.0'
    normalized_entry['contract_role'] = contract_role
    normalized_entry['affects_export'] = affects_export
    normalized_entry['effect_summary'] = effect_summary
    normalized_entry['editor_mode'] = editor_mode
    normalized_entry['storage_type'] = storage_type
    normalized_entry['storage_file'] = storage_file
    normalized_entry['storage'] = storage
    normalized_entry['schema'] = schema
    normalized_entry['config'] = config
    normalized_entry['fields'] = fields
    normalized_entry['primary_field'] = (
        raw_entry.get('primary_field')
        or base_entry.get('primary_field')
        or (fields[0].get('name') if fields else '')
    )
    normalized_entry['display_fields'] = (
        copy.deepcopy(raw_entry.get('display_fields'))
        if isinstance(raw_entry.get('display_fields'), list)
        else copy.deepcopy(base_entry.get('display_fields'))
        if isinstance(base_entry.get('display_fields'), list)
        else [field.get('name') for field in fields if field.get('visible_in_list', True)]
    )
    normalized_entry['capabilities'] = capabilities
    normalized_entry['apis'] = apis
    normalized_entry['detail_api'] = apis.get('detail_api')
    normalized_entry['save_api'] = apis.get('save_api')
    normalized_entry['template_api'] = apis.get('template_api')
    normalized_entry['import_api'] = apis.get('import_api')
    normalized_entry['export_api'] = apis.get('export_api')
    normalized_entry['supports_detail'] = bool(capabilities.get('detail', False))
    normalized_entry['supports_crud'] = bool(capabilities.get('crud', False))
    normalized_entry['supports_import_export'] = bool(capabilities.get('import_export', False))
    normalized_entry['supports_frontend_edit'] = bool(capabilities.get('frontend_edit', False))
    normalized_entry['template_filename'] = (
        raw_entry.get('template_filename')
        or config.get('template_filename')
        or schema.get('template_filename')
        or ''
    )
    normalized_entry['export_filename'] = (
        raw_entry.get('export_filename')
        or config.get('export_filename')
        or schema.get('export_filename')
        or ''
    )
    return normalized_entry


def _load_public_config_definitions_payload():
    """读取公共配置定义层文件。"""
    payload = _load_json_config_payload(
        PUBLIC_CONFIG_DEFINITIONS_FILE,
        {
            'version': '1.0',
            'last_updated': '',
            'notes': '',
            'items': [],
        }
    )
    payload['items'] = _normalize_public_config_center_items(payload.get('items'))
    return payload


def _save_public_config_definitions_payload(payload):
    """保存公共配置定义层文件。"""
    normalized_payload = copy.deepcopy(payload) if isinstance(payload, dict) else {}
    normalized_payload['version'] = str(normalized_payload.get('version') or '2.0')
    normalized_payload['last_updated'] = datetime.now().strftime('%Y-%m-%d')
    normalized_payload['notes'] = str(normalized_payload.get('notes') or '')
    normalized_payload['items'] = _normalize_public_config_center_items(normalized_payload.get('items'))

    os.makedirs(os.path.dirname(PUBLIC_CONFIG_DEFINITIONS_FILE), exist_ok=True)
    with open(PUBLIC_CONFIG_DEFINITIONS_FILE, 'w', encoding='utf-8') as f:
        json.dump(normalized_payload, f, ensure_ascii=False, indent=2)

    return normalized_payload


def _normalize_public_config_type_save_fields(raw_fields):
    """校验并标准化前端提交的字段定义。"""
    if not isinstance(raw_fields, list) or not raw_fields:
        raise ValueError('至少需要配置一个字段')

    allowed_types = {'text', 'textarea', 'number', 'boolean', 'select'}
    normalized_fields = []
    seen_names = set()
    for index, raw_field in enumerate(raw_fields, start=1):
        if not isinstance(raw_field, dict):
            continue

        field_name = str(raw_field.get('name') or '').strip()
        if not re.match(r'^[A-Za-z][A-Za-z0-9_]*$', field_name):
            raise ValueError(f'第 {index} 个字段标识无效，只能使用字母、数字和下划线，且需以字母开头')

        normalized_name = field_name.lower()
        if normalized_name in seen_names:
            raise ValueError(f'字段标识重复: {field_name}')
        seen_names.add(normalized_name)

        field_type = str(raw_field.get('type') or 'text').strip().lower()
        if field_type not in allowed_types:
            raise ValueError(f'字段 {field_name} 的类型不支持')

        normalized_field = {
            'name': normalized_name,
            'label': str(raw_field.get('label') or field_name).strip(),
            'type': field_type,
            'required': bool(raw_field.get('required', False)),
            'searchable': bool(raw_field.get('searchable', field_type != 'boolean')),
            'editable': bool(raw_field.get('editable', True)),
            'visible_in_list': bool(raw_field.get('visible_in_list', True)),
            'width': int(raw_field.get('width') or 180),
        }

        placeholder = str(raw_field.get('placeholder') or '').strip()
        if placeholder:
            normalized_field['placeholder'] = placeholder

        description = str(raw_field.get('description') or '').strip()
        if description:
            normalized_field['description'] = description

        if 'default_value' in raw_field:
            normalized_field['default_value'] = _coerce_public_config_field_value(
                normalized_field,
                raw_field.get('default_value')
            )
        elif field_type == 'boolean':
            normalized_field['default_value'] = True

        if field_type == 'select':
            raw_options = raw_field.get('options')
            if not isinstance(raw_options, list):
                raw_options = []
            normalized_field['options'] = [
                copy.deepcopy(option)
                for option in raw_options
                if isinstance(option, (str, int, float, dict))
            ]

        normalized_fields.append(normalized_field)

    if not normalized_fields:
        raise ValueError('至少需要配置一个有效字段')

    return normalized_fields


def _build_public_config_type_default_apis(config_key):
    """构建通用公共配置接口地址。"""
    return {
        'detail_api': f'/api/public-configs/{config_key}/records',
        'save_api': f'/api/public-configs/{config_key}/records',
        'template_api': f'/api/public-configs/{config_key}/template',
        'import_api': f'/api/public-configs/{config_key}/import',
        'export_api': f'/api/public-configs/{config_key}/export',
    }


def _build_public_config_type_default_schema(config_key, editor_mode, fields, capabilities):
    """构建映射表类公共配置的默认 schema。"""
    field_labels = {
        field.get('name'): field.get('label') or field.get('name')
        for field in fields
        if field.get('name')
    }
    source_field = next((field for field in fields if field.get('name') == 'source'), fields[0])
    target_field = next((field for field in fields if field.get('name') == 'target'), fields[min(1, len(fields) - 1)])
    enabled_field = next((field for field in fields if field.get('type') == 'boolean'), None)

    field_meta = {
        'source_label': source_field.get('label') or source_field.get('name') or '原值',
        'target_label': target_field.get('label') or target_field.get('name') or '目标值',
    }
    if enabled_field:
        field_meta['enabled_label'] = enabled_field.get('label') or enabled_field.get('name')

    return {
        'entry_type': 'source_target_mapping' if {'source', 'target'}.issubset(field_labels.keys()) else 'custom_mapping',
        'supports_enable_flag': bool(enabled_field),
        'supports_import_export': bool(capabilities.get('import_export', True)),
        'supports_frontend_edit': bool(capabilities.get('frontend_edit', True)),
        'editor_mode': editor_mode or config_key,
        'field_meta': field_meta,
        'import_columns': [
            field.get('label') or field.get('name')
            for field in fields
            if field.get('name')
        ],
    }


def _normalize_public_config_type_definition_for_save(raw_item, existing_item=None):
    """将新增/更新请求标准化为可落盘的公共配置类型定义。"""
    raw_entry = copy.deepcopy(raw_item) if isinstance(raw_item, dict) else {}
    existing_entry = copy.deepcopy(existing_item) if isinstance(existing_item, dict) else {}

    config_key = str(raw_entry.get('key') or existing_entry.get('key') or '').strip()
    if not re.match(r'^[a-z][a-z0-9_]*$', config_key):
        raise ValueError('类型标识只能使用小写字母、数字和下划线，且需以小写字母开头')

    fields = _normalize_public_config_type_save_fields(raw_entry.get('fields', existing_entry.get('fields')))
    display_name = str(raw_entry.get('name') or existing_entry.get('name') or config_key).strip()
    category = str(raw_entry.get('category') or existing_entry.get('category') or 'custom').strip()
    config_type = str(raw_entry.get('config_type') or raw_entry.get('type') or existing_entry.get('config_type') or existing_entry.get('type') or 'mapping_table').strip()
    affects_export = bool(raw_entry.get('affects_export', existing_entry.get('affects_export', False)))
    contract_role = str(
        raw_entry.get('contract_role')
        or existing_entry.get('contract_role')
        or ('runtime' if affects_export else 'reserve')
    ).strip()
    editor_mode = str(raw_entry.get('editor_mode') or existing_entry.get('editor_mode') or config_key).strip()

    raw_capabilities = copy.deepcopy(existing_entry.get('capabilities')) if isinstance(existing_entry.get('capabilities'), dict) else {}
    if isinstance(raw_entry.get('capabilities'), dict):
        raw_capabilities.update(copy.deepcopy(raw_entry.get('capabilities')))
    capabilities = {
        'detail': bool(raw_capabilities.get('detail', True)),
        'crud': bool(raw_capabilities.get('crud', True)),
        'import_export': bool(raw_capabilities.get('import_export', True)),
        'frontend_edit': bool(raw_capabilities.get('frontend_edit', True)),
        'template_download': bool(raw_capabilities.get('template_download', True)),
    }

    storage = copy.deepcopy(existing_entry.get('storage')) if isinstance(existing_entry.get('storage'), dict) else {}
    if isinstance(raw_entry.get('storage'), dict):
        storage.update(copy.deepcopy(raw_entry.get('storage')))
    storage_file = (
        storage.get('file')
        or raw_entry.get('storage_file')
        or (raw_entry.get('config') or {}).get('storage_file')
        or (existing_entry.get('config') or {}).get('storage_file')
        or f'config/{config_key}.json'
    )
    resolved_storage_file = storage_file if os.path.isabs(storage_file) else os.path.join(Config.BASE_DIR, storage_file)
    abs_base_dir = os.path.normcase(os.path.abspath(Config.BASE_DIR))
    abs_storage_file = os.path.normcase(os.path.abspath(resolved_storage_file))
    if os.path.commonpath([abs_base_dir, abs_storage_file]) != abs_base_dir:
        raise ValueError('存储文件路径超出项目目录')
    storage['type'] = storage.get('type') or raw_entry.get('storage_type') or existing_entry.get('storage_type') or 'json_file'
    storage['file'] = storage_file
    storage['record_path'] = storage.get('record_path') or 'mappings'

    raw_apis = copy.deepcopy(existing_entry.get('apis')) if isinstance(existing_entry.get('apis'), dict) else {}
    raw_apis.update(_build_public_config_type_default_apis(config_key))
    if isinstance(raw_entry.get('apis'), dict):
        raw_apis.update(copy.deepcopy(raw_entry.get('apis')))

    raw_config = copy.deepcopy(existing_entry.get('config')) if isinstance(existing_entry.get('config'), dict) else {}
    if isinstance(raw_entry.get('config'), dict):
        raw_config.update(copy.deepcopy(raw_entry.get('config')))
    raw_config['storage_file'] = storage_file
    raw_config['template_filename'] = raw_config.get('template_filename') or f'{display_name}_导入模板.xlsx'
    raw_config['export_filename'] = raw_config.get('export_filename') or f'{display_name}_导出.xlsx'
    raw_config.setdefault('mappings', [])

    primary_field = str(raw_entry.get('primary_field') or existing_entry.get('primary_field') or fields[0].get('name')).strip()
    field_names = {field.get('name') for field in fields}
    if primary_field not in field_names:
        primary_field = fields[0].get('name')

    raw_display_fields = raw_entry.get('display_fields', existing_entry.get('display_fields'))
    if isinstance(raw_display_fields, list):
        display_fields = [
            str(field_name).strip()
            for field_name in raw_display_fields
            if str(field_name).strip() in field_names
        ]
    else:
        display_fields = []
    if not display_fields:
        display_fields = [
            field.get('name')
            for field in fields
            if field.get('visible_in_list', True)
        ]

    schema = copy.deepcopy(existing_entry.get('schema')) if isinstance(existing_entry.get('schema'), dict) else {}
    if isinstance(raw_entry.get('schema'), dict):
        schema.update(copy.deepcopy(raw_entry.get('schema')))
    schema.update(_build_public_config_type_default_schema(config_key, editor_mode, fields, capabilities))

    return {
        'key': config_key,
        'title': str(raw_entry.get('title') or existing_entry.get('title') or '公共配置').strip(),
        'name': display_name,
        'type': config_type,
        'category': category,
        'description': str(raw_entry.get('description') or existing_entry.get('description') or '').strip(),
        'enabled': bool(raw_entry.get('enabled', existing_entry.get('enabled', True))),
        'version': str(raw_entry.get('version') or existing_entry.get('version') or '1.0'),
        'contract_role': contract_role,
        'affects_export': affects_export,
        'effect_summary': str(
            raw_entry.get('effect_summary')
            or existing_entry.get('effect_summary')
            or ('保存后会直接影响已接线报表导出。' if contract_role == 'runtime' else '当前仅作为公共配置储备，具体使用方式可后续决定。')
        ).strip(),
        'editor_mode': editor_mode,
        'storage_type': storage.get('type') or 'json_file',
        'storage': storage,
        'capabilities': capabilities,
        'apis': raw_apis,
        'primary_field': primary_field,
        'display_fields': display_fields,
        'fields': fields,
        'schema': schema,
        'config': raw_config,
    }


def _load_public_config_bindings_payload():
    """读取公共配置绑定层文件。"""
    payload = _load_json_config_payload(
        PUBLIC_CONFIG_BINDINGS_FILE,
        {
            'version': '1.0',
            'last_updated': '',
            'notes': '',
            'items': [],
        }
    )
    payload['items'] = _normalize_public_config_center_items(payload.get('items'))
    return payload


def _save_public_config_bindings_payload(payload):
    """保存公共配置绑定层文件。"""
    normalized_payload = copy.deepcopy(payload) if isinstance(payload, dict) else {}
    normalized_payload['version'] = str(normalized_payload.get('version') or '1.0')
    normalized_payload['last_updated'] = datetime.now().strftime('%Y-%m-%d')
    normalized_payload['notes'] = str(normalized_payload.get('notes') or '')
    normalized_payload['items'] = _normalize_public_config_center_items(
        normalized_payload.get('items')
    )

    os.makedirs(os.path.dirname(PUBLIC_CONFIG_BINDINGS_FILE), exist_ok=True)
    with open(PUBLIC_CONFIG_BINDINGS_FILE, 'w', encoding='utf-8') as f:
        json.dump(normalized_payload, f, ensure_ascii=False, indent=2)

    return normalized_payload


def _normalize_dynamic_config_query_definitions_payload(payload):
    """标准化第三级动态配置查询定义层。"""
    raw_payload = copy.deepcopy(payload) if isinstance(payload, dict) else {}
    if isinstance(raw_payload.get('data'), dict):
        raw_payload = raw_payload.get('data') or {}

    raw_definitions = raw_payload.get('definitions')
    if not isinstance(raw_definitions, dict):
        raw_definitions = raw_payload if isinstance(raw_payload, dict) else {}

    definitions = {}
    for source_type, definition in raw_definitions.items():
        normalized_source_type = str(source_type or '').strip()
        if not normalized_source_type or not isinstance(definition, dict):
            continue
        definitions[normalized_source_type] = copy.deepcopy(definition)

    return {
        'version': str(raw_payload.get('version') or '1.0'),
        'last_updated': str(raw_payload.get('last_updated') or ''),
        'notes': str(raw_payload.get('notes') or ''),
        'definitions': definitions,
    }


def _load_dynamic_config_query_definitions_payload():
    """读取第三级动态配置查询定义层。"""
    payload = _load_json_config_payload(
        DYNAMIC_CONFIG_QUERY_DEFINITIONS_FILE,
        {
            'version': '1.0',
            'last_updated': '',
            'notes': '',
            'definitions': {},
        }
    )
    return _normalize_dynamic_config_query_definitions_payload(payload)


def _save_dynamic_config_query_definitions_payload(payload):
    """保存第三级动态配置查询定义层。"""
    normalized_payload = _normalize_dynamic_config_query_definitions_payload(payload)
    normalized_payload['last_updated'] = datetime.now().strftime('%Y-%m-%d')

    os.makedirs(os.path.dirname(DYNAMIC_CONFIG_QUERY_DEFINITIONS_FILE), exist_ok=True)
    with open(DYNAMIC_CONFIG_QUERY_DEFINITIONS_FILE, 'w', encoding='utf-8') as f:
        json.dump(normalized_payload, f, ensure_ascii=False, indent=2)

    return normalized_payload


def _build_public_config_registry():
    """构建统一公共配置类型注册表。"""
    registry = {
        key: _normalize_public_config_type_definition({'key': key}, value) or copy.deepcopy(value)
        for key, value in PUBLIC_CONFIG_REGISTRY.items()
    }

    payload = _load_public_config_definitions_payload()
    for item in payload.get('items', []):
        config_key = str(item.get('key', '') or '').strip()
        if not config_key:
            continue
        normalized_entry = _normalize_public_config_type_definition(item, registry.get(config_key, {}))
        if normalized_entry:
            registry[config_key] = normalized_entry

    return registry


def _resolve_public_config_registry_entry(config_key):
    registry = _build_public_config_registry()
    entry = registry.get(config_key)
    if not isinstance(entry, dict):
        raise ValueError(f'未找到公共配置类型: {config_key}')
    return registry, copy.deepcopy(entry)


def _resolve_public_config_storage_file(config_key, registry_entry=None):
    entry = registry_entry if isinstance(registry_entry, dict) else _build_public_config_registry().get(config_key, {})
    storage_file = (
        entry.get('storage_file')
        or (entry.get('storage', {}) or {}).get('file')
        or (entry.get('config', {}) or {}).get('storage_file')
    )
    if not storage_file:
        return None
    if os.path.isabs(storage_file):
        return storage_file
    return os.path.join(Config.BASE_DIR, storage_file)


def _default_public_config_storage_payload(config_key, registry_entry):
    return {
        'key': config_key,
        'title': registry_entry.get('title') or '公共配置',
        'name': registry_entry.get('name') or config_key,
        'description': registry_entry.get('description') or '',
        'version': registry_entry.get('version') or '1.0',
        'last_updated': datetime.now().strftime('%Y-%m-%d'),
        'total_count': 0,
        'mappings': []
    }


def _coerce_public_config_field_value(field, raw_value):
    field_type = str((field or {}).get('type') or 'text').strip().lower()
    default_value = (field or {}).get('default_value')

    if raw_value is None or raw_value == '':
        if field_type == 'boolean':
            return bool(default_value) if default_value is not None else False
        if field_type == 'number':
            return default_value if default_value is not None else ''
        return default_value if default_value is not None else ''

    if field_type == 'boolean':
        if isinstance(raw_value, str):
            value = raw_value.strip().lower()
            if value in ('1', 'true', 'yes', 'y', 'on', '是'):
                return True
            if value in ('0', 'false', 'no', 'n', 'off', '否'):
                return False
        return bool(raw_value)

    if field_type == 'number':
        try:
            numeric_value = float(raw_value)
        except (TypeError, ValueError):
            return raw_value
        return int(numeric_value) if numeric_value.is_integer() else numeric_value

    return str(raw_value).strip()


def _get_public_config_fields(registry_entry):
    fields = registry_entry.get('fields')
    if isinstance(fields, list) and fields:
        return _normalize_public_config_field_definitions(fields)
    schema_fields = (registry_entry.get('schema') or {}).get('fields')
    return _normalize_public_config_field_definitions(schema_fields)


def _normalize_public_config_mapping_records(raw_mappings, fields):
    normalized_records = []
    field_list = fields if isinstance(fields, list) else []

    if isinstance(raw_mappings, dict):
        for index, (source, target) in enumerate(raw_mappings.items(), start=1):
            record = {
                'id': index,
                'source': str(source).strip(),
                'target': str(target).strip(),
                'enabled': True
            }
            normalized_records.append(record)
        return normalized_records

    if not isinstance(raw_mappings, list):
        return []

    for index, item in enumerate(raw_mappings, start=1):
        if not isinstance(item, dict):
            continue
        record = {'id': item.get('id', index)}
        for field in field_list:
            field_name = field.get('name')
            if not field_name:
                continue
            record[field_name] = _coerce_public_config_field_value(field, item.get(field_name))
        for key, value in item.items():
            if key == 'id' or any(field.get('name') == key for field in field_list):
                continue
            record[key] = value
        normalized_records.append(record)

    return normalized_records


def _load_public_config_storage_payload(config_key, registry_entry=None):
    entry = registry_entry if isinstance(registry_entry, dict) else _resolve_public_config_registry_entry(config_key)[1]
    storage_file = _resolve_public_config_storage_file(config_key, entry)
    payload = _default_public_config_storage_payload(config_key, entry)
    if storage_file and os.path.exists(storage_file):
        stored_payload = _load_json_config_payload(storage_file, payload)
        if isinstance(stored_payload, dict):
            payload.update(copy.deepcopy(stored_payload))

    field_list = _get_public_config_fields(entry)
    payload['mappings'] = _normalize_public_config_mapping_records(payload.get('mappings', []), field_list)
    payload['key'] = config_key
    payload['title'] = entry.get('title') or payload.get('title') or '公共配置'
    payload['name'] = entry.get('name') or payload.get('name') or config_key
    payload['description'] = entry.get('description') or payload.get('description') or ''
    payload['version'] = payload.get('version') or entry.get('version') or '1.0'
    payload['last_updated'] = payload.get('last_updated') or datetime.now().strftime('%Y-%m-%d')
    payload['total_count'] = len(payload['mappings'])
    return payload


def _save_public_config_storage_payload(config_key, payload, registry_entry=None):
    entry = registry_entry if isinstance(registry_entry, dict) else _resolve_public_config_registry_entry(config_key)[1]
    storage_file = _resolve_public_config_storage_file(config_key, entry)
    if not storage_file:
        raise ValueError('当前公共配置未登记存储文件')

    abs_base_dir = os.path.normcase(os.path.abspath(Config.BASE_DIR))
    abs_storage_file = os.path.normcase(os.path.abspath(storage_file))
    if os.path.commonpath([abs_base_dir, abs_storage_file]) != abs_base_dir:
        raise ValueError('公共配置存储路径超出项目目录')

    normalized_payload = copy.deepcopy(payload) if isinstance(payload, dict) else {}
    normalized_payload['key'] = config_key
    normalized_payload['title'] = entry.get('title') or normalized_payload.get('title') or '公共配置'
    normalized_payload['name'] = entry.get('name') or normalized_payload.get('name') or config_key
    normalized_payload['description'] = entry.get('description') or normalized_payload.get('description') or ''
    normalized_payload['version'] = normalized_payload.get('version') or entry.get('version') or '1.0'
    normalized_payload['last_updated'] = datetime.now().strftime('%Y-%m-%d')
    normalized_payload['mappings'] = _normalize_public_config_mapping_records(
        normalized_payload.get('mappings', []),
        _get_public_config_fields(entry)
    )
    normalized_payload['total_count'] = len(normalized_payload['mappings'])

    os.makedirs(os.path.dirname(storage_file), exist_ok=True)
    with open(storage_file, 'w', encoding='utf-8') as f:
        json.dump(normalized_payload, f, ensure_ascii=False, indent=2)

    return normalized_payload


def _build_public_config_record_payload(record_id, request_data, existing_record, registry_entry):
    field_list = _get_public_config_fields(registry_entry)
    record = copy.deepcopy(existing_record) if isinstance(existing_record, dict) else {}
    record['id'] = record_id

    for field in field_list:
        field_name = field.get('name')
        if not field_name:
            continue
        if field_name in request_data:
            record[field_name] = _coerce_public_config_field_value(field, request_data.get(field_name))
        elif field_name not in record:
            record[field_name] = _coerce_public_config_field_value(field, field.get('default_value'))

    missing_required = []
    for field in field_list:
        if not field.get('required'):
            continue
        field_name = field.get('name')
        value = record.get(field_name)
        if field.get('type') == 'boolean':
            continue
        if value in (None, ''):
            missing_required.append(field.get('label') or field_name)

    if missing_required:
        raise ValueError(f"缺少必填字段: {', '.join(missing_required)}")

    return record


def _get_public_config_mapping_conflict_meta(registry_entry):
    """获取映射冲突校验使用的源字段与目标字段元信息。"""
    field_list = _get_public_config_fields(registry_entry)
    if not field_list:
        return None

    source_field = next(
        (field for field in field_list if str(field.get('name') or '').strip() == 'source'),
        None
    )
    if not source_field:
        primary_field_name = str(registry_entry.get('primary_field') or '').strip()
        source_field = next(
            (field for field in field_list if str(field.get('name') or '').strip() == primary_field_name),
            None
        )
    if not source_field:
        source_field = field_list[0]

    target_field = next(
        (
            field for field in field_list
            if str(field.get('name') or '').strip() == 'target'
        ),
        None
    )
    if not target_field:
        target_field = next(
            (
                field for field in field_list
                if str(field.get('name') or '').strip() != str(source_field.get('name') or '').strip()
                and str(field.get('type') or '').strip().lower() != 'boolean'
            ),
            None
        )

    if not source_field or not target_field:
        return None

    return {
        'source_field_name': str(source_field.get('name') or '').strip(),
        'source_field_label': str(source_field.get('label') or source_field.get('name') or '原名称').strip(),
        'target_field_name': str(target_field.get('name') or '').strip(),
        'target_field_label': str(target_field.get('label') or target_field.get('name') or '映射名称').strip(),
    }


def _validate_public_config_mapping_conflict(candidate_record, existing_records, registry_entry, exclude_record_id=None):
    """校验同一个原名称不能映射到多个不同目标名称。"""
    if not isinstance(candidate_record, dict):
        return

    conflict_meta = _get_public_config_mapping_conflict_meta(registry_entry)
    if not isinstance(conflict_meta, dict):
        return

    source_field_name = conflict_meta.get('source_field_name')
    target_field_name = conflict_meta.get('target_field_name')
    source_field_label = conflict_meta.get('source_field_label') or source_field_name or '原名称'
    target_field_label = conflict_meta.get('target_field_label') or target_field_name or '映射名称'

    source_value = str(candidate_record.get(source_field_name, '') or '').strip()
    target_value = str(candidate_record.get(target_field_name, '') or '').strip()
    if not source_value or not target_value:
        return

    normalized_exclude_id = None if exclude_record_id is None else str(exclude_record_id).strip()
    for record in existing_records if isinstance(existing_records, list) else []:
        if not isinstance(record, dict):
            continue

        record_id = str(record.get('id', '')).strip()
        if normalized_exclude_id is not None and record_id == normalized_exclude_id:
            continue

        existing_source_value = str(record.get(source_field_name, '') or '').strip()
        if existing_source_value != source_value:
            continue

        existing_target_value = str(record.get(target_field_name, '') or '').strip()
        if not existing_target_value or existing_target_value == target_value:
            continue

        raise ValueError(
            f'映射冲突：同一个{source_field_label}“{source_value}”'
            f'已映射到“{existing_target_value}”，不能再映射到“{target_value}”，请修改后重试'
        )


def _build_public_config_excel_fields(registry_entry):
    field_list = _get_public_config_fields(registry_entry)
    if field_list:
        return field_list
    return [
        {'name': 'source', 'label': '原值', 'type': 'text', 'required': True},
        {'name': 'target', 'label': '目标值', 'type': 'text', 'required': True},
        {'name': 'enabled', 'label': '是否启用', 'type': 'boolean', 'required': False, 'default_value': True},
    ]


def _convert_public_config_excel_cell_value(field, value):
    if str((field or {}).get('type') or '').strip().lower() == 'boolean':
        return '是' if value is not False else '否'
    return value


def _extract_public_config_row_value(row, field):
    field_name = str(field.get('name') or '').strip()
    field_label = str(field.get('label') or field_name).strip()
    for candidate in (field_label, field_name):
        if candidate in row:
            return row.get(candidate)
    return None


def _parse_optional_bool_query_value(raw_value):
    """将查询参数解析为可选布尔值。"""
    if raw_value is None:
        return None

    value = str(raw_value).strip().lower()
    if value in ('1', 'true', 'yes', 'y', 'on'):
        return True
    if value in ('0', 'false', 'no', 'n', 'off'):
        return False
    return None


def _collect_public_config_usages(config_key, mapping_config=None):
    config = mapping_config if isinstance(mapping_config, dict) else _load_mapping_config_payload()
    reports = config.get('reports') or {}
    usages = []

    for category, category_reports in reports.items():
        if not isinstance(category_reports, dict):
            continue

        for report_code, report_config in category_reports.items():
            mapping_rules = report_config.get('mapping_rules') or []
            target_columns = []

            for rule in mapping_rules:
                if rule.get('mapping_config') != config_key:
                    continue

                target_column = str(rule.get('target_column', '')).strip()
                if target_column and target_column not in target_columns:
                    target_columns.append(target_column)

            if not target_columns:
                continue

            usages.append({
                'category': category,
                'report_code': report_code,
                'report_name': report_config.get('name') or report_code,
                'target_columns': target_columns,
            })

    usages.sort(key=lambda item: (item['category'], item['report_code']))
    return usages


def _collect_logic_rule_usages(rule_id, mapping_config=None):
    config = mapping_config if isinstance(mapping_config, dict) else _load_mapping_config_payload()
    reports = config.get('reports') or {}
    usages = []

    for category, category_reports in reports.items():
        if not isinstance(category_reports, dict):
            continue

        for report_code, report_config in category_reports.items():
            mapping_rules = report_config.get('mapping_rules') or []
            if not isinstance(mapping_rules, list):
                continue

            for rule in mapping_rules:
                if not isinstance(rule, dict):
                    continue
                if str(rule.get('logic_rule_id') or '').strip() != str(rule_id or '').strip():
                    continue

                usages.append({
                    'category': category,
                    'report_code': report_code,
                    'report_name': report_config.get('name') or report_code,
                    'target_column': str(rule.get('target_column') or '').strip(),
                    'target_name': str(rule.get('target_name') or '').strip(),
                })

    usages.sort(key=lambda item: (item['category'], item['report_code'], item['target_column']))
    return usages


def _collect_all_logic_rule_usage_targets(mapping_config=None):
    config = mapping_config if isinstance(mapping_config, dict) else _load_mapping_config_payload()
    reports = config.get('reports') or {}
    targets = []

    for category, category_reports in reports.items():
        if not isinstance(category_reports, dict):
            continue

        for report_code, report_config in category_reports.items():
            mapping_rules = report_config.get('mapping_rules') or []
            if not isinstance(mapping_rules, list):
                continue

            for rule in mapping_rules:
                if not isinstance(rule, dict):
                    continue
                logic_rule_id = str(rule.get('logic_rule_id') or '').strip()
                if not logic_rule_id:
                    continue

                targets.append({
                    'logic_rule_id': logic_rule_id,
                    'logic_rule_name': str(rule.get('logic_rule_name') or '').strip(),
                    'category': category,
                    'report_code': report_code,
                    'report_name': report_config.get('name') or report_code,
                    'target_column': str(rule.get('target_column') or '').strip(),
                    'target_name': str(rule.get('target_name') or '').strip(),
                })

    targets.sort(key=lambda item: (item['logic_rule_id'], item['category'], item['report_code'], item['target_column']))
    return targets


def _normalize_logic_rule_signature(source):
    if not isinstance(source, dict):
        return {}

    source_type = str(source.get('source_type') or source.get('type') or '').strip()
    source_field = source.get('source_field') if isinstance(source.get('source_field'), dict) else {}
    source_value = source.get('source_value')
    if source_value in (None, ''):
        source_value = source_field.get('field_index')

    def _normalized_conditions(conditions, include_result=True):
        normalized = []
        for item in conditions or []:
            if not isinstance(item, dict):
                continue
            row = {
                'field_index': str(item.get('field_index') if item.get('field_index') not in (None, '') else '').strip(),
                'match': str(item.get('match') or '').strip(),
                'operator': 'regex' if item.get('regex') else str(item.get('operator') or 'contains').strip(),
            }
            if include_result:
                row['result'] = str(item.get('result') or '').strip()
            normalized.append(row)
        return normalized

    def _normalized_groups(groups):
        normalized = []
        for group in groups or []:
            if not isinstance(group, dict):
                continue
            normalized.append({
                'logic': str(group.get('logic') or 'AND').strip(),
                'result': str(group.get('result') or '').strip(),
                'conditions': _normalized_conditions(group.get('conditions') or [], include_result=False)
            })
        return normalized

    return {
        'source_type': source_type,
        'source_value': '' if source_value in (None, '') else str(source_value).strip(),
        'default': str(source.get('default') or '').strip(),
        'logic': str(source.get('logic') or '').strip(),
        'group_logic': str(source.get('group_logic') or '').strip(),
        'result': str(source.get('result') or '').strip(),
        'conditions': _normalized_conditions(source.get('conditions') or []),
        'groups': _normalized_groups(source.get('groups') or []),
    }


def _logic_rule_signature_matches(logic_rule, mapping_rule):
    left = _normalize_logic_rule_signature(logic_rule)
    right = _normalize_logic_rule_signature(mapping_rule)
    if not left or not right:
        return False
    if left.get('source_type') != right.get('source_type'):
        return False
    return left == right


def _collect_logic_rule_usage_details(rule_id, mapping_config=None, logic_rules_payload=None):
    config = mapping_config if isinstance(mapping_config, dict) else _load_mapping_config_payload()
    logic_rules = logic_rules_payload if isinstance(logic_rules_payload, dict) else _load_logic_rules_payload()
    target_rule = copy.deepcopy((logic_rules.get('rules') or {}).get(str(rule_id or '').strip()) or {})
    if not target_rule:
        return {'explicit_usages': [], 'inferred_usages': []}

    reports = config.get('reports') or {}
    explicit_usages = []
    inferred_usages = []

    for category, category_reports in reports.items():
        if not isinstance(category_reports, dict):
            continue

        for report_code, report_config in category_reports.items():
            mapping_rules = report_config.get('mapping_rules') or []
            if not isinstance(mapping_rules, list):
                continue

            for rule in mapping_rules:
                if not isinstance(rule, dict):
                    continue

                usage_item = {
                    'category': category,
                    'report_code': report_code,
                    'report_name': report_config.get('name') or report_code,
                    'target_column': str(rule.get('target_column') or '').strip(),
                    'target_name': str(rule.get('target_name') or '').strip(),
                    'source_type': str(rule.get('source_type') or '').strip(),
                }

                if str(rule.get('logic_rule_id') or '').strip() == str(rule_id or '').strip():
                    explicit_usages.append(usage_item)
                    continue

                if _logic_rule_signature_matches(target_rule, rule):
                    inferred_usages.append(usage_item)

    explicit_usages.sort(key=lambda item: (item['category'], item['report_code'], item['target_column']))
    inferred_usages.sort(key=lambda item: (item['category'], item['report_code'], item['target_column']))
    return {
        'explicit_usages': explicit_usages,
        'inferred_usages': inferred_usages
    }


def _sync_logic_rule_usage_metadata(rule_id, mapping_config=None, logic_rules_payload=None):
    logic_rules = logic_rules_payload if isinstance(logic_rules_payload, dict) else _load_logic_rules_payload()
    rules_map = logic_rules.get('rules') or {}
    if str(rule_id or '').strip() not in rules_map:
        return logic_rules

    usages = _collect_logic_rule_usages(rule_id, mapping_config)
    target_rule = rules_map[str(rule_id).strip()]
    target_rule['usage_count'] = len(usages)
    target_rule['used_in'] = copy.deepcopy(usages)
    return logic_rules


def _sync_all_logic_rule_usage_metadata(mapping_config=None, logic_rules_payload=None):
    logic_rules = logic_rules_payload if isinstance(logic_rules_payload, dict) else _load_logic_rules_payload()
    rules_map = logic_rules.get('rules') or {}
    config = mapping_config if isinstance(mapping_config, dict) else _load_mapping_config_payload()

    for rule_id, rule in rules_map.items():
        usages = _collect_logic_rule_usages(rule_id, config)
        rule['usage_count'] = len(usages)
        rule['used_in'] = copy.deepcopy(usages)

    return logic_rules


def _apply_logic_rule_to_mapping_rule(logic_rule, existing_rule):
    updated_rule = copy.deepcopy(existing_rule) if isinstance(existing_rule, dict) else {}
    rule_type = str(logic_rule.get('type') or '').strip()

    updated_rule['source_type'] = rule_type
    updated_rule['logic_rule_id'] = str(logic_rule.get('id') or '').strip()
    updated_rule['logic_rule_name'] = str(logic_rule.get('name') or '').strip()
    updated_rule['logic_rule_applied_at'] = datetime.now().isoformat(timespec='seconds')

    source_field = logic_rule.get('source_field') if isinstance(logic_rule.get('source_field'), dict) else {}
    if source_field:
        updated_rule['source_field'] = copy.deepcopy(source_field)
        field_index = source_field.get('field_index')
        if field_index not in (None, ''):
            updated_rule['source_value'] = str(field_index)

    field_mappings = [
        'conditions',
        'groups',
        'logic',
        'group_logic',
        'result',
        'default',
        'match_mode',
        'lookup_key',
        'lookup_field',
        'use_mapping',
        'mapping_config',
        'fallback_to_raw_lookup',
        'transform',
        'name_column_index',
        'reference_column',
        'expansion_rules',
        'notes'
    ]

    for key in field_mappings:
        if key in logic_rule:
            updated_rule[key] = copy.deepcopy(logic_rule.get(key))

    if rule_type == 'conditional':
        updated_rule.pop('groups', None)
        updated_rule.pop('group_logic', None)
        updated_rule.pop('result', None)
    elif rule_type == 'conditional_groups':
        updated_rule.pop('conditions', None)
        updated_rule.pop('logic', None)
        updated_rule.pop('result', None)
    elif rule_type == 'multi_conditional':
        updated_rule.pop('groups', None)
        updated_rule.pop('group_logic', None)

    return updated_rule


def _build_registered_public_config_entry(config_key, config_data=None, mapping_config=None, registry=None):
    registry_map = registry if isinstance(registry, dict) else _build_public_config_registry()
    registry_entry = copy.deepcopy(registry_map.get(config_key, {}))
    raw_entry = copy.deepcopy(config_data) if isinstance(config_data, dict) else {}
    usages = _collect_public_config_usages(config_key, mapping_config)
    affects_export = registry_entry.get('affects_export', raw_entry.get('affects_export', False))
    contract_role = registry_entry.get('contract_role') or raw_entry.get('contract_role')
    if not contract_role:
        contract_role = 'runtime' if affects_export else 'reserve'
    effect_summary = registry_entry.get('effect_summary') or raw_entry.get('effect_summary')
    if not effect_summary:
        effect_summary = (
            '保存后会直接影响已接线报表导出'
            if contract_role == 'runtime'
            else '当前仅为预留或储备配置，不直接影响导出'
        )

    entry = copy.deepcopy(raw_entry)
    entry['key'] = config_key
    entry['title'] = registry_entry.get('title') or entry.get('title') or '公共配置'
    entry['name'] = registry_entry.get('name') or entry.get('name') or config_key
    entry['description'] = registry_entry.get('description') or entry.get('description') or ''
    entry['contract_role'] = contract_role
    entry['affects_export'] = bool(affects_export)
    entry['effect_summary'] = effect_summary
    entry['config_type'] = registry_entry.get('config_type') or entry.get('config_type') or 'generic'
    entry['category'] = registry_entry.get('category') or entry.get('category') or 'general'
    entry['editor_mode'] = registry_entry.get('editor_mode') or entry.get('editor_mode') or 'generic'
    entry['supports_detail'] = bool(registry_entry.get('supports_detail', entry.get('supports_detail', False)))
    entry['supports_crud'] = bool(registry_entry.get('supports_crud', entry.get('supports_crud', False)))
    entry['supports_import_export'] = bool(
        registry_entry.get('supports_import_export', entry.get('supports_import_export', False))
    )
    entry['detail_api'] = registry_entry.get('detail_api', entry.get('detail_api'))
    entry['save_api'] = registry_entry.get('save_api', entry.get('save_api'))
    entry['template_api'] = registry_entry.get('template_api', entry.get('template_api'))
    entry['import_api'] = registry_entry.get('import_api', entry.get('import_api'))
    entry['export_api'] = registry_entry.get('export_api', entry.get('export_api'))
    entry['storage_type'] = registry_entry.get('storage_type', entry.get('storage_type'))
    entry['storage_file'] = registry_entry.get('storage_file', entry.get('storage_file'))
    entry['storage'] = copy.deepcopy(registry_entry.get('storage', entry.get('storage', {})))
    entry['fields'] = copy.deepcopy(registry_entry.get('fields', entry.get('fields', [])))
    entry['primary_field'] = registry_entry.get('primary_field', entry.get('primary_field'))
    entry['display_fields'] = copy.deepcopy(
        registry_entry.get('display_fields', entry.get('display_fields', []))
    )
    entry['capabilities'] = copy.deepcopy(registry_entry.get('capabilities', entry.get('capabilities', {})))
    entry['apis'] = copy.deepcopy(registry_entry.get('apis', entry.get('apis', {})))
    entry['schema'] = copy.deepcopy(registry_entry.get('schema', entry.get('schema', {})))
    entry['template_filename'] = registry_entry.get('template_filename', entry.get('template_filename', ''))
    entry['export_filename'] = registry_entry.get('export_filename', entry.get('export_filename', ''))
    entry['version'] = entry.get('version') or '1.0'
    entry['last_updated'] = entry.get('last_updated') or datetime.now().strftime('%Y-%m-%d')
    entry['total_count'] = int(entry.get('total_count') or 0)
    entry['used_by'] = usages
    entry['usage_count'] = len(usages)
    return entry


def _build_public_config_entries(mapping_config=None):
    config = copy.deepcopy(mapping_config) if isinstance(mapping_config, dict) else _load_mapping_config_payload()
    public_config = copy.deepcopy(config.get('public_config') or {})
    entries = {}
    registry = _build_public_config_registry()

    for config_key in registry.keys():
        raw_entry = public_config.get(config_key)
        registry_entry = registry.get(config_key, {})
        if registry_entry.get('config_type') == 'mapping_table':
            raw_entry = _build_public_mapping_registry_entry(config_key, raw_entry)
        entries[config_key] = _build_registered_public_config_entry(
            config_key,
            raw_entry,
            config,
            registry,
        )

    for config_key, raw_entry in public_config.items():
        if config_key in entries:
            continue
        entries[config_key] = _build_registered_public_config_entry(
            config_key,
            raw_entry,
            config,
            registry,
        )

    return entries


def _load_business_system_name_mapping_pairs():
    """优先读取专用映射文件，缺失时回退到 mapping_config.json。"""
    if os.path.exists(BUSINESS_SYSTEM_NAME_MAPPING_FILE):
        try:
            with open(BUSINESS_SYSTEM_NAME_MAPPING_FILE, 'r', encoding='utf-8-sig') as f:
                business_mapping = json.load(f)
            return _normalize_business_system_name_mapping_pairs(
                business_mapping.get('mappings', business_mapping)
            )
        except Exception:
            pass

    try:
        with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8-sig') as f:
            config = json.load(f)
        public_config = config.get('public_config', {})
        business_mapping = public_config.get(PUBLIC_CONFIG_BUSINESS_SYSTEM_NAME_MAPPING_KEY, {})
        return _normalize_business_system_name_mapping_pairs(
            business_mapping.get('mappings', business_mapping)
        )
    except Exception:
        return {}

def _write_business_system_name_mapping_file(file_path, payload):
    os.makedirs(os.path.dirname(file_path), exist_ok=True)
    with open(file_path, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)


def _save_business_system_name_mapping_file(mappings):
    """将映射字典同步写入专用映射文件。"""
    payload = {
        'name': '公共配置',
        'description': '用于将合并结果表中的业务系统名称转换为数据概览表中的标准名称',
        'version': '1.0',
        'last_updated': datetime.now().strftime('%Y-%m-%d'),
        'total_count': 0,
        'mappings': []
    }

    if os.path.exists(BUSINESS_SYSTEM_NAME_MAPPING_FILE):
        try:
            with open(BUSINESS_SYSTEM_NAME_MAPPING_FILE, 'r', encoding='utf-8-sig') as f:
                existing = json.load(f)
            if isinstance(existing, dict):
                for key in ('description', 'version'):
                    if existing.get(key):
                        payload[key] = existing[key]
        except Exception:
            pass

    payload['mappings'] = [
        {
            'id': idx,
            'source': source,
            'target': target,
            'enabled': True
        }
        for idx, (source, target) in enumerate(mappings.items(), 1)
    ]
    payload['total_count'] = len(payload['mappings'])

    _write_business_system_name_mapping_file(BUSINESS_SYSTEM_NAME_MAPPING_FILE, payload)


# 已删除 get_data_files_db_connection() - 不再需要数据文件元信息数据库

# ==================== 首页和基础路由 ====================

@common_bp.route('/')
def index():
    """首页"""
    return send_file(os.path.join(Config.BASE_DIR, 'web', 'index.html'))


# ==================== 列配置路由 ====================

@common_bp.route('/api/columns', methods=['GET'])
def get_columns_api():
    """获取列定义"""
    try:
        with open(PathConfig.ASSETS_COLUMNS_JSON, 'r', encoding='utf-8') as f:
            data = json.load(f)
            if isinstance(data, dict) and isinstance(data.get('全部'), dict):
                data = copy.deepcopy(data)
                data.pop('columns', None)
            return jsonify(data)
    except:
        return jsonify({'columns': []})


@common_bp.route('/api/columns', methods=['POST'])
def update_columns():
    """更新列配置"""
    data = request.json
    try:
        with open(PathConfig.ASSETS_COLUMNS_JSON, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ==================== 统计信息路由 ====================

@common_bp.route('/api/stats', methods=['GET'])
def get_stats():
    """获取统计信息"""
    try:
        from core.utils import ensure_database_and_table

        db_path = DatabaseConfig.ASSETS_DB

        # 确保数据库和表存在
        ensure_database_and_table(
            db_path,
            'assets',
            '''CREATE TABLE IF NOT EXISTS assets (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                "业务系统" TEXT,
                "所属系统类型（文字）" TEXT
            )'''
        )

        conn = get_db_connection()
        cursor = conn.cursor()

        # 总资产数
        cursor.execute('SELECT COUNT(*) as total FROM assets')
        total_assets = cursor.fetchone()['total']

        # 按系统类型统计
        cursor.execute('''
            SELECT "所属系统类型（文字）" as type, COUNT(*) as count
            FROM assets
            GROUP BY "所属系统类型（文字）"
        ''')
        by_type = [dict(row) for row in cursor.fetchall()]

        # 按业务系统统计
        cursor.execute('''
            SELECT "业务系统" as system, COUNT(*) as count
            FROM assets
            GROUP BY "业务系统"
            ORDER BY count DESC
            LIMIT 10
        ''')
        top_systems = [dict(row) for row in cursor.fetchall()]

        # 业务系统数量（去重）
        cursor.execute('''
            SELECT COUNT(DISTINCT "业务系统") as system_count
            FROM assets
            WHERE "业务系统" IS NOT NULL AND "业务系统" != ''
        ''')
        system_count = cursor.fetchone()['system_count']

        conn.close()

        return jsonify({
            'total': total_assets,
            'systemCount': system_count,
            'total_assets': total_assets,  # 兼容旧版本
            'by_type': by_type,
            'top_systems': top_systems
        })
    except Exception as e:
        return jsonify({
            'error': f'获取统计信息失败: {str(e)}',
            'total': 0,
            'systemCount': 0,
            'total_assets': 0,
            'by_type': [],
            'top_systems': []
        }), 500


# ==================== 导入预览路由 ====================

@common_bp.route('/api/import/preview', methods=['POST'])
def preview_import():
    """预览导入的Excel文件"""
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': '未选择文件'}), 400

    file = request.files['file']

    try:
        from openpyxl import load_workbook

        workbook = load_workbook(file, read_only=True, data_only=True)
        worksheet = workbook.active

        rows_iter = worksheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            workbook.close()
            return jsonify({'success': False, 'error': 'Excel 文件为空'}), 400

        columns = [
            str(cell) if cell is not None else f'列{i + 1}'
            for i, cell in enumerate(header_row)
        ]
        expected_headers = _get_workspace_template_headers_safe('device')
        missing_columns, extra_columns = _compare_import_headers(columns, expected_headers)

        preview_rows = []
        total_rows = 0

        for row in rows_iter:
            total_rows += 1
            if len(preview_rows) < 10:
                row_values = list(row)
                normalized_row = {
                    columns[i]: '' if i >= len(row_values) or row_values[i] is None else str(row_values[i])
                    for i in range(len(columns))
                }
                preview_rows.append(normalized_row)

        workbook.close()

        preview_data = {
            'success': True,
            'columns': columns,
            'rows': preview_rows,
            'total_columns': len(columns),
            'rowCount': total_rows,
            'missing_columns': missing_columns,
            'extra_columns': extra_columns,
            'warning_message': _build_import_schema_warning('device', missing_columns, extra_columns)
            if (missing_columns or extra_columns) else ''
        }

        return jsonify(preview_data)
    except Exception as e:
        return jsonify({'success': False, 'error': f'读取文件失败: {str(e)}'}), 400


@common_bp.route('/api/import/confirm', methods=['POST'])
def confirm_import():
    """确认导入Excel文件到assets表（高效批量导入模式）"""
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': '未选择文件'}), 400

    file = request.files['file']
    file_name = file.filename

    try:
        from core.utils import ensure_database_and_table
        import threading
        import time
        from datetime import datetime

        # 读取Excel文件
        df = pd.read_excel(file)
        columns = df.columns.tolist()
        total_rows = len(df)
        expected_headers = _get_workspace_template_headers_safe('device')
        missing_columns, extra_columns = _compare_import_headers(columns, expected_headers)
        warning_message = _build_import_schema_warning('device', missing_columns, extra_columns) if (missing_columns or extra_columns) else ''
        force_import = _is_force_import_requested()
        if missing_columns and not force_import:
            return jsonify({
                'success': False,
                'requires_confirmation': True,
                'warning_type': 'missing_columns',
                'missing_columns': missing_columns,
                'extra_columns': extra_columns,
                'warning_message': warning_message
            })

        # 构建CREATE TABLE语句
        column_defs = ', '.join([f'"{col}" TEXT' for col in columns])
        create_table_sql = (
            'CREATE TABLE IF NOT EXISTS assets ('
            'id INTEGER PRIMARY KEY AUTOINCREMENT'
            f'{", " if column_defs else ""}{column_defs})'
        )

        # 确保数据库和表存在
        ensure_database_and_table(
            DatabaseConfig.ASSETS_DB,
            'assets',
            create_table_sql
        )

        clear_import_progress()
        _set_asset_import_progress('running', 0, total_rows, '开始导入...')

        def do_import():
            """后台线程执行导入"""
            start_time = time.time()

            try:
                conn = get_db_connection()
                cursor = conn.cursor()

                # 删除旧表并创建新表
                cursor.execute('DROP TABLE IF EXISTS assets')
                cursor.execute(create_table_sql)

                # 优化配置
                BATCH_SIZE = 50000  # 每批5万行

                # 准备列名和占位符
                col_names = ', '.join([f'"{col}"' for col in columns])
                placeholders = ', '.join(['?'] * len(columns))

                # 分批插入数据
                for i in range(0, total_rows, BATCH_SIZE):
                    batch_df = df.iloc[i:i+BATCH_SIZE]

                    # 批量构建数据
                    data_list = []
                    for _, row in batch_df.iterrows():
                        row_data = tuple(str(v) if pd.notna(v) else '' for v in row.values)
                        data_list.append(row_data)

                    # 使用 executemany 批量插入
                    cursor.executemany(
                        f'INSERT INTO assets ({col_names}) VALUES ({placeholders})',
                        data_list
                    )
                    conn.commit()

                    # 更新进度
                    current_count = min(i + BATCH_SIZE, total_rows)
                    percent = round((current_count / total_rows) * 100, 1) if total_rows else 100
                    _set_asset_import_progress(
                        'running',
                        current_count,
                        total_rows,
                        f'正在导入... {percent}% ({current_count:,} / {total_rows:,})'
                    )

                conn.close()

                # 计算耗时
                elapsed_time = round(time.time() - start_time, 1)
                rows_per_sec = round(total_rows / elapsed_time, 0) if elapsed_time > 0 else 0

                # 更新列配置，兼容当前按“全部”分组的结构
                column_data = get_columns()
                grouped_config = column_data.get('全部') if isinstance(column_data.get('全部'), dict) else None
                target_config = grouped_config if grouped_config is not None else column_data

                new_column_config = []
                existing_cols = {
                    col['name']: col for col in target_config.get('columns', [])
                    if isinstance(col, dict) and col.get('name')
                }

                for col_name in columns:
                    if col_name in existing_cols:
                        new_column_config.append(existing_cols[col_name])
                    else:
                        new_column_config.append({
                            'name': col_name,
                            'visible': True,
                            'width': 120
                        })

                target_config['columns'] = new_column_config
                target_config['sourceFile'] = file_name
                if grouped_config is not None:
                    column_data['全部'] = target_config
                    column_data.pop('columns', None)
                    column_data.pop('sourceFile', None)
                else:
                    column_data = target_config

                with open(PathConfig.ASSETS_COLUMNS_JSON, 'w', encoding='utf-8') as f:
                    json.dump(column_data, f, ensure_ascii=False, indent=2)

                # 完成状态
                _set_asset_import_progress(
                    'completed',
                    total_rows,
                    total_rows,
                    f'导入完成！共导入 {total_rows:,} 条记录，耗时 {elapsed_time} 秒（{rows_per_sec:,} 条/秒）'
                )

                # 记录操作日志
                try:
                    log_operation = get_log_operation()
                    log_operation(
                        operation_type='data_import',
                        page_type='device',
                        operation_desc='导入数据概览',
                        file_name=file_name,
                        record_count=total_rows,
                        details=f'耗时 {elapsed_time} 秒'
                    )
                except Exception as log_err:
                    print(f'[日志记录失败] {log_err}')

            except Exception as e:
                import traceback
                traceback.print_exc()
                if 'conn' in locals():
                    conn.close()
                _set_asset_import_progress(
                    'error',
                    import_progress.get('current', 0),
                    total_rows,
                    f'导入失败: {str(e)}',
                    str(e)
                )

        # 在后台线程中执行导入
        thread = threading.Thread(target=do_import)
        thread.daemon = True
        thread.start()

        return jsonify({
            'success': True,
            'message': f'开始导入 {total_rows:,} 条记录，请查看进度条',
            'missing_columns': missing_columns,
            'extra_columns': extra_columns,
            'warning_message': warning_message
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': f'导入失败: {str(e)}'}), 500


def get_columns():
    """获取列配置"""
    try:
        with open(PathConfig.ASSETS_COLUMNS_JSON, 'r', encoding='utf-8') as f:
            return json.load(f)
    except:
        return {'columns': []}


# ==================== 通用导出路由 ====================

@common_bp.route('/api/export', methods=['GET'])
def export_data():
    """导出assets数据到Excel"""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('SELECT * FROM assets')
        rows = cursor.fetchall()

        # 获取列名
        columns = [desc[0] for desc in cursor.description]
        conn.close()

        # 转换为DataFrame
        data = [dict(row) for row in rows]
        df = pd.DataFrame(data)

        # 导出
        from datetime import datetime
        filename = f'assets_export_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        df.to_excel(filename, index=False)

        # 记录操作日志
        try:
            log_operation = get_log_operation()
            log_operation(
                operation_type='data_export',
                page_type='device',
                operation_desc='导出数据概览',
                file_name=filename,
                record_count=len(data)
            )
        except Exception as log_err:
            print(f'[日志记录失败] {log_err}')

        return jsonify({'success': True, 'filename': filename})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ==================== 文件下载路由 ====================

@common_bp.route('/api/download-source')
def download_source():
    """下载源文件"""
    file_type = request.args.get('type', 'assets')
    filename = (request.args.get('filename', '') or request.args.get('file', '')).strip()

    if not filename:
        return jsonify({'error': '缺少文件名'}), 400

    try:
        file_path, safe_filename = find_source_file_path(filename, file_type)

        if not file_path:
            return jsonify({'error': '文件不存在'}), 404

        return send_file(file_path, as_attachment=True, download_name=safe_filename)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ==================== 存储文件管理路由 ====================

@common_bp.route('/api/stored-files', methods=['GET'])
def get_stored_files():
    """获取所有存储的文件列表"""
    try:
        upload_folder = PathConfig.UPLOAD_FOLDER
        if not os.path.exists(upload_folder):
            return jsonify({'files': []})

        files = []
        for filename in os.listdir(upload_folder):
            file_path = os.path.join(upload_folder, filename)
            if os.path.isfile(file_path):
                files.append({
                    'filename': filename,
                    'size': os.path.getsize(file_path),
                    'modified': os.path.getmtime(file_path)
                })

        # 按修改时间降序排序
        files.sort(key=lambda x: x['modified'], reverse=True)

        return jsonify({'files': files})
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@common_bp.route('/api/stored-files/download/<filename>')
def download_stored_file(filename):
    """下载存储的文件"""
    try:
        file_path, safe_filename = resolve_safe_file_path(PathConfig.UPLOAD_FOLDER, filename)

        if not os.path.exists(file_path):
            return jsonify({'error': '文件不存在'}), 404

        return send_file(file_path, as_attachment=True, download_name=safe_filename)
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500


@common_bp.route('/api/stored-files/<filename>', methods=['DELETE'])
def delete_stored_file(filename):
    """删除存储的文件"""
    try:
        blocked_response = reject_cross_site_delete_request()
        if blocked_response:
            return blocked_response

        file_path, _ = resolve_safe_file_path(PathConfig.UPLOAD_FOLDER, filename)

        if not os.path.exists(file_path):
            return jsonify({'error': '文件不存在'}), 404

        os.remove(file_path)

        # 记录操作日志
        try:
            log_operation = get_log_operation()
            log_operation(
                operation_type='file_delete',
                page_type='file_manage',
                operation_desc='删除存储文件',
                file_name=filename,
                details='从uploaded_files目录删除'
            )
        except Exception as log_err:
            print(f'[日志记录失败] {log_err}')

        return jsonify({'success': True, 'message': '文件已删除'})
    except ValueError as e:
        return jsonify({'error': str(e)}), 400
    except Exception as e:
        return jsonify({'error': str(e)}), 500


# ==================== 工程文件管理路由 ====================

@common_bp.route('/api/project-files', methods=['GET'])
def get_project_files():
    """获取工程文件列表（模板、数据库、数据处理、合并结果、数据概览、填报数据）"""
    try:
        import glob
        from datetime import datetime
        import re

        # 新的数据结构：templates保持扁平，其他类型按月份分组
        files = {
            'templates': [],  # 扁平结构
            'databases': {},  # 按月份分组: {'2026-03': [files...]}
            'data_process': [],  # 扁平结构
            'merge_results': {},
            'assets': {},
            'exports': {}
        }

        # 辅助函数：从文件路径或修改时间提取月份标签
        def extract_month_tag(file_path, modified_time):
            """尝试从路径或时间提取月份标签，格式：YYYY-MM"""
            # 先尝试从路径中提取
            path_parts = file_path.replace('\\', '/').split('/')
            for part in path_parts:
                # 匹配 2026-03 或 202603 格式
                if re.match(r'\d{4}-\d{2}', part):
                    return part
                if re.match(r'\d{6}', part):
                    return f"{part[:4]}-{part[4:6]}"

            # 如果路径中没有，使用修改时间的月份
            return datetime.fromtimestamp(modified_time).strftime('%Y-%m')

        # 辅助函数：按月份分组文件
        def group_by_month(file_list):
            """将文件列表按月份分组"""
            grouped = {}
            for file in file_list:
                month_tag = extract_month_tag(file['path'], file['modified'])
                if month_tag not in grouped:
                    grouped[month_tag] = []
                grouped[month_tag].append(file)
            return grouped

        # 1. 模板文件（Templates目录）- 保持扁平结构
        templates_dir = PathConfig.TEMPLATES_DIR
        if os.path.exists(templates_dir):
            for category in os.listdir(templates_dir):
                category_path = os.path.join(templates_dir, category)
                if os.path.isdir(category_path):
                    for filename in os.listdir(category_path):
                        if filename.endswith(('.xlsx', '.xls')) and not filename.startswith('~'):
                            file_path = os.path.join(category_path, filename)
                            files['templates'].append({
                                'filename': filename,
                                'category': category,
                                'path': file_path,
                                'size': os.path.getsize(file_path),
                                'modified': os.path.getmtime(file_path),
                                'type': 'template'
                            })

        # 2. 数据库文件（data目录中的.db文件）- 按月份分组
        data_dir = os.path.join(Config.BASE_DIR, 'data')
        databases_list = []
        if os.path.exists(data_dir):
            for filename in os.listdir(data_dir):
                if filename.endswith('.db'):
                    file_path = os.path.join(data_dir, filename)
                    databases_list.append({
                        'filename': filename,
                        'path': file_path,
                        'size': os.path.getsize(file_path),
                        'modified': os.path.getmtime(file_path),
                        'type': 'database'
                    })
        files['databases'] = group_by_month(databases_list)

        # 3. 合并结果文件 - 按月份分组
        merge_results_list = []

        # 3.1 添加exports目录中的合并结果导出文件
        exports_dir = PathConfig.EXPORT_FOLDER
        if os.path.exists(exports_dir):
            for filename in os.listdir(exports_dir):
                # 包含"合并结果"的文件
                if '合并结果' in filename and filename.endswith(('.xlsx', '.xls')) and not filename.startswith('~'):
                    file_path = os.path.join(exports_dir, filename)
                    merge_results_list.append({
                        'filename': filename,
                        'source': 'exports',
                        'path': file_path,
                        'size': os.path.getsize(file_path),
                        'modified': os.path.getmtime(file_path),
                        'type': 'merge_result',
                        'is_current': False,
                        'month_tag': ''
                    })

        # 3.2 添加SGExportFiles中的导出文件（排除业支上报文件）
        try:
            yezhi_dir = PathConfig.YEZHI_EXPORT_DIR
            if os.path.exists(yezhi_dir):
                for filename in os.listdir(yezhi_dir):
                    # 排除业支上报文件（i_10600_等前缀）
                    if filename.startswith('i_10600_'):
                        continue

                    if filename.endswith(('.xlsx', '.xls')) and not filename.startswith('~'):
                        file_path = os.path.join(yezhi_dir, filename)
                        # 避免重复（按完整路径去重，避免不同目录同名文件被误去重）
                        if not any(os.path.abspath(f.get('path', '')) == os.path.abspath(file_path) for f in merge_results_list):
                            merge_results_list.append({
                                'filename': filename,
                                'source': 'SGExportFiles',
                                'path': file_path,
                                'size': os.path.getsize(file_path),
                                'modified': os.path.getmtime(file_path),
                                'type': 'merge_result',
                                'is_current': False,
                                'month_tag': ''
                            })
        except Exception as e:
            print(f'[WARNING] 读取SGExportFiles文件失败: {e}')

        files['merge_results'] = group_by_month(merge_results_list)

        # 4. 数据概览文件 - 按月份分组
        assets_list = []

        # 4.1 添加exports目录中的数据概览导出文件
        if os.path.exists(exports_dir):
            for filename in os.listdir(exports_dir):
                # 包含"数据概览"或"资产"的文件
                if ('数据概览' in filename or '资产' in filename) and filename.endswith(('.xlsx', '.xls')) and not filename.startswith('~'):
                    file_path = os.path.join(exports_dir, filename)
                    # 避免与合并结果重复
                    if not any(f['filename'] == filename for f in merge_results_list):
                        assets_list.append({
                            'filename': filename,
                            'source': 'exports',
                            'path': file_path,
                            'size': os.path.getsize(file_path),
                            'modified': os.path.getmtime(file_path),
                            'type': 'asset',
                            'is_current': False,
                            'month_tag': ''
                        })

        # 4.2 uploaded_files中的asset文件
        if os.path.exists(PathConfig.UPLOAD_FOLDER):
            for filename in os.listdir(PathConfig.UPLOAD_FOLDER):
                if 'asset' in filename.lower() and filename.endswith(('.xlsx', '.xls')):
                    file_path = os.path.join(PathConfig.UPLOAD_FOLDER, filename)
                    # 避免重复（按完整路径去重，避免不同目录同名文件被误去重）
                    if not any(os.path.abspath(f.get('path', '')) == os.path.abspath(file_path) for f in assets_list):
                        assets_list.append({
                            'filename': filename,
                            'source': 'uploaded_files',
                            'path': file_path,
                            'size': os.path.getsize(file_path),
                            'modified': os.path.getmtime(file_path),
                            'type': 'asset',
                            'is_current': False,
                            'month_tag': ''
                        })

        files['assets'] = group_by_month(assets_list)

        # 5. 数据处理结果文件（temp_data_process）- 扁平结构
        data_process_list = []
        data_process_base_dir = os.path.join(Config.BASE_DIR, 'temp_data_process')
        data_process_modes = {
            'allinone': '一键处理',
            'split': '数据拆分',
            'csv': 'CSV转换',
            'dedup': '数据去重'
        }
        for mode_dir_name, mode_label in data_process_modes.items():
            mode_dir = os.path.join(data_process_base_dir, mode_dir_name)
            if not os.path.exists(mode_dir):
                continue

            for root, _, filenames in os.walk(mode_dir):
                for filename in filenames:
                    if filename.startswith('~') or filename.startswith('.'):
                        continue

                    file_path = os.path.join(root, filename)
                    if not os.path.isfile(file_path):
                        continue

                    relative_dir = os.path.relpath(root, mode_dir)
                    if relative_dir in ('.', ''):
                        file_note = mode_label
                    else:
                        normalized_relative_dir = relative_dir.replace('\\', '/')
                        file_note = f'{mode_label} / {normalized_relative_dir}'

                    data_process_list.append({
                        'filename': filename,
                        'source': 'temp_data_process',
                        'path': file_path,
                        'size': os.path.getsize(file_path),
                        'modified': os.path.getmtime(file_path),
                        'type': 'data_process',
                        'note': file_note,
                        'is_current': False
                    })

        files['data_process'] = data_process_list

        # 6. 填报数据导出文件 - 按月份分组
        exports_list = []
        # SGExportFiles通用导出目录
        exports_dir = PathConfig.EXPORT_FOLDER
        if os.path.exists(exports_dir):
            for filename in os.listdir(exports_dir):
                if filename.endswith(('.xlsx', '.xls')) and not filename.startswith('~'):
                    file_path = os.path.join(exports_dir, filename)
                    exports_list.append({
                        'filename': filename,
                        'path': file_path,
                        'size': os.path.getsize(file_path),
                        'modified': os.path.getmtime(file_path),
                        'type': 'export',
                        'month_tag': ''  # 会在group_by_month中从路径提取
                    })

        # YeZhi导出目录
        yezhi_dir = PathConfig.YEZHI_EXPORT_DIR
        if os.path.exists(yezhi_dir):
            for filename in os.listdir(yezhi_dir):
                if filename.endswith(('.xlsx', '.xls')) and not filename.startswith('~'):
                    file_path = os.path.join(yezhi_dir, filename)
                    # 避免重复（按完整路径去重，避免不同目录同名文件被误去重）
                    if not any(os.path.abspath(f.get('path', '')) == os.path.abspath(file_path) for f in exports_list):
                        exports_list.append({
                            'filename': filename,
                            'source': 'YeZhi',
                            'path': file_path,
                            'size': os.path.getsize(file_path),
                            'modified': os.path.getmtime(file_path),
                            'type': 'export',
                            'month_tag': ''  # 会在group_by_month中从路径提取
                        })

        # SMC导出目录
        smc_dir = PathConfig.SMC_EXPORT_DIR
        if os.path.exists(smc_dir):
            for filename in os.listdir(smc_dir):
                if filename.endswith(('.xlsx', '.xls')) and not filename.startswith('~'):
                    file_path = os.path.join(smc_dir, filename)
                    # 避免重复（按完整路径去重，避免不同目录同名文件被误去重）
                    if not any(os.path.abspath(f.get('path', '')) == os.path.abspath(file_path) for f in exports_list):
                        exports_list.append({
                            'filename': filename,
                            'source': 'SMC',
                            'path': file_path,
                            'size': os.path.getsize(file_path),
                            'modified': os.path.getmtime(file_path),
                            'type': 'export',
                            'month_tag': ''  # 会在group_by_month中从路径提取
                        })

        # XinAn导出目录
        xinan_dir = PathConfig.XINAN_EXPORT_DIR
        if os.path.exists(xinan_dir):
            for filename in os.listdir(xinan_dir):
                if filename.endswith(('.xlsx', '.xls')) and not filename.startswith('~'):
                    file_path = os.path.join(xinan_dir, filename)
                    # 避免重复（按完整路径去重，避免不同目录同名文件被误去重）
                    if not any(os.path.abspath(f.get('path', '')) == os.path.abspath(file_path) for f in exports_list):
                        exports_list.append({
                            'filename': filename,
                            'source': 'XinAn',
                            'path': file_path,
                            'size': os.path.getsize(file_path),
                            'modified': os.path.getmtime(file_path),
                            'type': 'export',
                            'month_tag': ''  # 会在group_by_month中从路径提取
                        })
        files['exports'] = group_by_month(exports_list)

        # 格式化文件大小
        def format_size(size_bytes):
            for unit in ['B', 'KB', 'MB', 'GB']:
                if size_bytes < 1024.0:
                    return f"{size_bytes:.1f}{unit}"
                size_bytes /= 1024.0
            return f"{size_bytes:.1f}TB"

        # 格式化时间和大小，并排序
        # 1. templates是列表，直接排序
        files['templates'].sort(key=lambda x: x['modified'], reverse=True)
        for file in files['templates']:
            file['modified_formatted'] = datetime.fromtimestamp(file['modified']).strftime('%Y-%m-%d %H:%M')
            file['size_formatted'] = format_size(file['size'])

        files['data_process'].sort(key=lambda x: x['modified'], reverse=True)
        for file in files['data_process']:
            file['modified_formatted'] = datetime.fromtimestamp(file['modified']).strftime('%Y-%m-%d %H:%M')
            file['size_formatted'] = format_size(file['size'])

        # 3. 其他类型是字典（按月份分组），需要分别排序和格式化
        for category in ['databases', 'merge_results', 'assets', 'exports']:
            for month_tag in files[category]:
                files[category][month_tag].sort(key=lambda x: x['modified'], reverse=True)
                for file in files[category][month_tag]:
                    file['modified_formatted'] = datetime.fromtimestamp(file['modified']).strftime('%Y-%m-%d %H:%M')
                    file['size_formatted'] = format_size(file['size'])

        # 统计信息
        def count_grouped_files(grouped_data):
            """计算分组数据的文件总数"""
            if isinstance(grouped_data, list):
                return len(grouped_data)
            return sum(len(files) for files in grouped_data.values())

        stats = {
            'total_files': count_grouped_files(files['templates']) +
                          count_grouped_files(files['data_process']) +
                          sum(count_grouped_files(files[key]) for key in ['databases', 'merge_results', 'assets', 'exports']),
            'templates_count': len(files['templates']),
            'data_process_count': len(files['data_process']),
            'databases_count': count_grouped_files(files['databases']),
            'merge_results_count': count_grouped_files(files['merge_results']),
            'assets_count': count_grouped_files(files['assets']),
            'exports_count': count_grouped_files(files['exports'])
        }

        return jsonify({
            'success': True,
            'files': files,
            'stats': stats
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


@common_bp.route('/api/project-files/delete', methods=['POST'])
def delete_project_file():
    """删除工程文件"""
    try:
        blocked_response = reject_cross_site_delete_request()
        if blocked_response:
            return blocked_response

        data = request.json
        file_type = data.get('type')
        filename = data.get('filename')
        file_path = data.get('path')

        if not file_path:
            return jsonify({'error': '缺少文件路径'}), 400

        # 安全检查：确保文件路径在允许的目录内
        import os
        base_dir = os.path.abspath(Config.BASE_DIR)
        abs_file_path = os.path.abspath(file_path)

        # 检查文件路径是否在项目目录内
        if not is_path_within(base_dir, abs_file_path):
            return jsonify({'error': '不允许删除项目目录外的文件'}), 403

        # 保护关键数据库文件
        protected_files = ['assets.db', 'merge_results.db', 'data_files.db']
        if filename in protected_files:
            return jsonify({'error': f'不能删除关键数据库文件: {filename}'}), 403

        # 检查文件是否存在
        if not os.path.exists(abs_file_path):
            return jsonify({'error': '文件不存在'}), 404

        # 删除文件
        try:
            os.remove(abs_file_path)
        except Exception as e:
            return jsonify({'error': f'删除失败: {str(e)}'}), 500

        # 记录操作日志
        try:
            log_operation = get_log_operation()
            log_operation(
                operation_type='file_delete',
                page_type='file_manage',
                operation_desc='删除工程文件',
                file_name=filename,
                details=f'类型: {file_type}'
            )
        except Exception as log_err:
            print(f'[日志记录失败] {log_err}')

        return jsonify({
            'success': True,
            'message': f'文件 "{filename}" 已删除'
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


@common_bp.route('/api/project-files/download/<file_type>/<path:filename>', methods=['GET'])
def download_project_file(file_type, filename):
    """下载工程文件"""
    try:
        import os
        from flask import send_file

        base_dirs = get_project_file_base_dirs(file_type)
        if not base_dirs:
            return jsonify({'error': f'不支持的文件类型: {file_type}'}), 400

        # 如果 filename 包含路径分隔符，视为完整路径，允许新前端按真实路径下载。
        if '\\' in filename or '/' in filename:
            file_path = os.path.abspath(filename)
            base_dir = next(
                (candidate for candidate in base_dirs if is_path_within(candidate, file_path)),
                None
            )
            if not base_dir:
                return jsonify({'error': '不允许下载项目目录外的文件'}), 403
        else:
            file_path = None
            base_dir = None
            for search_dir in base_dirs:
                test_path = os.path.join(search_dir, filename)
                if os.path.exists(test_path):
                    file_path = test_path
                    base_dir = search_dir
                    break

            if not file_path:
                file_path = os.path.join(base_dirs[0], filename)
                base_dir = base_dirs[0]

        # 安全检查：确保文件路径在允许的目录内
        abs_file_path = os.path.abspath(file_path)
        abs_base_dir = os.path.abspath(base_dir)

        if not is_path_within(abs_base_dir, abs_file_path):
            return jsonify({'error': '不允许下载项目目录外的文件'}), 403

        # 检查文件是否存在
        if not os.path.exists(abs_file_path):
            return jsonify({'error': f'文件不存在: {filename}'}), 404

        # 检查是否为目录
        if os.path.isdir(abs_file_path):
            return jsonify({'error': '不能下载目录'}), 400

        # 从完整路径中提取实际文件名用于下载
        actual_filename = os.path.basename(abs_file_path)

        # 发送文件
        return send_file(
            abs_file_path,
            as_attachment=True,
            download_name=actual_filename,
            mimetype='application/octet-stream'
        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


# ==================== 模板文件上传管理 ====================

@common_bp.route('/api/templates/upload', methods=['POST'])
def upload_template():
    """上传模板文件到Templates目录"""
    try:
        if 'file' not in request.files:
            return jsonify({'error': '未选择文件'}), 400

        file = request.files['file']
        category = request.form.get('category', '业支上报')

        # 验证分类
        valid_categories = ['业支上报', 'SMC上报', '信安上报', '共用']
        if category not in valid_categories:
            return jsonify({'error': f'无效的分类，有效分类: {", ".join(valid_categories)}'}), 400

        # 验证文件类型
        if not file.filename.endswith(('.xlsx', '.xls')):
            return jsonify({'error': '只支持Excel文件（.xlsx, .xls）'}), 400

        # 保存文件到对应分类目录
        category_dir = os.path.join(PathConfig.TEMPLATES_DIR, category)
        if not os.path.exists(category_dir):
            os.makedirs(category_dir)

        file_path = os.path.join(category_dir, file.filename)
        file.save(file_path)

        # 记录操作日志
        try:
            log_operation = get_log_operation()
            log_operation(
                operation_type='file_upload',
                page_type='file_manage',
                operation_desc='上传模板文件',
                file_name=file.filename,
                details=f'分类: {category}'
            )
        except Exception as log_err:
            print(f'[日志记录失败] {log_err}')

        return jsonify({
            'success': True,
            'message': f'模板 "{file.filename}" 上传成功',
            'category': category,
            'filename': file.filename
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


@common_bp.route('/api/templates/categories', methods=['GET'])
def get_template_categories():
    """获取模板分类列表"""
    try:
        categories = []
        for category_name in ['业支上报', 'SMC上报', '信安上报', '共用']:
            category_dir = os.path.join(PathConfig.TEMPLATES_DIR, category_name)
            file_count = 0
            if os.path.exists(category_dir):
                file_count = len([f for f in os.listdir(category_dir) if f.endswith(('.xlsx', '.xls')) and not f.startswith('~')])

            categories.append({
                'name': category_name,
                'path': category_dir,
                'file_count': file_count
            })

        return jsonify({
            'success': True,
            'categories': categories
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


# ==================== 数据删除路由 ====================

@common_bp.route('/api/data/delete', methods=['POST'])
def delete_data():
    """删除数据表格及相关文件"""
    try:
        blocked_response = reject_cross_site_delete_request()
        if blocked_response:
            return blocked_response

        data = request.get_json()
        data_type = data.get('type', 'device')  # 'device' 或 'merge'

        print(f'[DELETE DATA] 开始删除 {data_type} 类型的数据')

        # 根据类型删除不同的数据
        if data_type == 'device':
            result = delete_device_data()
        elif data_type == 'merge':
            result = delete_merge_data()
        else:
            return jsonify({'error': f'无效的数据类型: {data_type}'}), 400

        if result['success']:
            print(f'[DELETE DATA] 删除成功')

            # 记录操作日志
            try:
                log_operation = get_log_operation()
                log_operation(
                    operation_type='data_delete',
                    page_type=data_type,
                    operation_desc=f'删除{data_type}数据',
                    details=result['message']
                )
            except Exception as log_err:
                print(f'[日志记录失败] {log_err}')

            return jsonify({
                'success': True,
                'message': result['message']
            })
        else:
            return jsonify({'error': result['error']}), 500

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'error': str(e)}), 500


def delete_device_data():
    """删除设备管理数据"""
    try:
        deleted_files = []

        # 1. 删除 assets.db 数据库文件
        db_path = DatabaseConfig.ASSETS_DB
        if os.path.exists(db_path):
            os.remove(db_path)
            deleted_files.append(db_path)
            print(f'[DELETE] 已删除数据库文件: {db_path}')

        # 2. 删除uploaded_files目录下的assets相关文件
        upload_dir = PathConfig.UPLOAD_FOLDER
        if os.path.exists(upload_dir):
            for filename in os.listdir(upload_dir):
                if 'asset' in filename.lower() and filename.endswith(('.xlsx', '.xls', '.db')):
                    file_path = os.path.join(upload_dir, filename)
                    os.remove(file_path)
                    deleted_files.append(file_path)
                    print(f'[DELETE] 已删除文件: {file_path}')

        # 3. 删除 DataFiles/assets 目录下的所有源文件
        datafiles_assets_dir = os.path.join(Config.BASE_DIR, 'DataFiles', 'assets')
        if os.path.exists(datafiles_assets_dir):
            for root, _, files in os.walk(datafiles_assets_dir):
                for filename in files:
                    file_path = os.path.join(root, filename)
                    os.remove(file_path)
                    deleted_files.append(file_path)
                    print(f'[DELETE] 已删除文件: {file_path}')

        return {
            'success': True,
            'message': f'成功删除 {len(deleted_files)} 个文件'
        }

    except Exception as e:
        print(f'[ERROR] 删除设备数据失败: {str(e)}')
        import traceback
        traceback.print_exc()
        return {
            'success': False,
            'error': str(e)
        }


def delete_merge_data():
    """删除合并结果数据"""
    try:
        deleted_files = []

        # 1. 删除 merge_results.db 数据库文件
        db_path = DatabaseConfig.MERGE_RESULTS_DB
        if os.path.exists(db_path):
            os.remove(db_path)
            deleted_files.append(db_path)
            print(f'[DELETE] 已删除数据库文件: {db_path}')

        # 2. 删除uploaded_files目录下的merge相关文件
        upload_dir = PathConfig.UPLOAD_FOLDER
        if os.path.exists(upload_dir):
            for filename in os.listdir(upload_dir):
                if 'merge' in filename.lower() and filename.endswith(('.xlsx', '.xls', '.db')):
                    file_path = os.path.join(upload_dir, filename)
                    os.remove(file_path)
                    deleted_files.append(file_path)
                    print(f'[DELETE] 已删除文件: {file_path}')

        # 3. 删除 DataFiles/merge_results 目录下的所有源文件
        datafiles_merge_dir = os.path.join(Config.BASE_DIR, 'DataFiles', 'merge_results')
        if os.path.exists(datafiles_merge_dir):
            for root, _, files in os.walk(datafiles_merge_dir):
                for filename in files:
                    file_path = os.path.join(root, filename)
                    os.remove(file_path)
                    deleted_files.append(file_path)
                    print(f'[DELETE] 已删除文件: {file_path}')

        return {
            'success': True,
            'message': f'成功删除 {len(deleted_files)} 个文件'
        }

    except Exception as e:
        print(f'[ERROR] 删除合并数据失败: {str(e)}')
        import traceback
        traceback.print_exc()
        return {
            'success': False,
            'error': str(e)
        }



MAPPING_CONFIG_FILE = os.path.join(Config.BASE_DIR, 'config', 'mapping_config.json')
LOGIC_RULES_FILE = os.path.join(Config.BASE_DIR, 'config', 'logic_rules.json')
CUSTOM_REPORTS_FILE = os.path.join(Config.BASE_DIR, 'config', 'custom_reports.json')
BUILTIN_REPORT_CODES = {
    'yeji': {
        '10001', '10002', '10004',
        'i_10600_10001', 'i_10600_10002', 'i_10600_10004'
    },
    'smc': {
        '附件三', '附件五',
        'attachment_3', 'attachment_5'
    },
    'xinan': {
        '00000', '10001',
        'i_10600_00000', 'i_10600_10001'
    }
}


def _mapping_candidate_keys(category, report_code):
    """Build candidate keys for resolving an existing mapping config."""
    rc = str(report_code or '').strip()
    if not rc:
        return []

    rc_short = rc.replace('i_10600_', '', 1) if rc.startswith('i_10600_') else rc

    if category == 'yeji':
        preferred_keys = [f'i_10600_{rc_short}', rc]
    elif category == 'smc':
        zh_to_en = {'附件三': 'attachment_3', '附件五': 'attachment_5'}
        en_to_zh = {'attachment_3': '附件三', 'attachment_5': '附件五'}
        canonical = zh_to_en.get(rc, rc)
        preferred_keys = [canonical, rc, en_to_zh.get(canonical, canonical)]
    elif category == 'xinan':
        preferred_keys = [rc_short, rc, f'i_10600_{rc_short}']
    else:
        preferred_keys = [rc]

    common_keys = [f'i_10600_{rc_short}', f'i_{rc_short}', f'smc_{rc_short}', rc]
    return list(dict.fromkeys([key for key in preferred_keys + common_keys if key]))


def _resolve_mapping_config_key(mapping_config, category, report_code):
    """Resolve the actual persisted mapping-config key for a report code."""
    category_reports = (mapping_config.get('reports') or {}).get(category) or {}
    if not category_reports:
        return None

    rc = str(report_code or '').strip()
    if not rc:
        return None

    rc_short = rc.replace('i_10600_', '', 1) if rc.startswith('i_10600_') else rc

    for key in _mapping_candidate_keys(category, report_code):
        if key in category_reports:
            return key

    for key in category_reports:
        if key.endswith(f'_{rc_short}') or key == rc:
            return key

    return None


def _normalize_mapping_storage_key(category, report_code):
    """Normalize a report code to the storage key used in mapping_config.json."""
    rc = str(report_code or '').strip()
    if not rc:
        return ''

    rc_short = rc.replace('i_10600_', '', 1) if rc.startswith('i_10600_') else rc

    if category == 'yeji':
        if rc.startswith('i_10600_'):
            return rc
        if rc_short.isdigit() and len(rc_short) == 5:
            return f'i_10600_{rc_short}'
    elif category == 'smc':
        return {'附件三': 'attachment_3', '附件五': 'attachment_5'}.get(rc, rc)
    elif category == 'xinan':
        return rc_short

    return rc


def _load_custom_reports_payload():
    """读取自定义报表配置。"""
    if not os.path.exists(CUSTOM_REPORTS_FILE):
        return {
            'yeji': [],
            'smc': [],
            'xinan': []
        }

    with open(CUSTOM_REPORTS_FILE, 'r', encoding='utf-8') as f:
        payload = json.load(f)

    if not isinstance(payload, dict):
        return {
            'yeji': [],
            'smc': [],
            'xinan': []
        }

    return payload


def _find_report_metadata(category, report_code, actual_report_code=''):
    """查找指定报表的展示层元数据。"""
    custom_reports = _load_custom_reports_payload()
    report_list = custom_reports.get(category) or []
    candidate_codes = {
        str(report_code or '').strip(),
        str(actual_report_code or '').strip()
    }

    for candidate in _mapping_candidate_keys(category, report_code):
        candidate_codes.add(str(candidate or '').strip())

    normalized_candidates = {code for code in candidate_codes if code}
    for report in report_list:
        if not isinstance(report, dict):
            continue
        stored_codes = {
            str(report.get('code') or '').strip(),
            str(report.get('original_code') or report.get('code') or '').strip()
        }
        if normalized_candidates.intersection({code for code in stored_codes if code}):
            return copy.deepcopy(report)

    return None


def _collect_report_public_config_bindings(category, report_code, actual_report_code=''):
    """收集某个报表当前登记的公共配置绑定。"""
    payload = _load_public_config_bindings_payload()
    items = payload.get('items') or []
    candidate_codes = {
        str(report_code or '').strip(),
        str(actual_report_code or '').strip()
    }
    for candidate in _mapping_candidate_keys(category, report_code):
        candidate_codes.add(str(candidate or '').strip())
    normalized_candidates = {code for code in candidate_codes if code}

    bindings = []
    for item in items:
        if not isinstance(item, dict):
            continue
        if str(item.get('category') or '').strip() != str(category or '').strip():
            continue
        stored_code = str(item.get('report_code') or '').strip()
        if stored_code not in normalized_candidates:
            continue
        bindings.append(copy.deepcopy(item))
    return bindings


def _normalize_report_package_binding(binding, category, report_code, position):
    """标准化单报表包中的公共配置绑定。"""
    raw_binding = copy.deepcopy(binding) if isinstance(binding, dict) else {}
    target_column = str(raw_binding.get('target_column') or '').strip().upper()
    config_key = str(raw_binding.get('config_key') or '').strip()
    if not target_column or not config_key:
        return None

    normalized_binding = raw_binding
    normalized_binding['category'] = str(category or '').strip()
    normalized_binding['report_code'] = str(report_code or '').strip()
    normalized_binding['target_column'] = target_column
    normalized_binding['config_key'] = config_key
    normalized_binding['enabled'] = bool(normalized_binding.get('enabled', True))
    normalized_binding['priority'] = int(normalized_binding.get('priority') or 10)
    normalized_binding['conditions'] = (
        copy.deepcopy(normalized_binding.get('conditions'))
        if isinstance(normalized_binding.get('conditions'), dict)
        else {}
    )
    normalized_binding['remarks'] = str(normalized_binding.get('remarks') or '')
    normalized_binding['fallback_policy'] = str(
        normalized_binding.get('fallback_policy') or 'continue_with_raw_value'
    )
    normalized_binding['id'] = str(
        normalized_binding.get('id')
        or f'bind-{category}-{report_code}-{target_column.lower()}-{position:03d}'
    ).strip()
    return normalized_binding


def _is_builtin_report_reference(category, *report_codes):
    builtin_codes = BUILTIN_REPORT_CODES.get(category, set())
    for report_code in report_codes:
        normalized = str(report_code or '').strip()
        if normalized and normalized in builtin_codes:
            return True
    return False


def _sync_mapping_config_for_custom_report(category, existing_report, original_code, target_code, report_name):
    """Keep mapping config reachable when a mapping report code changes."""
    existing_code = ''
    if existing_report:
        existing_code = str(existing_report.get('code') or '').strip()

    should_sync = bool(
        (existing_code and existing_code != target_code) or
        (not existing_code and original_code != target_code)
    )
    if not should_sync or not os.path.exists(MAPPING_CONFIG_FILE):
        return

    with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8') as f:
        mapping_config = json.load(f)

    category_reports = mapping_config.setdefault('reports', {}).setdefault(category, {})
    source_candidates = []
    if existing_code and existing_code != target_code:
        source_candidates.append(existing_code)
    if original_code and original_code != target_code:
        source_candidates.append(original_code)

    source_key = None
    source_is_existing_custom_key = False
    normalized_existing_code = _normalize_mapping_storage_key(category, existing_code) if existing_code else ''
    for source_code in list(dict.fromkeys(source_candidates)):
        resolved_key = _resolve_mapping_config_key(mapping_config, category, source_code)
        if resolved_key:
            source_key = resolved_key
            source_is_existing_custom_key = bool(normalized_existing_code and normalized_existing_code == resolved_key)
            break

    if not source_key:
        return

    target_key = _normalize_mapping_storage_key(category, target_code)
    if not target_key or target_key == source_key:
        return

    migrated_config = copy.deepcopy(category_reports[source_key])
    if report_name:
        migrated_config['name'] = report_name
    migrated_config['category'] = category
    category_reports[target_key] = migrated_config

    if source_is_existing_custom_key and source_key in category_reports:
        del category_reports[source_key]

    with open(MAPPING_CONFIG_FILE, 'w', encoding='utf-8') as f:
        json.dump(mapping_config, f, ensure_ascii=False, indent=2)


def _is_conditional_mapping_rule(rule):
    """判断映射规则是否包含条件逻辑。"""
    if not isinstance(rule, dict):
        return False

    source_type = str(rule.get('source_type') or '').strip()
    if source_type in {'conditional', 'multi_conditional', 'conditional_groups'}:
        return True

    return bool(rule.get('conditions') or rule.get('groups'))


def _merge_conditional_rule_fields(source_rule, target_rule):
    """将条件逻辑字段从来源规则合并到目标规则。"""
    merged_rule = copy.deepcopy(target_rule if isinstance(target_rule, dict) else {})
    source_rule = source_rule if isinstance(source_rule, dict) else {}

    copied_fields = (
        'source_type',
        'source_value',
        'conditions',
        'default',
        'logic',
        'result',
        'groups',
        'group_logic',
        'remarks',
        'description',
        'transform',
    )
    for field in copied_fields:
        if field in source_rule:
            merged_rule[field] = copy.deepcopy(source_rule[field])

    return merged_rule


def _merge_conditions_only_mapping_config(source_mapping_config, target_mapping_config):
    """仅把条件逻辑合并到目标映射配置，保留目标模板和基础映射。"""
    merged_config = copy.deepcopy(target_mapping_config if isinstance(target_mapping_config, dict) else {})
    target_rules = merged_config.setdefault('mapping_rules', [])
    source_rules = [
        copy.deepcopy(rule)
        for rule in (source_mapping_config.get('mapping_rules') or [])
        if _is_conditional_mapping_rule(rule)
    ]

    if not source_rules:
        return merged_config

    for index, target_rule in enumerate(target_rules):
        target_column = str(target_rule.get('target_column') or '').strip()
        target_name = str(target_rule.get('target_name') or '').strip()
        matched_source_rule = None

        for source_rule in source_rules:
            source_column = str(source_rule.get('target_column') or '').strip()
            source_name = str(source_rule.get('target_name') or '').strip()
            if target_column and source_column and target_column == source_column:
                matched_source_rule = source_rule
                break
            if target_name and source_name and target_name == source_name:
                matched_source_rule = source_rule
                break

        if matched_source_rule:
            target_rules[index] = _merge_conditional_rule_fields(matched_source_rule, target_rule)

    return merged_config


def _apply_custom_report_overrides(base_report, target_code, target_name, target_report_payload=None):
    """复制自定义报表对象，并应用目标报表覆盖字段。"""
    copied_report = copy.deepcopy(base_report if isinstance(base_report, dict) else {})
    overrides = target_report_payload if isinstance(target_report_payload, dict) else {}

    copied_report['original_code'] = target_code
    copied_report['code'] = target_code
    copied_report['name'] = target_name
    copied_report['created_at'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    allowed_override_fields = (
        'description',
        'records',
        'notes',
        'tag',
        'status',
        'type',
        'template',
        'export',
        'data_source',
    )
    for field in allowed_override_fields:
        if field in overrides:
            copied_report[field] = copy.deepcopy(overrides[field])

    return copied_report


def _normalize_template_header_name(value):
    text = '' if value is None else str(value)
    text = text.replace('\r', ' ').replace('\n', ' ').strip()
    text = re.sub(r'^\*+', '', text)
    text = re.sub(r'[（(](必填|选填|必填项|选填项|可选)[)）]', '', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text


def _normalize_template_compare_key(value):
    return re.sub(r'\s+', '', _normalize_template_header_name(value)).lower()


def _analyze_template_file_payload(template_base_name):
    """分析模板文件并返回结构化列信息。"""
    if not template_base_name:
        raise ValueError('缺少template_base_name参数')

    found = False
    template_file_path = ''
    template_category_name = ''

    for cat_name in ('业支上报', 'SMC上报', '信安上报'):
        dir_path = os.path.join(PathConfig.TEMPLATES_DIR, cat_name)
        if os.path.exists(dir_path):
            for filename in os.listdir(dir_path):
                if filename.startswith(template_base_name) and filename.endswith('.xlsx'):
                    template_file_path = os.path.join(dir_path, filename)
                    template_category_name = cat_name
                    found = True
                    break
        if found:
            break

    if not found:
        raise FileNotFoundError(f'找不到模板文件: {template_base_name}')

    rel_path = f"Templates/{template_category_name}/{os.path.basename(template_file_path)}"

    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    wb = load_workbook(template_file_path)
    ws = wb.active

    def is_numeric_like(value):
        if value is None or isinstance(value, bool):
            return False
        if isinstance(value, (int, float)):
            return True
        return re.fullmatch(r'\d+(?:\.\d+)?', str(value).strip()) is not None

    def is_description_like(value):
        raw_text = '' if value is None else str(value)
        if not raw_text.strip():
            return False

        description_keywords = (
            '填写', '填报', '必填', '选填', '样例', '例如',
            '请参考', '请根据', '说明', '顺次增加', '多个'
        )
        if any(keyword in raw_text for keyword in description_keywords):
            return True
        if '\n' in raw_text:
            return True
        if len(raw_text.strip()) >= 20:
            return True
        if '：' in raw_text or ':' in raw_text:
            return True
        return False

    def detect_header_row():
        max_scan_rows = min(ws.max_row or 0, 10)
        max_scan_cols = ws.max_column or 0
        best_row = 1
        best_score = float('-inf')

        for row_idx in range(1, max_scan_rows + 1):
            non_empty_count = 0
            numeric_count = 0
            description_count = 0
            short_label_count = 0
            example_count = 0

            for col_idx in range(1, max_scan_cols + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is None or str(value).strip() == '':
                    continue

                non_empty_count += 1
                normalized_text = _normalize_template_header_name(value)

                if is_numeric_like(value):
                    numeric_count += 1
                    continue

                if '示例' in str(value) or '样例' in str(value):
                    example_count += 1

                if is_description_like(value):
                    description_count += 1

                if normalized_text and len(normalized_text) <= 20:
                    short_label_count += 1

            if non_empty_count < 2:
                continue

            text_count = non_empty_count - numeric_count
            score = (
                non_empty_count * 6 +
                text_count * 4 +
                short_label_count * 2 -
                numeric_count * 8 -
                description_count * 6 -
                example_count * 3
            )

            if score > best_score:
                best_score = score
                best_row = row_idx

        return best_row

    def detect_data_start_row(header_row):
        max_scan_cols = ws.max_column or 0
        max_row = ws.max_row or header_row

        for row_idx in range(header_row + 1, max_row + 1):
            has_value = False
            for col_idx in range(1, max_scan_cols + 1):
                value = ws.cell(row_idx, col_idx).value
                if value is not None and str(value).strip() != '':
                    has_value = True
                    break

            if not has_value:
                return row_idx

        return max_row + 1

    header_row = detect_header_row()
    data_start_row = detect_data_start_row(header_row)

    columns_info = []
    max_col = ws.max_column if ws.max_column else 0
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(header_row, col_idx)
        raw_value = cell.value
        normalized_name = _normalize_template_header_name(raw_value)
        columns_info.append({
            'index': col_idx,
            'letter': get_column_letter(col_idx),
            'name': normalized_name if normalized_name else f'列{col_idx}',
            'raw_name': str(raw_value) if raw_value else '',
            'value': str(raw_value) if raw_value else ''
        })

    wb.close()

    return {
        'template_file': rel_path,
        'header_row': header_row,
        'data_start_row': data_start_row,
        'total_columns': max_col,
        'columns': columns_info
    }


def get_mapping_config():
    """获取映射配置"""
    try:
        # 获取查询参数
        category = request.args.get('category')  # 'yeji', 'smc', 'xinan' or None (all)
        report_code = request.args.get('report_code')  # specific report code or None
        include_public = request.args.get('include_public', 'false').lower() == 'true'  # 是否包含公共配置

        config = _load_mapping_config_payload()

        # 如果请求包含公共配置
        if include_public:
            config['public_config'] = _build_public_config_entries(config)
            return jsonify({
                'success': True,
                'config': config  # 返回完整配置，包括public_config
            })

        # 如果指定了类别，只返回该类别的配置
        if category:
            if category in config['reports']:
                result = {
                    category: config['reports'][category]
                }
                # 如果还指定了报表代码，只返回该报表
                if report_code and report_code in config['reports'][category]:
                    result = {
                        category: {
                            report_code: config['reports'][category][report_code]
                        }
                    }
            else:
                result = {}
        else:
            result = config['reports']

        return jsonify({
            'success': True,
            'data': result
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def get_public_configs():
    """返回公共配置中心列表。"""
    try:
        mapping_config = _load_mapping_config_payload()
        public_config_entries = _build_public_config_entries(mapping_config)
        public_config_list = sorted(
            public_config_entries.values(),
            key=lambda item: item.get('key', '')
        )
        return jsonify({
            'success': True,
            'data': public_config_list
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def get_public_config_types():
    """返回统一公共配置类型注册表。"""
    try:
        registry = _build_public_config_registry()
        items = sorted(
            registry.values(),
            key=lambda item: item.get('key', '')
        )

        key_filter = str(request.args.get('key', '') or '').strip()
        type_filter = str(request.args.get('type', '') or '').strip()
        category_filter = str(request.args.get('category', '') or '').strip()
        enabled_filter = _parse_optional_bool_query_value(request.args.get('enabled'))

        if key_filter:
            items = [
                item for item in items
                if str(item.get('key', '') or '').strip() == key_filter
            ]

        if type_filter:
            items = [
                item for item in items
                if str(item.get('config_type', item.get('type', '')) or '').strip() == type_filter
            ]

        if category_filter:
            items = [
                item for item in items
                if str(item.get('category', '') or '').strip() == category_filter
            ]

        if enabled_filter is not None:
            items = [
                item for item in items
                if bool(item.get('enabled', False)) == enabled_filter
            ]

        payload = _load_public_config_definitions_payload()
        return jsonify({
            'success': True,
            'data': items,
            'total': len(items),
            'meta': {
                'version': payload.get('version') or '1.0',
                'last_updated': payload.get('last_updated') or '',
                'notes': payload.get('notes') or '',
                'registered_total': len(registry),
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def create_public_config_type():
    """新增公共配置类型定义。"""
    try:
        request_data = request.get_json(silent=True) or {}
        config_key = str(request_data.get('key') or '').strip()
        if not config_key:
            return jsonify({'success': False, 'error': '缺少类型标识 key'}), 400

        payload = _load_public_config_definitions_payload()
        existing_keys = {
            str(item.get('key') or '').strip()
            for item in payload.get('items', [])
            if isinstance(item, dict)
        }
        registry = _build_public_config_registry()
        if config_key in existing_keys or config_key in registry:
            return jsonify({'success': False, 'error': f'公共配置类型已存在: {config_key}'}), 409

        definition_item = _normalize_public_config_type_definition_for_save(request_data)
        payload['items'] = list(payload.get('items') or []) + [definition_item]
        payload['items'] = sorted(
            _normalize_public_config_center_items(payload.get('items')),
            key=lambda item: str(item.get('key') or '')
        )
        saved_payload = _save_public_config_definitions_payload(payload)

        return jsonify({
            'success': True,
            'message': '公共配置类型创建成功',
            'data': definition_item,
            'meta': {
                'version': saved_payload.get('version') or '2.0',
                'last_updated': saved_payload.get('last_updated') or '',
                'total_items': len(saved_payload.get('items') or []),
            }
        })
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def update_public_config_type(config_key):
    """更新公共配置类型定义。"""
    try:
        normalized_key = str(config_key or '').strip()
        request_data = request.get_json(silent=True) or {}
        payload = _load_public_config_definitions_payload()
        registry = _build_public_config_registry()

        existing_index = next(
            (
                index for index, item in enumerate(payload.get('items', []))
                if str(item.get('key') or '').strip() == normalized_key
            ),
            -1
        )
        existing_item = payload['items'][existing_index] if existing_index >= 0 else registry.get(normalized_key)
        if not isinstance(existing_item, dict):
            return jsonify({'success': False, 'error': f'未找到公共配置类型: {normalized_key}'}), 404

        requested_key = str(request_data.get('key') or normalized_key).strip()
        if requested_key != normalized_key:
            return jsonify({'success': False, 'error': '更新时不允许修改类型标识 key'}), 400

        definition_item = _normalize_public_config_type_definition_for_save(
            {
                **request_data,
                'key': normalized_key
            },
            existing_item
        )

        if existing_index >= 0:
            payload['items'][existing_index] = definition_item
        else:
            payload['items'] = list(payload.get('items') or []) + [definition_item]

        payload['items'] = sorted(
            _normalize_public_config_center_items(payload.get('items')),
            key=lambda item: str(item.get('key') or '')
        )
        saved_payload = _save_public_config_definitions_payload(payload)

        return jsonify({
            'success': True,
            'message': '公共配置类型更新成功',
            'data': definition_item,
            'meta': {
                'version': saved_payload.get('version') or '2.0',
                'last_updated': saved_payload.get('last_updated') or '',
                'total_items': len(saved_payload.get('items') or []),
            }
        })
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def handle_public_config_records(config_key):
    """通用公共配置记录接口。"""
    try:
        _, registry_entry = _resolve_public_config_registry_entry(config_key)
        payload = _load_public_config_storage_payload(config_key, registry_entry)

        if request.method == 'GET':
            return jsonify({
                'success': True,
                'data': payload
            })

        request_data = request.get_json(silent=True) or {}
        current_records = payload.get('mappings', [])
        next_id = max(
            [int(item.get('id', 0)) for item in current_records if str(item.get('id', '')).isdigit()],
            default=0
        ) + 1
        new_record = _build_public_config_record_payload(next_id, request_data, None, registry_entry)
        _validate_public_config_mapping_conflict(new_record, current_records, registry_entry)
        current_records.append(new_record)
        payload['mappings'] = current_records
        saved_payload = _save_public_config_storage_payload(config_key, payload, registry_entry)

        return jsonify({
            'success': True,
            'data': new_record,
            'total_count': saved_payload.get('total_count', len(saved_payload.get('mappings', []))),
            'message': '添加成功'
        })
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def handle_public_config_record_item(config_key, record_id):
    """通用公共配置单条记录接口。"""
    try:
        _, registry_entry = _resolve_public_config_registry_entry(config_key)
        payload = _load_public_config_storage_payload(config_key, registry_entry)
        current_records = payload.get('mappings', [])

        target_index = next(
            (index for index, item in enumerate(current_records) if int(item.get('id', 0)) == int(record_id)),
            -1
        )
        if target_index < 0:
            return jsonify({'success': False, 'error': '记录不存在'}), 404

        if request.method == 'DELETE':
            del current_records[target_index]
            payload['mappings'] = current_records
            _save_public_config_storage_payload(config_key, payload, registry_entry)
            return jsonify({'success': True, 'message': '删除成功'})

        request_data = request.get_json(silent=True) or {}
        updated_record = _build_public_config_record_payload(
            record_id,
            request_data,
            current_records[target_index],
            registry_entry
        )
        _validate_public_config_mapping_conflict(
            updated_record,
            current_records,
            registry_entry,
            exclude_record_id=record_id
        )
        current_records[target_index] = updated_record
        payload['mappings'] = current_records
        _save_public_config_storage_payload(config_key, payload, registry_entry)
        return jsonify({'success': True, 'data': updated_record, 'message': '更新成功'})
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def import_public_config_records(config_key):
    """通用公共配置 Excel 导入接口。"""
    try:
        _, registry_entry = _resolve_public_config_registry_entry(config_key)
        if 'file' not in request.files:
            return jsonify({'success': False, 'error': '未找到上传文件'}), 400

        file = request.files['file']
        if not file.filename:
            return jsonify({'success': False, 'error': '未选择文件'}), 400
        if not file.filename.endswith(('.xlsx', '.xls')):
            return jsonify({'success': False, 'error': '仅支持 Excel 文件（.xlsx/.xls）'}), 400

        file.stream.seek(0)
        df = pd.read_excel(file.stream, engine='openpyxl')

        fields = _build_public_config_excel_fields(registry_entry)
        required_headers = [
            field.get('label') or field.get('name')
            for field in fields
            if field.get('required')
        ]
        missing_headers = [
            header for header in required_headers
            if header not in df.columns and header not in [field.get('name') for field in fields]
        ]
        if missing_headers:
            return jsonify({
                'success': False,
                'error': f"Excel 文件缺少必填列: {', '.join(missing_headers)}"
            }), 400

        payload = _load_public_config_storage_payload(config_key, registry_entry)
        current_records = payload.get('mappings', [])
        primary_field = str(registry_entry.get('primary_field') or (fields[0].get('name') if fields else 'source')).strip()
        next_id = max(
            [int(item.get('id', 0)) for item in current_records if str(item.get('id', '')).isdigit()],
            default=0
        )

        imported_count = 0
        updated_count = 0
        skipped_count = 0

        for _, row in df.iterrows():
            record_data = {}
            for field in fields:
                field_name = field.get('name')
                if not field_name:
                    continue
                row_value = _extract_public_config_row_value(row, field)
                if pd.isna(row_value):
                    row_value = None
                record_data[field_name] = row_value

            primary_value = record_data.get(primary_field)
            if primary_value in (None, ''):
                skipped_count += 1
                continue

            existing_index = next(
                (
                    index for index, item in enumerate(current_records)
                    if str(item.get(primary_field, '')).strip() == str(primary_value).strip()
                ),
                -1
            )

            if existing_index >= 0:
                record_id = int(current_records[existing_index].get('id', existing_index + 1))
                candidate_record = _build_public_config_record_payload(
                    record_id,
                    record_data,
                    current_records[existing_index],
                    registry_entry
                )
                _validate_public_config_mapping_conflict(
                    candidate_record,
                    current_records,
                    registry_entry,
                    exclude_record_id=record_id
                )
                current_records[existing_index] = candidate_record
                updated_count += 1
                continue

            next_id += 1
            candidate_record = _build_public_config_record_payload(next_id, record_data, None, registry_entry)
            _validate_public_config_mapping_conflict(candidate_record, current_records, registry_entry)
            current_records.append(
                candidate_record
            )
            imported_count += 1

        payload['mappings'] = current_records
        saved_payload = _save_public_config_storage_payload(config_key, payload, registry_entry)
        return jsonify({
            'success': True,
            'imported_count': imported_count,
            'updated_count': updated_count,
            'skipped_count': skipped_count,
            'total_count': saved_payload.get('total_count', len(saved_payload.get('mappings', []))),
            'message': f'导入成功：新增 {imported_count} 条，更新 {updated_count} 条，跳过 {skipped_count} 条'
        })
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': f'导入失败: {str(e)}'}), 500


def export_public_config_records(config_key):
    """通用公共配置 Excel 导出接口。"""
    try:
        _, registry_entry = _resolve_public_config_registry_entry(config_key)
        payload = _load_public_config_storage_payload(config_key, registry_entry)
        fields = _build_public_config_excel_fields(registry_entry)
        mappings = payload.get('mappings', [])
        rows = []
        for item in mappings:
            row = {}
            for field in fields:
                label = field.get('label') or field.get('name')
                row[label] = _convert_public_config_excel_cell_value(
                    field,
                    item.get(field.get('name'))
                )
            rows.append(row)

        df = pd.DataFrame(rows, columns=[field.get('label') or field.get('name') for field in fields])
        filename = os.path.basename(
            str(
                payload.get('export_filename')
                or registry_entry.get('export_filename')
                or f"{payload.get('name') or config_key}_导出.xlsx"
            )
        ) or f"{config_key}_导出.xlsx"
        output = BytesIO()

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            sheet_name = str(payload.get('name') or config_key)[:31]
            df.to_excel(writer, index=False, sheet_name=sheet_name)

        output.seek(0)

        response = send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Content-Disposition'] = (
            f"attachment; filename=\"public-config-export.xlsx\"; "
            f"filename*=UTF-8''{quote(filename)}"
        )
        return response
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def download_public_config_template(config_key):
    """通用公共配置模板下载接口。"""
    try:
        _, registry_entry = _resolve_public_config_registry_entry(config_key)
        payload = _load_public_config_storage_payload(config_key, registry_entry)
        fields = _build_public_config_excel_fields(registry_entry)
        headers = [field.get('label') or field.get('name') for field in fields]
        filename = os.path.basename(
            str(
                payload.get('template_filename')
                or registry_entry.get('template_filename')
                or f"{payload.get('name') or config_key}_导入模板.xlsx"
            )
        ) or f"{config_key}_导入模板.xlsx"
        output = BytesIO()

        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            pd.DataFrame(columns=headers).to_excel(
                writer,
                index=False,
                sheet_name=str(payload.get('name') or config_key)[:31]
            )
            workbook = writer.book
            info_sheet = workbook.create_sheet('填写说明')
            info_rows = [
                [f"{payload.get('name') or config_key} 导入模板"],
                [''],
                ['字段说明：']
            ]
            for field in fields:
                requirement = '必填' if field.get('required') else '可选'
                info_rows.append([
                    f"{field.get('label') or field.get('name')}（{field.get('name')}）",
                    f"类型：{field.get('type') or 'text'}，{requirement}"
                ])

            for row_index, row_values in enumerate(info_rows, start=1):
                for column_index, cell_value in enumerate(row_values, start=1):
                    info_sheet.cell(row=row_index, column=column_index, value=cell_value)

        output.seek(0)

        response = send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response.headers['Content-Disposition'] = (
            f"attachment; filename=\"public-config-template.xlsx\"; "
            f"filename*=UTF-8''{quote(filename)}"
        )
        return response
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def get_public_config_center_definitions():
    """返回公共配置中心定义层列表。"""
    try:
        payload = _load_public_config_definitions_payload()
        items = payload.get('items') or []

        key_filter = str(request.args.get('key', '') or '').strip()
        type_filter = str(request.args.get('type', '') or '').strip()
        enabled_filter = _parse_optional_bool_query_value(request.args.get('enabled'))

        if key_filter:
            items = [
                item for item in items
                if str(item.get('key', '') or '').strip() == key_filter
            ]

        if type_filter:
            items = [
                item for item in items
                if str(item.get('type', '') or '').strip() == type_filter
            ]

        if enabled_filter is not None:
            items = [
                item for item in items
                if bool(item.get('enabled', False)) == enabled_filter
            ]

        return jsonify({
            'success': True,
            'data': items,
            'total': len(items),
            'meta': {
                'version': payload.get('version') or '1.0',
                'last_updated': payload.get('last_updated') or '',
                'notes': payload.get('notes') or '',
                'total_items': len(payload.get('items') or []),
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def get_public_config_center_bindings():
    """返回公共配置中心绑定层列表。"""
    try:
        payload = _load_public_config_bindings_payload()
        items = payload.get('items') or []

        config_key_filter = str(request.args.get('config_key', '') or '').strip()
        category_filter = str(request.args.get('category', '') or '').strip()
        report_code_filter = str(request.args.get('report_code', '') or '').strip()
        target_column_filter = str(request.args.get('target_column', '') or '').strip().upper()
        enabled_filter = _parse_optional_bool_query_value(request.args.get('enabled'))

        if config_key_filter:
            items = [
                item for item in items
                if str(item.get('config_key', '') or '').strip() == config_key_filter
            ]

        if category_filter:
            items = [
                item for item in items
                if str(item.get('category', '') or '').strip() == category_filter
            ]

        if report_code_filter:
            items = [
                item for item in items
                if str(item.get('report_code', '') or '').strip() == report_code_filter
            ]

        if target_column_filter:
            items = [
                item for item in items
                if str(item.get('target_column', '') or '').strip().upper() == target_column_filter
            ]

        if enabled_filter is not None:
            items = [
                item for item in items
                if bool(item.get('enabled', False)) == enabled_filter
            ]

        return jsonify({
            'success': True,
            'data': items,
            'total': len(items),
            'meta': {
                'version': payload.get('version') or '1.0',
                'last_updated': payload.get('last_updated') or '',
                'notes': payload.get('notes') or '',
                'total_items': len(payload.get('items') or []),
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def get_dynamic_config_query_definitions():
    """获取第三级动态配置查询定义层。"""
    try:
        payload = _load_dynamic_config_query_definitions_payload()
        return jsonify({
            'success': True,
            'data': payload,
            'meta': {
                'has_saved_definitions': bool(payload.get('definitions')),
                'total_definitions': len(payload.get('definitions') or {}),
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def save_dynamic_config_query_definitions():
    """保存第三级动态配置查询定义层。"""
    try:
        request_data = request.get_json(silent=True) or {}
        normalized_payload = _normalize_dynamic_config_query_definitions_payload(request_data)
        if not isinstance(normalized_payload.get('definitions'), dict):
            return jsonify({'success': False, 'error': 'definitions 格式无效'}), 400

        saved_payload = _save_dynamic_config_query_definitions_payload(normalized_payload)
        return jsonify({
            'success': True,
            'message': '动态配置查询定义已保存',
            'data': saved_payload,
            'meta': {
                'total_definitions': len(saved_payload.get('definitions') or {}),
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def export_report_mapping_package(category, report_code):
    """导出单报表完整映射包。"""
    try:
        mapping_config = _load_mapping_config_payload()
        actual_report_code = _resolve_mapping_config_key(mapping_config, category, report_code)
        if not actual_report_code:
            return jsonify({'success': False, 'error': '未找到该报表配置'}), 404

        category_reports = (mapping_config.get('reports') or {}).get(category) or {}
        report_mapping_config = copy.deepcopy(category_reports.get(actual_report_code) or {})
        if not report_mapping_config:
            return jsonify({'success': False, 'error': '报表映射配置为空'}), 404

        bindings = _collect_report_public_config_bindings(category, report_code, actual_report_code)
        referenced_public_configs = sorted({
            str(rule.get('mapping_config') or '').strip()
            for rule in (report_mapping_config.get('mapping_rules') or [])
            if isinstance(rule, dict) and str(rule.get('mapping_config') or '').strip()
        })

        package_payload = {
            'package_type': 'report_mapping_package',
            'version': '1.0',
            'exported_at': datetime.now().isoformat(timespec='seconds'),
            'category': str(category or '').strip(),
            'requested_report_code': str(report_code or '').strip(),
            'actual_report_code': actual_report_code,
            'report_meta': _find_report_metadata(category, report_code, actual_report_code),
            'mapping_config': report_mapping_config,
            'public_config_bindings': bindings,
            'dynamic_config_query': _load_dynamic_config_query_definitions_payload(),
            'referenced_public_configs': referenced_public_configs,
        }

        return jsonify({
            'success': True,
            'data': package_payload
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def import_report_mapping_package(category, report_code):
    """导入单报表完整映射包。"""
    try:
        request_data = request.get_json(silent=True) or {}
        package_payload = (
            request_data.get('package')
            if isinstance(request_data.get('package'), dict)
            else request_data
        )
        if not isinstance(package_payload, dict):
            return jsonify({'success': False, 'error': '缺少可导入的报表包数据'}), 400

        report_mapping_config = copy.deepcopy(
            package_payload.get('mapping_config')
            or package_payload.get('config')
            or {}
        )
        if not isinstance(report_mapping_config, dict):
            return jsonify({'success': False, 'error': '报表包中的 mapping_config 格式无效'}), 400
        if not isinstance(report_mapping_config.get('mapping_rules'), list):
            return jsonify({'success': False, 'error': '报表包缺少 mapping_rules'}), 400

        normalized_storage_key = _normalize_mapping_storage_key(category, report_code)
        if not normalized_storage_key:
            return jsonify({'success': False, 'error': '目标报表编码无效'}), 400

        mapping_payload = _load_mapping_config_payload()
        category_reports = mapping_payload.setdefault('reports', {}).setdefault(category, {})
        report_mapping_config['category'] = report_mapping_config.get('category') or category

        imported_report_meta = package_payload.get('report_meta')
        if not str(report_mapping_config.get('name') or '').strip():
            if isinstance(imported_report_meta, dict):
                report_mapping_config['name'] = str(
                    imported_report_meta.get('name')
                    or imported_report_meta.get('display_name')
                    or report_code
                ).strip()
            else:
                report_mapping_config['name'] = str(report_code or '').strip()

        category_reports[normalized_storage_key] = report_mapping_config

        alias_keys = set(_mapping_candidate_keys(category, report_code))

        for alias in alias_keys:
            if alias and alias != normalized_storage_key and alias in category_reports:
                del category_reports[alias]

        with open(MAPPING_CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump(mapping_payload, f, ensure_ascii=False, indent=2)

        imported_bindings = package_payload.get('public_config_bindings')
        saved_bindings_count = 0
        if isinstance(imported_bindings, list):
            binding_payload = _load_public_config_bindings_payload()
            binding_items = binding_payload.get('items') or []
            next_binding_items = []
            removable_codes = set(alias_keys)
            removable_codes.add(normalized_storage_key)
            normalized_removable_codes = {code for code in removable_codes if code}
            for item in binding_items:
                if not isinstance(item, dict):
                    continue
                if (
                    str(item.get('category') or '').strip() == str(category or '').strip()
                    and str(item.get('report_code') or '').strip() in normalized_removable_codes
                ):
                    continue
                next_binding_items.append(copy.deepcopy(item))

            normalized_bindings = []
            for index, binding in enumerate(imported_bindings, start=1):
                normalized_binding = _normalize_report_package_binding(
                    binding, category, normalized_storage_key, index
                )
                if normalized_binding:
                    normalized_bindings.append(normalized_binding)

            binding_payload['items'] = next_binding_items + normalized_bindings
            _save_public_config_bindings_payload(binding_payload)
            saved_bindings_count = len(normalized_bindings)

        dynamic_config_payload = (
            package_payload.get('frontend_dynamic_config_query')
            or package_payload.get('dynamic_config_query')
            or package_payload.get('dynamic_definitions')
        )
        saved_dynamic_config_payload = None
        if isinstance(dynamic_config_payload, dict):
            saved_dynamic_config_payload = _save_dynamic_config_query_definitions_payload(
                dynamic_config_payload
            )

        return jsonify({
            'success': True,
            'message': f'已导入 {category}/{normalized_storage_key} 报表包',
            'data': {
                'actual_report_code': normalized_storage_key,
                'mapping_config': report_mapping_config,
                'bindings_count': saved_bindings_count,
                'dynamic_config_query': saved_dynamic_config_payload,
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def get_report_mapping_config(category, report_code):
    """获取特定报表的映射配置

    支持短代码匹配：如果传入 10001，会自动查找 i_10600_10001
    """
    try:
        with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8') as f:
            config = json.load(f)

        # 归一化并优先匹配“导出实际使用”的键名
        if category in config['reports']:
            category_reports = config['reports'][category]
            rc = str(report_code).strip()
            if rc.startswith('i_10600_'):
                rc_short = rc.replace('i_10600_', '', 1)
            else:
                rc_short = rc

            preferred_keys = []
            if category == 'yeji':
                preferred_keys = [f'i_10600_{rc_short}', rc]
            elif category == 'smc':
                zh_to_en = {'附件三': 'attachment_3', '附件五': 'attachment_5'}
                en_to_zh = {'attachment_3': '附件三', 'attachment_5': '附件五'}
                canonical = zh_to_en.get(rc, rc)
                preferred_keys = [canonical, rc, en_to_zh.get(canonical, canonical)]
            elif category == 'xinan':
                preferred_keys = [rc_short, rc, f'i_10600_{rc_short}']
            else:
                preferred_keys = [rc]

            for key in list(dict.fromkeys(preferred_keys)):
                if key in category_reports:
                    return jsonify({
                        'success': True,
                        'data': category_reports[key],
                        'matched_key': key
                    })

            # 兜底：尝试常见前缀和后缀模糊匹配
            possible_keys = [f'i_10600_{rc_short}', f'i_{rc_short}', f'smc_{rc_short}', rc]
            for key in possible_keys:
                if key in category_reports:
                    return jsonify({
                        'success': True,
                        'data': category_reports[key],
                        'matched_key': key
                    })

            for key in category_reports:
                if key.endswith(f'_{rc_short}') or key == rc:
                    return jsonify({
                        'success': True,
                        'data': category_reports[key],
                        'matched_key': key
                    })

        return jsonify({'success': False, 'error': '未找到该报表配置'}), 404
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/export-scope/merge-source-level-stats', methods=['GET'])
def get_merge_export_scope_source_level_stats():
    """获取合并结果整体筛选弹窗的统计视图数据。"""
    try:
        result = _collect_merge_export_scope_source_level_stats()
        return jsonify({'success': True, **result})
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/export-scope/merge-source-level-stats/export', methods=['GET'])
def export_merge_export_scope_source_level_stats():
    """导出合并结果整体筛选统计表。"""
    try:
        result = _collect_merge_export_scope_source_level_stats()
        items = result.get('items') or []
        summary = result.get('summary') or {}
        level_totals = summary.get('level_totals') or {}
        registry_map = {}
        conn = None
        try:
            conn = sqlite3.connect(DatabaseConfig.MERGE_RESULTS_DB)
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()
            cursor.execute(
                '''
                SELECT source_name, last_mode, ever_incremental, last_import_time
                FROM merge_source_registry
                '''
            )
            for row in cursor.fetchall():
                source_name = str(row['source_name'] or '').strip()
                if not source_name:
                    continue
                last_mode = str(row['last_mode'] or '').strip().lower()
                ever_incremental = bool(int(row['ever_incremental'] or 0))
                if last_mode == 'incremental':
                    update_mode_label = '增量更新'
                elif last_mode == 'full' and ever_incremental:
                    update_mode_label = '曾增量，最近全量'
                elif last_mode == 'full':
                    update_mode_label = '仅全量'
                else:
                    update_mode_label = '未标注'
                registry_map[source_name] = {
                    'update_mode_label': update_mode_label,
                    'last_import_time': str(row['last_import_time'] or '').strip()
                }
        except Exception:
            registry_map = {}
        finally:
            if conn:
                conn.close()

        export_rows = []
        ratio_rows = []
        for item in items:
            source_name = str(item.get('source_name') or '').strip()
            source_label = str(item.get('source_label') or '')
            total_count = int(item.get('total_count') or 0)
            level_1_count = int(item.get('level_1_count') or 0)
            level_2_count = int(item.get('level_2_count') or 0)
            level_3_count = int(item.get('level_3_count') or 0)
            level_4_count = int(item.get('level_4_count') or 0)
            registry = registry_map.get(source_name) or registry_map.get(source_label) or {}
            update_mode_label = str(registry.get('update_mode_label') or '未标注')
            last_import_time = str(registry.get('last_import_time') or '')
            export_rows.append({
                '数据源名称': source_label,
                '总行数': total_count,
                '1级': level_1_count,
                '2级': level_2_count,
                '3级': level_3_count,
                '4级': level_4_count,
                '更新标注': update_mode_label,
                '最近导入': last_import_time
            })
            ratio_rows.append({
                '数据源名称': source_label,
                '总行数': total_count,
                '1级占比(%)': round((level_1_count / total_count * 100.0), 2) if total_count > 0 else 0.0,
                '2级占比(%)': round((level_2_count / total_count * 100.0), 2) if total_count > 0 else 0.0,
                '3级占比(%)': round((level_3_count / total_count * 100.0), 2) if total_count > 0 else 0.0,
                '4级占比(%)': round((level_4_count / total_count * 100.0), 2) if total_count > 0 else 0.0,
                '更新标注': update_mode_label,
                '最近导入': last_import_time
            })

        # 底部追加“总计”行（数值列汇总）。
        total_rows = int(summary.get('total_rows') or 0)
        total_level_1 = int(level_totals.get('level_1_count') or 0)
        total_level_2 = int(level_totals.get('level_2_count') or 0)
        total_level_3 = int(level_totals.get('level_3_count') or 0)
        total_level_4 = int(level_totals.get('level_4_count') or 0)
        export_rows.append({
            '数据源名称': '总计',
            '总行数': total_rows,
            '1级': total_level_1,
            '2级': total_level_2,
            '3级': total_level_3,
            '4级': total_level_4,
            '更新标注': '-',
            '最近导入': ''
        })
        ratio_rows.append({
            '数据源名称': '总计',
            '总行数': total_rows,
            '1级占比(%)': round((total_level_1 / total_rows * 100.0), 2) if total_rows > 0 else 0.0,
            '2级占比(%)': round((total_level_2 / total_rows * 100.0), 2) if total_rows > 0 else 0.0,
            '3级占比(%)': round((total_level_3 / total_rows * 100.0), 2) if total_rows > 0 else 0.0,
            '4级占比(%)': round((total_level_4 / total_rows * 100.0), 2) if total_rows > 0 else 0.0,
            '更新标注': '-',
            '最近导入': ''
        })

        count_df = pd.DataFrame(
            export_rows,
            columns=[
                '数据源名称', '总行数', '1级', '2级', '3级', '4级', '更新标注', '最近导入'
            ]
        )
        ratio_df = pd.DataFrame(
            ratio_rows,
            columns=[
                '数据源名称', '总行数', '1级占比(%)', '2级占比(%)', '3级占比(%)', '4级占比(%)', '更新标注', '最近导入'
            ]
        )
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            count_df.to_excel(writer, index=False, sheet_name='数量统计')
            ratio_df.to_excel(writer, index=False, sheet_name='占比统计')
            workbook = writer.book
            for sheet_name in ('数量统计', '占比统计'):
                worksheet = workbook[sheet_name]
                worksheet.freeze_panes = 'A2'
        output.seek(0)

        filename = f'merge_scope_stats_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'

        try:
            log_operation = get_log_operation()
            log_operation(
                operation_type='data_export',
                page_type='merge',
                operation_desc='导出整体筛选统计',
                file_name=filename,
                record_count=len(items)
            )
        except Exception as log_err:
            print(f'[日志记录失败] {log_err}')

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/merge/fine-scope-policy', methods=['GET'])
def get_merge_fine_scope_policy():
    """获取合并结果精细化筛减策略参数。"""
    try:
        return jsonify({'success': True, 'policy': _load_merge_fine_scope_policy()})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/merge/fine-scope-policy', methods=['POST'])
def save_merge_fine_scope_policy():
    """保存合并结果精细化筛减策略参数。"""
    try:
        payload = request.get_json(silent=True) or {}
        policy = _save_merge_fine_scope_policy(payload.get('policy') or {})
        return jsonify({'success': True, 'policy': policy})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/merge/fine-scope-policy/preview', methods=['POST'])
def preview_merge_fine_scope_policy():
    """预演精细化筛减策略，仅计算不改库。"""
    try:
        payload = request.get_json(silent=True) or {}
        raw_policy = payload.get('policy')
        policy = _normalize_merge_fine_scope_policy(raw_policy) if isinstance(raw_policy, dict) else _load_merge_fine_scope_policy()
        stats_result = _collect_merge_export_scope_source_level_stats()
        preview = _build_merge_fine_scope_preview(stats_result, policy)
        return jsonify({'success': True, **preview})
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/merge/fine-scope-policy/export', methods=['GET'])
def export_merge_fine_scope_policy():
    """导出精细化筛减策略参数为 Excel。"""
    try:
        policy = _load_merge_fine_scope_policy()
        export_rows = [
            {'参数键': 'enabled', '参数值': '1' if policy.get('enabled') else '0', '说明': '是否启用精细化筛减（1启用，0关闭）'},
            {'参数键': 'max_target_rows', '参数值': policy.get('max_target_rows', 950000), '说明': '目标总行数 MAX'},
            {'参数键': 'tolerance_rows', '参数值': policy.get('tolerance_rows', 50000), '说明': '目标容差行数（最终允许 MAX±容差）'},
            {'参数键': 'manual_ratio_enabled', '参数值': '1' if policy.get('manual_ratio_enabled') else '0', '说明': '是否启用手动比例C（1启用，0自动计算）'},
            {'参数键': 'manual_ratio', '参数值': policy.get('manual_ratio', 0.0), '说明': '手动比例C（0~1）'},
            {'参数键': 'target_mode', '参数值': policy.get('target_mode', 'absolute'), '说明': '目标模式：absolute(按行数) 或 percent(按比例)'},
            {'参数键': 'target_min', '参数值': policy.get('target_min', 900000), '说明': '导出目标最小行数'},
            {'参数键': 'target_max', '参数值': policy.get('target_max', 950000), '说明': '导出目标最大行数'},
            {'参数键': 'preferred_target', '参数值': policy.get('preferred_target', 950000), '说明': '优先逼近目标行数'},
            {'参数键': 'target_min_percent', '参数值': policy.get('target_min_percent', 90.0), '说明': '导出目标最小比例(0-100)'},
            {'参数键': 'target_max_percent', '参数值': policy.get('target_max_percent', 95.0), '说明': '导出目标最大比例(0-100)'},
            {'参数键': 'preferred_target_percent', '参数值': policy.get('preferred_target_percent', 95.0), '说明': '优先逼近目标比例(0-100)'},
            {'参数键': 'small_source_threshold', '参数值': policy.get('small_source_threshold', 10000), '说明': '数据源总行数小于等于该值时不压缩'},
            {'参数键': 'rounding', '参数值': policy.get('rounding', 'largest_remainder'), '说明': '取整方式：largest_remainder 或 floor'},
            {'参数键': 'keep_at_least_one_when_nonzero', '参数值': '1' if policy.get('keep_at_least_one_when_nonzero') else '0', '说明': '分组非零时是否至少保留1条（1是，0否）'},
            {'参数键': 'fallback_when_small_exceeds_max', '参数值': policy.get('fallback_when_small_exceeds_max', 'warn_only'), '说明': '小源保留已超上限时策略：warn_only 或 compress_all'},
        ]

        df = pd.DataFrame(export_rows, columns=['参数键', '参数值', '说明'])
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='参数配置')
        output.seek(0)

        filename = f'merge_fine_scope_policy_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/merge/fine-scope-policy/import', methods=['POST'])
def import_merge_fine_scope_policy():
    """从 Excel 导入精细化筛减策略参数。"""
    try:
        if 'file' not in request.files:
            return jsonify({'success': False, 'error': '未选择文件'}), 400
        file = request.files['file']
        file_name = os.path.basename(str(file.filename or '').strip())
        if not file_name:
            return jsonify({'success': False, 'error': '文件名不能为空'}), 400

        df = pd.read_excel(file, sheet_name='参数配置')
        required_columns = {'参数键', '参数值'}
        if not required_columns.issubset(set(df.columns)):
            return jsonify({'success': False, 'error': '参数配置表缺少“参数键/参数值”列'}), 400

        raw_policy = {}
        for _, row in df.iterrows():
            key = str(row.get('参数键') or '').strip()
            if not key:
                continue
            value = row.get('参数值')
            if key in {'enabled', 'keep_at_least_one_when_nonzero', 'manual_ratio_enabled'}:
                text = str(value).strip().lower()
                raw_policy[key] = text in {'1', 'true', 'yes', 'y', 'on'}
            else:
                raw_policy[key] = value

        policy = _save_merge_fine_scope_policy(raw_policy)
        return jsonify({'success': True, 'policy': policy, 'message': '参数导入成功'})
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/logic-rules/<rule_id>/usages', methods=['GET'])
def get_logic_rule_usages(rule_id):
    """获取某条规则当前被哪些报表列引用。"""
    try:
        mapping_config = _load_mapping_config_payload()
        logic_rules = _sync_logic_rule_usage_metadata(rule_id, mapping_config, _load_logic_rules_payload())
        _save_logic_rules_payload(logic_rules)
        usages = _collect_logic_rule_usages(rule_id, mapping_config)
        return jsonify({
            'success': True,
            'data': {
                'rule_id': str(rule_id or '').strip(),
                'usages': usages,
                'usage_count': len(usages)
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/logic-rules/<rule_id>/usage-detail', methods=['GET'])
def get_logic_rule_usage_detail(rule_id):
    """获取某条规则的详细引用情况，包含明确引用与结构匹配。"""
    try:
        mapping_config = _load_mapping_config_payload()
        logic_rules = _load_logic_rules_payload()
        detail = _collect_logic_rule_usage_details(rule_id, mapping_config, logic_rules)
        explicit_usages = detail.get('explicit_usages') or []
        inferred_usages = detail.get('inferred_usages') or []
        return jsonify({
            'success': True,
            'data': {
                'rule_id': str(rule_id or '').strip(),
                'explicit_usages': explicit_usages,
                'explicit_count': len(explicit_usages),
                'inferred_usages': inferred_usages,
                'inferred_count': len(inferred_usages),
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/logic-rules/<rule_id>/apply', methods=['POST'])
def apply_logic_rule_to_reports(rule_id):
    """将规则库中的规则应用到已引用列或指定报表列。"""
    try:
        payload = request.get_json(silent=True) or {}
        logic_rules_payload = _load_logic_rules_payload()
        rules_map = logic_rules_payload.get('rules') or {}
        logic_rule = copy.deepcopy(rules_map.get(str(rule_id or '').strip()) or {})
        if not logic_rule:
            return jsonify({'success': False, 'error': '未找到对应规则'}), 404

        mapping_payload = _load_mapping_config_payload()
        targets = payload.get('targets') if isinstance(payload.get('targets'), list) else []
        linked_only = payload.get('linked_only') is True
        applied_targets = []

        if linked_only and not targets:
            linked_usages = _collect_logic_rule_usages(rule_id, mapping_payload)
            targets = [
                {
                    'category': item.get('category'),
                    'report_code': item.get('report_code'),
                    'target_column': item.get('target_column'),
                    'target_name': item.get('target_name')
                }
                for item in linked_usages
            ]

        for target in targets:
            if not isinstance(target, dict):
                continue

            category = str(target.get('category') or '').strip()
            report_code = str(target.get('report_code') or '').strip()
            target_column = str(target.get('target_column') or '').strip()
            target_name = str(target.get('target_name') or '').strip()
            if not category or not report_code:
                continue

            actual_report_code = _resolve_mapping_config_key(mapping_payload, category, report_code)
            if not actual_report_code:
                continue

            report_mapping_config = ((mapping_payload.get('reports') or {}).get(category) or {}).get(actual_report_code)
            if not isinstance(report_mapping_config, dict):
                continue

            mapping_rules = report_mapping_config.get('mapping_rules') or []
            if not isinstance(mapping_rules, list):
                continue

            for index, existing_rule in enumerate(mapping_rules):
                if not isinstance(existing_rule, dict):
                    continue

                is_match = False
                if linked_only:
                    is_match = str(existing_rule.get('logic_rule_id') or '').strip() == str(rule_id or '').strip()
                elif target_column:
                    is_match = str(existing_rule.get('target_column') or '').strip() == target_column
                elif target_name:
                    is_match = str(existing_rule.get('target_name') or '').strip() == target_name

                if not is_match:
                    continue

                mapping_rules[index] = _apply_logic_rule_to_mapping_rule(logic_rule, existing_rule)
                applied_targets.append({
                    'category': category,
                    'report_code': actual_report_code,
                    'report_name': report_mapping_config.get('name') or actual_report_code,
                    'target_column': str(existing_rule.get('target_column') or '').strip(),
                    'target_name': str(existing_rule.get('target_name') or '').strip(),
                })

        if not applied_targets:
            return jsonify({'success': False, 'error': '未找到可应用的报表列'}), 400

        _save_mapping_config_payload(mapping_payload)
        logic_rules_payload = _sync_logic_rule_usage_metadata(rule_id, mapping_payload, logic_rules_payload)
        _save_logic_rules_payload(logic_rules_payload)

        return jsonify({
            'success': True,
            'data': {
                'rule_id': str(rule_id or '').strip(),
                'applied_count': len(applied_targets),
                'targets': applied_targets
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/logic-rules/apply-all/preview', methods=['GET'])
def preview_apply_all_logic_rules():
    """预览全局应用规则库将影响的报表列。"""
    try:
        mapping_payload = _load_mapping_config_payload()
        targets = _collect_all_logic_rule_usage_targets(mapping_payload)
        distinct_rule_ids = sorted({str(item.get('logic_rule_id') or '').strip() for item in targets if str(item.get('logic_rule_id') or '').strip()})
        distinct_reports = sorted({f"{item.get('category')}/{item.get('report_code')}" for item in targets})
        return jsonify({
            'success': True,
            'data': {
                'rule_count': len(distinct_rule_ids),
                'report_count': len(distinct_reports),
                'target_count': len(targets),
                'targets': targets
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/logic-rules/apply-all', methods=['POST'])
def apply_all_logic_rules():
    """将规则库中的最新内容应用到所有已建立引用关系的报表列。"""
    try:
        mapping_payload = _load_mapping_config_payload()
        logic_rules_payload = _load_logic_rules_payload()
        rules_map = logic_rules_payload.get('rules') or {}
        reports = mapping_payload.get('reports') or {}
        applied_targets = []

        for category, category_reports in reports.items():
            if not isinstance(category_reports, dict):
                continue

            for report_code, report_config in category_reports.items():
                mapping_rules = report_config.get('mapping_rules') or []
                if not isinstance(mapping_rules, list):
                    continue

                for index, existing_rule in enumerate(mapping_rules):
                    if not isinstance(existing_rule, dict):
                        continue
                    logic_rule_id = str(existing_rule.get('logic_rule_id') or '').strip()
                    if not logic_rule_id or logic_rule_id not in rules_map:
                        continue

                    mapping_rules[index] = _apply_logic_rule_to_mapping_rule(rules_map[logic_rule_id], existing_rule)
                    applied_targets.append({
                        'logic_rule_id': logic_rule_id,
                        'logic_rule_name': str(rules_map[logic_rule_id].get('name') or '').strip(),
                        'category': category,
                        'report_code': report_code,
                        'report_name': report_config.get('name') or report_code,
                        'target_column': str(existing_rule.get('target_column') or '').strip(),
                        'target_name': str(existing_rule.get('target_name') or '').strip(),
                    })

        if not applied_targets:
            return jsonify({'success': False, 'error': '当前没有任何已建立引用关系的规则可应用'}), 400

        _save_mapping_config_payload(mapping_payload)
        logic_rules_payload = _sync_all_logic_rule_usage_metadata(mapping_payload, logic_rules_payload)
        _save_logic_rules_payload(logic_rules_payload)

        return jsonify({
            'success': True,
            'data': {
                'rule_count': len({item['logic_rule_id'] for item in applied_targets}),
                'report_count': len({f"{item['category']}/{item['report_code']}" for item in applied_targets}),
                'applied_count': len(applied_targets),
                'targets': applied_targets
            }
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def update_report_mapping_config(category, report_code):
    """更新特定报表的映射配置"""
    try:
        data = request.json

        with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8') as f:
            config = json.load(f)

        # 更新配置
        if category not in config['reports']:
            config['reports'][category] = {}

        # 智能解析报表代码：统一写入导出实际使用的键，避免同报表多键并存
        category_reports = config['reports'][category]
        rc = str(report_code).strip()
        rc_short = rc.replace('i_10600_', '', 1) if rc.startswith('i_10600_') else rc

        actual_report_code = rc
        if category == 'yeji':
            # 业支内置报表统一使用 i_10600_xxxxx
            if rc.startswith('i_10600_'):
                actual_report_code = rc
            elif rc_short.isdigit() and len(rc_short) == 5:
                actual_report_code = f'i_10600_{rc_short}'
        elif category == 'smc':
            actual_report_code = {'附件三': 'attachment_3', '附件五': 'attachment_5'}.get(rc, rc)
        elif category == 'xinan':
            # 信安内置报表统一使用短码（00000/10001）
            actual_report_code = rc_short

        category_reports[actual_report_code] = data

        # 清理同一报表的别名键，避免配置分叉
        alias_keys = []
        if category == 'yeji' and actual_report_code.startswith('i_10600_'):
            alias_keys.append(actual_report_code.replace('i_10600_', '', 1))
            alias_keys.append(rc)
        elif category == 'smc':
            reverse_zh = {'attachment_3': '附件三', 'attachment_5': '附件五'}
            alias_keys.extend([rc, reverse_zh.get(actual_report_code, '')])
        elif category == 'xinan':
            alias_keys.extend([rc, f'i_10600_{actual_report_code}'])

        for alias in list(dict.fromkeys([k for k in alias_keys if k])):
            if alias != actual_report_code and alias in category_reports:
                del category_reports[alias]

        # 保存配置
        with open(MAPPING_CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump(config, f, ensure_ascii=False, indent=2)

        return jsonify({
            'success': True,
            'message': f'已更新 {category}/{actual_report_code} 的映射配置',
            'actual_report_code': actual_report_code
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def add_report_mapping_config(category, report_code):
    """添加新的报表映射配置"""
    try:
        data = request.json

        with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8') as f:
            config = json.load(f)

        # 添加新配置
        if category not in config['reports']:
            config['reports'][category] = {}

        if report_code in config['reports'][category]:
            return jsonify({'success': False, 'error': '该报表配置已存在'}), 400

        config['reports'][category][report_code] = data

        # 保存配置
        with open(MAPPING_CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump(config, f, ensure_ascii=False, indent=2)

        return jsonify({
            'success': True,
            'message': f'已添加 {category}/{report_code} 的映射配置'
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def delete_report_mapping_config(category, report_code):
    """删除报表映射配置"""
    try:
        blocked_response = reject_cross_site_delete_request()
        if blocked_response:
            return blocked_response

        with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8') as f:
            config = json.load(f)

        if category in config['reports'] and report_code in config['reports'][category]:
            del config['reports'][category][report_code]

            # 保存配置
            with open(MAPPING_CONFIG_FILE, 'w', encoding='utf-8') as f:
                json.dump(config, f, ensure_ascii=False, indent=2)

            return jsonify({
                'success': True,
                'message': f'已删除 {category}/{report_code} 的映射配置'
            })
        else:
            return jsonify({'success': False, 'error': '未找到该报表配置'}), 404
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/data-source-columns/<source_type>', methods=['GET'])
def get_data_source_columns(source_type):
    """获取数据源表的列信息
    Args:
        source_type: 数据源类型，'assets' 或 'merge_results'
    """
    try:
        columns = []

        if source_type == 'assets':
            # 数据概览表
            conn = get_db_connection()
            cursor = conn.cursor()
            cursor.execute('PRAGMA table_info(assets)')
            columns_info = cursor.fetchall()
            conn.close()

            columns = [
                {
                    'index': idx,
                    'name': col[1],
                    'type': col[2],
                    'display_name': f"索引{idx}: {col[1]}"
                }
                for idx, col in enumerate(columns_info)
            ]

        elif source_type == 'merge_results':
            # 合并结果表（JSON字段）
            # 从配置文件中读取merge_columns.json
            merge_columns_file = os.path.join(Config.BASE_DIR, 'config', 'merge_columns.json')
            if os.path.exists(merge_columns_file):
                with open(merge_columns_file, 'r', encoding='utf-8') as f:
                    merge_data = json.load(f)

                # merge_columns.json格式：{"columns": [...], "pageSize": ..., ...}
                merge_columns_list = merge_data.get('columns', [])

                columns = [
                    {
                        'index': idx,
                        'name': col.get('name', col.get('field_name', '')),
                        'type': col.get('type', 'text'),
                        'display_name': f"索引{idx}: {col.get('name', col.get('field_name', ''))}"
                    }
                    for idx, col in enumerate(merge_columns_list)
                ]
            else:
                # 如果配置文件不存在，使用默认列
                default_columns = [
                    '实例名/数据库名(schema)', '表名', '字段类型', '字段名称', '空值',
                    '主键', '扫描数据类型', '字段数据分类', '字段数据分级', '数据要素',
                    '业务系统名称', '数据源IP', '数据库IP', '表所属系统', '系统类型', '业务系统负责人',
                    '业务系统负责人联系方式'
                ]
                columns = [
                    {
                        'index': idx,
                        'name': col,
                        'type': 'text',
                        'display_name': f"索引{idx}: {col}"
                    }
                    for idx, col in enumerate(default_columns)
                ]
        else:
            return jsonify({'success': False, 'error': f'未知的数据源类型: {source_type}'}), 400

        return jsonify({
            'success': True,
            'source_type': source_type,
            'columns': columns
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


# ==================== 自定义报表管理路由 ====================

def get_custom_reports():
    """获取用户自定义报表配置"""
    try:
        # 确保配置文件存在
        if not os.path.exists(CUSTOM_REPORTS_FILE):
            default_config = {
                'yeji': [],
                'smc': [],
                'xinan': []
            }
            with open(CUSTOM_REPORTS_FILE, 'w', encoding='utf-8') as f:
                json.dump(default_config, f, ensure_ascii=False, indent=2)

        with open(CUSTOM_REPORTS_FILE, 'r', encoding='utf-8') as f:
            custom_reports = json.load(f)

        return jsonify({
            'success': True,
            'data': custom_reports
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def save_custom_report():
    """保存用户自定义报表"""
    try:
        data = request.get_json()

        # 验证必填字段
        required_fields = ['category', 'code', 'name']
        for field in required_fields:
            if field not in data:
                return jsonify({'success': False, 'error': f'缺少必填字段: {field}'}), 400

        category = data['category']
        code = data['code']
        name = data['name']
        original_code = (data.get('original_code') or code).strip()

        # 读取现有配置
        if not os.path.exists(CUSTOM_REPORTS_FILE):
            custom_reports = {
                'yeji': [],
                'smc': [],
                'xinan': []
            }
        else:
            with open(CUSTOM_REPORTS_FILE, 'r', encoding='utf-8') as f:
                custom_reports = json.load(f)

        # 检查是否已存在
        if category not in custom_reports:
            custom_reports[category] = []

        existing_index = -1
        existing_report = None
        for i, report in enumerate(custom_reports[category]):
            report_original_code = report.get('original_code') or report.get('code')
            if report.get('code') == code or report_original_code == original_code:
                existing_index = i
                existing_report = report
                break

        requested_status = (data.get('status') or '').strip()
        allowed_statuses = {'completed', 'pending'}
        report_status = requested_status if requested_status in allowed_statuses else 'completed'

        # 创建报表对象
        report_obj = {
            'original_code': original_code,
            'code': code,
            'name': name,
            'description': data.get('description', ''),
            'records': data.get('records', 0),
            'notes': data.get('notes', ''),
            'tag': data.get('tag', ''),
            'status': report_status,
            'type': data.get('type', 'display'),
            'template': data.get('template', ''),
            'export': data.get('export', ''),
            'data_source': data.get('data_source', ''),
            'created_at': (
                existing_report.get('created_at')
                if existing_report and existing_report.get('created_at')
                else datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            )
        }

        if report_obj['type'] == 'mapping':
            _sync_mapping_config_for_custom_report(
                category=category,
                existing_report=existing_report,
                original_code=original_code,
                target_code=code,
                report_name=name
            )

        # 添加或更新
        if existing_index >= 0:
            custom_reports[category][existing_index] = report_obj
        else:
            custom_reports[category].append(report_obj)

        # 保存配置
        with open(CUSTOM_REPORTS_FILE, 'w', encoding='utf-8') as f:
            json.dump(custom_reports, f, ensure_ascii=False, indent=2)

        return jsonify({
            'success': True,
            'message': f'已保存自定义报表: {name}',
            'data': report_obj
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def save_report_settings():
    """保存报表显示层设置，系统预置报表仅保存覆盖项，自定义报表可更新业务参数。"""
    try:
        data = request.get_json() or {}

        category = str(data.get('category') or '').strip()
        report_code = str(data.get('report_code') or '').strip()
        if not category or not report_code:
            return jsonify({'success': False, 'error': '缺少必填字段: category/report_code'}), 400

        if not os.path.exists(CUSTOM_REPORTS_FILE):
            custom_reports = {
                'yeji': [],
                'smc': [],
                'xinan': []
            }
        else:
            with open(CUSTOM_REPORTS_FILE, 'r', encoding='utf-8') as f:
                custom_reports = json.load(f)

        custom_reports.setdefault(category, [])

        existing_index = -1
        existing_report = None
        for i, report in enumerate(custom_reports[category]):
            stored_code = str(report.get('code') or '').strip()
            stored_original_code = str(report.get('original_code') or stored_code).strip()
            if report_code in {stored_code, stored_original_code}:
                existing_index = i
                existing_report = report
                break

        is_builtin_override = bool(existing_report and existing_report.get('is_builtin_override'))
        if existing_report is None:
            if not _is_builtin_report_reference(category, report_code):
                return jsonify({'success': False, 'error': '未找到目标报表'}), 404

            existing_report = {
                'original_code': report_code,
                'code': report_code,
                'name': str(data.get('name') or report_code).strip() or report_code,
                'status': str(data.get('status') or 'completed').strip() or 'completed',
                'type': str(data.get('type') or 'mapping').strip() or 'mapping',
                'is_builtin_override': True
            }
            custom_reports[category].append(existing_report)
            existing_index = len(custom_reports[category]) - 1
            is_builtin_override = True

        report_obj = copy.deepcopy(existing_report)
        report_obj['code'] = str(report_obj.get('code') or report_code).strip() or report_code
        report_obj['original_code'] = str(report_obj.get('original_code') or report_code).strip() or report_code

        if is_builtin_override:
            report_obj['is_builtin_override'] = True
            report_obj['name'] = str(report_obj.get('name') or data.get('name') or report_code).strip() or report_code
            report_obj['status'] = str(report_obj.get('status') or data.get('status') or 'completed').strip() or 'completed'
            report_obj['type'] = str(report_obj.get('type') or data.get('type') or 'mapping').strip() or 'mapping'
        else:
            for field in ('data_source', 'template', 'export'):
                if field in data:
                    report_obj[field] = str(data.get(field) or '').strip()

        for field in ('display_name', 'display_code', 'tag', 'notes'):
            if field not in data:
                continue
            field_value = str(data.get(field) or '').strip()
            if field_value:
                report_obj[field] = field_value
            else:
                report_obj.pop(field, None)

        custom_reports[category][existing_index] = report_obj

        with open(CUSTOM_REPORTS_FILE, 'w', encoding='utf-8') as f:
            json.dump(custom_reports, f, ensure_ascii=False, indent=2)

        return jsonify({
            'success': True,
            'data': report_obj
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


def copy_custom_report():
    """复制已有自定义报表，并同步复制映射配置。"""
    try:
        data = request.get_json(silent=True) or {}

        required_fields = ['source_category', 'source_code', 'target_code', 'target_name']
        for field in required_fields:
            if not str(data.get(field, '')).strip():
                return jsonify({'success': False, 'error': f'缺少必填字段: {field}'}), 400

        source_category = str(data.get('source_category', '')).strip()
        target_category = str(data.get('target_category') or source_category).strip()
        source_code = str(data.get('source_code', '')).strip()
        target_code = str(data.get('target_code', '')).strip()
        target_name = str(data.get('target_name', '')).strip()
        copy_mode = str(data.get('copy_mode') or 'full').strip().lower()
        target_report_payload = data.get('target_report_payload') if isinstance(data.get('target_report_payload'), dict) else {}
        target_mapping_config_payload = (
            copy.deepcopy(data.get('target_mapping_config'))
            if isinstance(data.get('target_mapping_config'), dict)
            else None
        )

        if copy_mode not in {'full', 'conditions_only'}:
            return jsonify({'success': False, 'error': 'copy_mode 仅支持 full 或 conditions_only'}), 400

        if source_category != target_category:
            return jsonify({'success': False, 'error': '当前仅支持同类别复制自定义报表'}), 400

        if _is_builtin_report_reference(target_category, target_code):
            return jsonify({'success': False, 'error': '目标报表编码为系统内置编码，请更换后重试'}), 400

        if not os.path.exists(CUSTOM_REPORTS_FILE):
            return jsonify({'success': False, 'error': '自定义报表配置文件不存在'}), 404

        with open(CUSTOM_REPORTS_FILE, 'r', encoding='utf-8-sig') as f:
            custom_reports = json.load(f)

        category_reports = custom_reports.get(source_category) or []
        source_report = None
        for report in category_reports:
            report_code = str(report.get('code') or '').strip()
            report_original_code = str(report.get('original_code') or report_code).strip()
            if source_code in {report_code, report_original_code}:
                source_report = report
                break

        if not source_report:
            return jsonify({'success': False, 'error': '未找到源自定义报表'}), 404

        for report in category_reports:
            report_code = str(report.get('code') or '').strip()
            report_original_code = str(report.get('original_code') or report_code).strip()
            if target_code in {report_code, report_original_code}:
                return jsonify({'success': False, 'error': '目标报表编码已存在'}), 409

        copied_report = _apply_custom_report_overrides(
            base_report=source_report,
            target_code=target_code,
            target_name=target_name,
            target_report_payload=target_report_payload
        )

        if copied_report.get('type') == 'mapping':
            mapping_config = _load_mapping_config_payload()
            source_mapping_key = _resolve_mapping_config_key(mapping_config, source_category, source_code)
            if not source_mapping_key:
                return jsonify({'success': False, 'error': '源报表缺少对应的映射配置，无法复制'}), 404

            source_mapping_reports = (mapping_config.get('reports') or {}).get(source_category) or {}
            target_mapping_reports = mapping_config.setdefault('reports', {}).setdefault(target_category, {})
            target_mapping_key = _normalize_mapping_storage_key(target_category, target_code)

            if not target_mapping_key:
                return jsonify({'success': False, 'error': '目标报表编码无效'}), 400
            if target_mapping_key in target_mapping_reports:
                return jsonify({'success': False, 'error': '目标报表映射配置已存在'}), 409

            source_mapping_config = copy.deepcopy(source_mapping_reports[source_mapping_key])
            if copy_mode == 'conditions_only':
                if not target_mapping_config_payload:
                    return jsonify({'success': False, 'error': 'conditions_only 模式缺少目标映射草稿'}), 400
                copied_mapping_config = _merge_conditions_only_mapping_config(
                    source_mapping_config,
                    target_mapping_config_payload
                )
            else:
                copied_mapping_config = source_mapping_config

            if target_mapping_config_payload:
                for field in (
                    'data_source',
                    'data_source_mode',
                    'primary_data_source',
                    'template_file',
                    'data_start_row',
                    'filter_config',
                    'source_table',
                    'notes',
                ):
                    if field in target_mapping_config_payload:
                        copied_mapping_config[field] = copy.deepcopy(target_mapping_config_payload[field])

            copied_mapping_config['name'] = target_name
            copied_mapping_config['category'] = target_category
            target_mapping_reports[target_mapping_key] = copied_mapping_config

            with open(MAPPING_CONFIG_FILE, 'w', encoding='utf-8') as f:
                json.dump(mapping_config, f, ensure_ascii=False, indent=2)

        custom_reports.setdefault(target_category, []).append(copied_report)

        with open(CUSTOM_REPORTS_FILE, 'w', encoding='utf-8') as f:
            json.dump(custom_reports, f, ensure_ascii=False, indent=2)

        return jsonify({
            'success': True,
            'message': f'已复制自定义报表: {target_name}',
            'data': copied_report
        })
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


def delete_custom_report(category, code):
    """删除用户自定义报表"""
    try:
        blocked_response = reject_cross_site_delete_request()
        if blocked_response:
            return blocked_response

        if not os.path.exists(CUSTOM_REPORTS_FILE):
            return jsonify({'success': False, 'error': '配置文件不存在'}), 404

        with open(CUSTOM_REPORTS_FILE, 'r', encoding='utf-8') as f:
            custom_reports = json.load(f)

        if category not in custom_reports:
            return jsonify({'success': False, 'error': '类别不存在'}), 404

        # 查找报表
        report_to_delete = None
        new_list = []
        for report in custom_reports[category]:
            if report['code'] == code:
                report_to_delete = report
            else:
                new_list.append(report)

        if not report_to_delete:
            return jsonify({'success': False, 'error': '未找到该报表'}), 404

        # 删除报表记录
        custom_reports[category] = new_list

        # 保存配置
        with open(CUSTOM_REPORTS_FILE, 'w', encoding='utf-8') as f:
            json.dump(custom_reports, f, ensure_ascii=False, indent=2)

        deleted_export_files = []
        mapping_deleted = False
        protected_system_report = _is_builtin_report_reference(
            category,
            code,
            report_to_delete.get('original_code'),
            report_to_delete.get('code')
        )

        if not protected_system_report:
            # 标准化代码（兼容中文/英文代码）
            normalized_code = str(code).strip()
            if category == 'smc':
                smc_code_map = {
                    '附件三': 'attachment_3',
                    '附件五': 'attachment_5'
                }
                normalized_code = smc_code_map.get(normalized_code, normalized_code)
            elif category in ('yeji', 'xinan'):
                if normalized_code.startswith('i_10600_'):
                    normalized_code = normalized_code.replace('i_10600_', '', 1)

            # 删除对应类别的导出文件
            export_dir_map = {
                'yeji': PathConfig.YEZHI_EXPORT_DIR,
                'smc': PathConfig.SMC_EXPORT_DIR,
                'xinan': PathConfig.XINAN_EXPORT_DIR
            }
            export_dir = export_dir_map.get(category)

            if export_dir and os.path.exists(export_dir):
                for filename in os.listdir(export_dir):
                    if not filename.lower().endswith(('.xlsx', '.xls')):
                        continue

                    matched = False
                    if category in ('yeji', 'xinan'):
                        prefix = f'i_10600_{normalized_code}_'
                        matched = filename.startswith(prefix)
                    elif category == 'smc':
                        smc_keyword_map = {
                            'attachment_3': ['附件三', '附件3', 'attachment_3'],
                            'attachment_5': ['附件五', '附件5', 'attachment_5']
                        }
                        keywords = smc_keyword_map.get(normalized_code, [normalized_code])
                        matched = any(keyword in filename for keyword in keywords)

                    if not matched:
                        continue

                    try:
                        file_path = os.path.join(export_dir, filename)
                        os.remove(file_path)
                        deleted_export_files.append(filename)
                    except Exception as e:
                        print(f'[WARNING] 删除导出文件失败 {filename}: {e}')

            # 如果报表有映射配置，也需要删除（兼容多种键名）
            if os.path.exists(MAPPING_CONFIG_FILE):
                try:
                    with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8') as f:
                        mapping_config = json.load(f)

                    if category in mapping_config['reports']:
                        category_reports = mapping_config['reports'][category]
                        candidate_keys = [code, normalized_code]

                        if category in ('yeji', 'xinan'):
                            candidate_keys.append(f'i_10600_{normalized_code}')
                        elif category == 'smc':
                            reverse_smc_map = {
                                'attachment_3': '附件三',
                                'attachment_5': '附件五'
                            }
                            candidate_keys.append(reverse_smc_map.get(normalized_code, normalized_code))

                        for key in list(dict.fromkeys(candidate_keys)):
                            if key in category_reports:
                                del category_reports[key]
                                mapping_deleted = True

                        # 保存映射配置
                        with open(MAPPING_CONFIG_FILE, 'w', encoding='utf-8') as f:
                            json.dump(mapping_config, f, ensure_ascii=False, indent=2)
                        if mapping_deleted:
                            print(f'[DELETE] 已删除映射配置: {category}/{code}')
                except Exception as e:
                    print(f'[WARNING] 删除映射配置失败: {e}')
        else:
            print(f'[DELETE] 检测到系统预置报表覆盖记录，仅删除 custom_reports 记录，保留映射与导出文件: {category}/{code}')

        return jsonify({
            'success': True,
            'message': f'已删除自定义报表: {category}/{code}',
            'mapping_deleted': mapping_deleted,
            'deleted_export_files': deleted_export_files,
            'protected_system_report': protected_system_report
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


def analyze_template_file():
    """分析Excel模板文件，返回列信息"""
    try:
        data = request.get_json()
        template_base_name = data.get('template_base_name')
        template_info = _analyze_template_file_payload(template_base_name)

        return jsonify({
            'success': True,
            'data': template_info
        })
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except FileNotFoundError as e:
        return jsonify({'success': False, 'error': str(e)}), 404
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


def diff_template_file():
    """比较来源模板与目标模板，并识别失效映射规则。"""
    try:
        data = request.get_json(silent=True) or {}
        source_category = str(data.get('source_category') or '').strip()
        source_code = str(data.get('source_code') or '').strip()
        template_base_name = str(data.get('template_base_name') or '').strip()

        if not source_category:
            return jsonify({'success': False, 'error': '缺少source_category参数'}), 400
        if not source_code:
            return jsonify({'success': False, 'error': '缺少source_code参数'}), 400
        if not template_base_name:
            return jsonify({'success': False, 'error': '缺少template_base_name参数'}), 400

        mapping_config = _load_mapping_config_payload()
        source_mapping_key = _resolve_mapping_config_key(mapping_config, source_category, source_code)
        if not source_mapping_key:
            return jsonify({'success': False, 'error': '未找到来源报表的映射配置'}), 404

        source_mapping_config = ((mapping_config.get('reports') or {}).get(source_category) or {}).get(source_mapping_key) or {}
        source_template_file = str(source_mapping_config.get('template_file') or '').strip()
        if not source_template_file:
            return jsonify({'success': False, 'error': '来源报表缺少模板文件配置'}), 400

        source_template_base_name = os.path.splitext(os.path.basename(source_template_file))[0]
        source_template_info = _analyze_template_file_payload(source_template_base_name)
        target_template_info = _analyze_template_file_payload(template_base_name)

        source_columns = source_template_info.get('columns') or []
        target_columns = target_template_info.get('columns') or []

        source_by_key = {
            _normalize_template_compare_key(item.get('name')): item
            for item in source_columns
            if _normalize_template_compare_key(item.get('name'))
        }
        target_by_key = {
            _normalize_template_compare_key(item.get('name')): item
            for item in target_columns
            if _normalize_template_compare_key(item.get('name'))
        }

        unchanged_keys = [key for key in source_by_key.keys() if key in target_by_key]
        added_keys = [key for key in target_by_key.keys() if key not in source_by_key]
        removed_keys = [key for key in source_by_key.keys() if key not in target_by_key]

        unchanged_columns = [copy.deepcopy(source_by_key[key]) for key in unchanged_keys]
        added_columns = [copy.deepcopy(target_by_key[key]) for key in added_keys]
        removed_columns = [copy.deepcopy(source_by_key[key]) for key in removed_keys]

        similar_columns = []
        for removed_key in removed_keys:
            removed_col = source_by_key[removed_key]
            removed_name = str(removed_col.get('name') or '').strip()
            best_match = None
            best_score = 0.0

            for added_key in added_keys:
                added_col = target_by_key[added_key]
                added_name = str(added_col.get('name') or '').strip()
                score = SequenceMatcher(None, removed_name, added_name).ratio()
                if removed_name and added_name and removed_name in added_name:
                    score = max(score, 0.9)
                if score > best_score:
                    best_score = score
                    best_match = added_col

            if best_match and best_score >= 0.5:
                similar_columns.append({
                    'source_name': removed_name,
                    'source_letter': removed_col.get('letter'),
                    'target_name': best_match.get('name'),
                    'target_letter': best_match.get('letter'),
                    'score': round(best_score, 3)
                })

        valid_target_keys = set(target_by_key.keys())
        invalid_mapping_rules = []
        invalid_conditional_rules = []
        for rule in source_mapping_config.get('mapping_rules') or []:
            target_name = str(rule.get('target_name') or '').strip()
            compare_key = _normalize_template_compare_key(target_name)
            if compare_key and compare_key not in valid_target_keys:
                invalid_item = {
                    'target_column': rule.get('target_column'),
                    'target_name': target_name,
                    'source_type': rule.get('source_type'),
                }
                invalid_mapping_rules.append(invalid_item)
                if _is_conditional_mapping_rule(rule):
                    invalid_conditional_rules.append(copy.deepcopy(invalid_item))

        return jsonify({
            'success': True,
            'data': {
                'source_template': source_template_info,
                'target_template': target_template_info,
                'unchanged_columns': unchanged_columns,
                'added_columns': added_columns,
                'removed_columns': removed_columns,
                'similar_columns': similar_columns,
                'invalid_mapping_rules': invalid_mapping_rules,
                'invalid_conditional_rules': invalid_conditional_rules,
            }
        })
    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except FileNotFoundError as e:
        return jsonify({'success': False, 'error': str(e)}), 404
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


# ==================== 业务系统映射字典相关API ====================

@common_bp.route('/api/get_assets_business_systems', methods=['GET'])
def get_assets_business_systems():
    """获取数据概览表中的所有业务系统名称"""
    try:
        conn = get_db_connection()
        cursor = conn.cursor()

        cursor.execute('SELECT DISTINCT "业务系统" FROM assets WHERE "业务系统" IS NOT NULL AND "业务系统" != "" ORDER BY "业务系统"')
        rows = cursor.fetchall()

        business_systems = [row[0] for row in rows]

        conn.close()

        return jsonify({
            'success': True,
            'business_systems': business_systems
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/get_merge_results_business_systems', methods=['GET'])
def get_merge_results_business_systems():
    """获取合并结果表中的所有业务系统名称（索引12）"""
    try:
        conn = sqlite3.connect(DatabaseConfig.MERGE_RESULTS_DB)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        # 检查表结构，判断是旧格式还是新格式
        cursor.execute('PRAGMA table_info(merge_results)')
        columns_info = cursor.fetchall()
        column_names = [col[1] for col in columns_info]
        use_json_format = 'json_data' in column_names and len(column_names) <= 3

        business_systems_set = set()

        if use_json_format:
            # 旧格式：使用 json_data 列
            cursor.execute('SELECT json_data FROM merge_results')
            rows = cursor.fetchall()
            for row in rows:
                try:
                    data = json.loads(row[0])
                    fields = list(data.values())
                    if len(fields) > 12 and fields[12]:
                        business_systems_set.add(fields[12])
                except:
                    pass
        else:
            # 新格式：直接使用列结构
            # 获取第12列（索引12）的列名（不包括id）
            data_columns = [col for col in column_names if col.lower() != 'id']
            if len(data_columns) > 12:
                column_name = data_columns[12]
                cursor.execute(f'SELECT DISTINCT "{column_name}" FROM merge_results WHERE "{column_name}" IS NOT NULL AND "{column_name}" != ""')
                rows = cursor.fetchall()
                for row in rows:
                    if row[0]:
                        business_systems_set.add(row[0])

        business_systems = sorted(list(business_systems_set))

        conn.close()

        return jsonify({
            'success': True,
            'business_systems': business_systems
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/merge/results', methods=['GET'])
def get_merge_results_paginated():
    """
    获取合并结果数据（分页）
    支持分页查询，避免一次性加载百万行数据导致性能问题
    """
    try:
        # 获取查询参数
        page = int(request.args.get('page', 1))
        pageSize = int(request.args.get('pageSize', 100))
        search = request.args.get('search', '').strip()
        category = request.args.get('category', '').strip()

        # 参数验证
        if page < 1:
            page = 1
        if pageSize < 1 or pageSize > 1000:
            pageSize = 100

        # 计算偏移量
        offset = (page - 1) * pageSize

        conn = sqlite3.connect(DatabaseConfig.MERGE_RESULTS_DB)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        # 检查表结构，判断是旧格式还是新格式
        cursor.execute('PRAGMA table_info(merge_results)')
        columns_info = cursor.fetchall()
        column_names = [col[1] for col in columns_info]
        use_json_format = 'json_data' in column_names and len(column_names) <= 3

        # 获取总记录数
        cursor.execute('SELECT COUNT(*) FROM merge_results')
        total = cursor.fetchone()[0]

        # 解析数据
        data = []

        if use_json_format:
            # 旧格式：使用 json_data 列
            sql = 'SELECT json_data FROM merge_results ORDER BY rowid LIMIT ? OFFSET ?'
            cursor.execute(sql, (pageSize, offset))
            rows = cursor.fetchall()
            for row in rows:
                try:
                    json_data = json.loads(row[0])
                    # 将JSON对象转换为数组格式，方便前端按索引访问
                    data.append(list(json_data.values()))
                except:
                    # 如果解析失败，返回空数组
                    data.append([])
        else:
            # 新格式：直接使用列结构
            # 获取所有列名（按顺序）
            columns_str = ', '.join([f'"{col}"' for col in column_names])
            query = f'SELECT {columns_str} FROM merge_results ORDER BY id LIMIT ? OFFSET ?'
            cursor.execute(query, (pageSize, offset))
            rows = cursor.fetchall()
            for row in rows:
                # 转换为字典，然后按列顺序提取值（不包括id）
                row_dict = dict(row)
                values = [row_dict.get(col) for col in column_names if col.lower() != 'id']
                data.append(values)

        conn.close()

        # 计算总页数
        total_pages = (total + pageSize - 1) // pageSize if total > 0 else 0

        return jsonify({
            'success': True,
            'data': data,
            'total': total,
            'page': page,
            'pageSize': pageSize,
            'totalPages': total_pages
        })

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


def save_business_system_name_mapping():
    """保存业务系统名称映射公共配置。"""
    try:
        data = request.get_json()
        mappings = _normalize_business_system_name_mapping_pairs((data or {}).get('mappings', {}))

        # 读取当前配置
        with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8') as f:
            config = json.load(f)

        # 更新映射字典
        if 'public_config' not in config:
            config['public_config'] = {}

        existing_config = config['public_config'].get(PUBLIC_CONFIG_BUSINESS_SYSTEM_NAME_MAPPING_KEY)
        config['public_config'][PUBLIC_CONFIG_BUSINESS_SYSTEM_NAME_MAPPING_KEY] = (
            _build_business_system_name_mapping_public_entry(mappings, existing_config)
        )

        # 保存配置
        with open(MAPPING_CONFIG_FILE, 'w', encoding='utf-8') as f:
            json.dump(config, f, ensure_ascii=False, indent=2)

        # 同步保存专用映射文件，确保导出链路读取到最新配置
        _save_business_system_name_mapping_file(mappings)

        return jsonify({
            'success': True,
            'message': '公共配置保存成功'
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500



# ==================== 大文件导入API ====================

@common_bp.route('/api/import/large', methods=['POST'])
def import_large_file():
    """大文件导入API（支持超大Excel文件）"""
    if 'file' not in request.files:
        return jsonify({'success': False, 'error': '未选择文件'}), 400

    file = request.files['file']
    data_type = request.form.get('type', 'device')  # 'device' 或 'merge'
    original_filename = file.filename
    file_ext = os.path.splitext(original_filename)[1].lower()

    if not file_ext:
        return jsonify({'success': False, 'error': '无法识别文件扩展名，请使用 .xlsx 文件'}), 400

    # LargeFileImporter 基于 openpyxl(read_only)，不支持 .xls
    if file_ext == '.xls':
        return jsonify({'success': False, 'error': '当前导入通道不支持 .xls，请先在 Excel 中另存为 .xlsx 后重试'}), 400

    if file_ext not in ['.xlsx', '.xlsm', '.xltx', '.xltm']:
        return jsonify({'success': False, 'error': '不支持的文件格式，请上传 .xlsx/.xlsm/.xltx/.xltm 文件'}), 400

    temp_saved_file_path = ''
    try:
        force_import = _is_force_import_requested()
        # 确定保存目录（DataFiles子目录）
        if data_type == 'device':
            save_dir = get_datafiles_source_dir('assets')
            columns_config_path = PathConfig.ASSETS_COLUMNS_JSON
        else:
            save_dir = get_datafiles_source_dir('merge_results')
            columns_config_path = PathConfig.MERGE_COLUMNS_JSON

        # 确保目录存在
        os.makedirs(save_dir, exist_ok=True)

        # 保留用户原始文件名（仅取basename避免路径注入）
        safe_original_filename = os.path.basename(original_filename or '')
        if not safe_original_filename:
            return jsonify({'success': False, 'error': '文件名无效'}), 400

        import_task_id = datetime.now().strftime('%Y%m%d%H%M%S%f')
        temp_saved_name = f'.tmp_{import_task_id}_{safe_original_filename}'
        temp_saved_file_path = os.path.join(save_dir, temp_saved_name)

        # 先保存临时文件，校验通过后再落到正式目录，避免无效文件残留
        file.save(temp_saved_file_path)
        columns = read_excel_headers(temp_saved_file_path)
        if not columns:
            _safe_remove_file(temp_saved_file_path)
            return jsonify({'success': False, 'error': '文件为空或无法读取表头'}), 400
        expected_headers = _get_workspace_template_headers_safe(data_type)
        missing_columns, extra_columns = _compare_import_headers(columns, expected_headers)
        warning_message = _build_import_schema_warning(data_type, missing_columns, extra_columns) if (missing_columns or extra_columns) else ''
        if missing_columns and not force_import:
            _safe_remove_file(temp_saved_file_path)
            return jsonify({
                'success': False,
                'requires_confirmation': True,
                'warning_type': 'missing_columns',
                'missing_columns': missing_columns,
                'extra_columns': extra_columns,
                'warning_message': warning_message
            })

        # 正式文件名加入时间戳，避免同名覆盖历史导入文件
        saved_filename = f'{import_task_id}_{safe_original_filename}'
        saved_file_path = os.path.join(save_dir, saved_filename)
        os.replace(temp_saved_file_path, saved_file_path)
        temp_saved_file_path = ''

        # 清除之前的进度
        clear_import_progress()

        # 根据类型选择数据库
        if data_type == 'device':
            db_path = DatabaseConfig.ASSETS_DB
            table_name = 'assets'
        else:
            db_path = DatabaseConfig.MERGE_RESULTS_DB
            table_name = 'merge_results'

        # 在后台线程中执行导入
        def import_in_background():
            importer = LargeFileImporter(db_path, table_name)
            try:
                result = importer.import_excel_file(
                    saved_file_path,
                    expected_headers=None,
                    allow_extra_columns=True
                )
                if result.get('success'):
                    sync_columns_config(columns_config_path, columns, saved_filename)
                # 导入成功，文件已保存到DataFiles目录，不需要删除
                print(f'[INFO] 导入完成，文件已保存至: {saved_file_path}')
            except Exception as e:
                print(f'[ERROR] 导入失败: {e}')
                # 导入失败时可以选择删除文件
                # if os.path.exists(saved_file_path):
                #     os.remove(saved_file_path)

        import threading
        thread = threading.Thread(target=import_in_background)
        thread.start()

        return jsonify({
            'success': True,
            'message': '已开始导入，请查看进度',
            'missing_columns': missing_columns,
            'extra_columns': extra_columns,
            'warning_message': warning_message
        })

    except Exception as e:
        _safe_remove_file(temp_saved_file_path)
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/import/progress', methods=['GET'])
def get_import_progress_api():
    """获取导入进度"""
    return jsonify(get_import_progress_func())


# ==================== 高级搜索路由 ====================

@common_bp.route('/api/assets/advanced-search', methods=['POST'])
def advanced_search_assets():
    """数据概览高级搜索API"""
    conn = None
    try:
        data = request.json
        conditions = data.get('conditions', [])
        page, page_size = normalize_advanced_search_pagination(
            data.get('page', 1),
            data.get('page_size', 20)
        )

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute('PRAGMA table_info(assets)')
        valid_db_columns = {col[1] for col in cursor.fetchall()}

        # 构建SQL查询
        where_clauses = []
        params = []

        for condition in conditions:
            field = condition.get('field')
            operator = validate_advanced_search_operator(condition.get('operator'))
            value = condition.get('value', '')
            db_field = resolve_advanced_search_field(field, valid_db_columns)

            if operator == 'contains':
                where_clauses.append(f'"{db_field}" LIKE ?')
                params.append(f'%{value}%')
            elif operator == 'equals':
                where_clauses.append(f'"{db_field}" = ?')
                params.append(value)
            elif operator == 'not_contains':
                where_clauses.append(f'"{db_field}" NOT LIKE ?')
                params.append(f'%{value}%')
            elif operator == 'not_equals':
                where_clauses.append(f'"{db_field}" != ?')
                params.append(value)
            elif operator == 'is_empty':
                where_clauses.append(f'("{db_field}" IS NULL OR "{db_field}" = "")')
            elif operator == 'not_empty':
                where_clauses.append(f'("{db_field}" IS NOT NULL AND "{db_field}" != "")')

        # 构建完整SQL
        base_sql = 'SELECT * FROM assets'
        count_sql = 'SELECT COUNT(*) as total FROM assets'

        if where_clauses:
            where_sql = ' AND '.join(where_clauses)
            base_sql += f' WHERE {where_sql}'
            count_sql += f' WHERE {where_sql}'

        # 获取总数
        cursor.execute(count_sql, params)
        total = cursor.fetchone()['total']

        # 获取分页数据
        offset = (page - 1) * page_size
        base_sql += ' ORDER BY id DESC LIMIT ? OFFSET ?'
        cursor.execute(base_sql, params + [page_size, offset])
        rows = cursor.fetchall()

        # 转换为字典列表
        data_list = [dict(row) for row in rows]

        return jsonify({
            'success': True,
            'data': data_list,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if conn:
            conn.close()


@common_bp.route('/api/merge/advanced-search', methods=['POST'])
def advanced_search_merge():
    """合并结果高级搜索API"""
    conn = None
    try:
        data = request.json
        conditions = data.get('conditions', [])
        category = data.get('category', '')
        page, page_size = normalize_advanced_search_pagination(
            data.get('page', 1),
            data.get('page_size', 20)
        )

        conn = sqlite3.connect(DatabaseConfig.MERGE_RESULTS_DB)
        conn.row_factory = sqlite3.Row
        cursor = conn.cursor()

        # 检查表结构，判断是旧格式还是新格式
        cursor.execute('PRAGMA table_info(merge_results)')
        columns_info = cursor.fetchall()
        column_names = [col[1] for col in columns_info]
        valid_db_columns = set(column_names) | set(ADVANCED_SEARCH_FIELD_MAPPING.values())
        use_json_format = 'json_data' in column_names and len(column_names) <= 3

        # 构建SQL查询
        where_clauses = []
        params = []

        # 添加分类筛选
        if category and category != 'all':
            if use_json_format:
                where_clauses.append('category = ?')
                params.append(category)

        for condition in conditions:
            field = condition.get('field')
            operator = validate_advanced_search_operator(condition.get('operator'))
            value = condition.get('value', '')
            db_field = resolve_advanced_search_field(field, valid_db_columns)

            if use_json_format:
                # 旧格式：使用 json_data 语法
                if operator == 'contains':
                    where_clauses.append(f'json_data->>"$.{db_field}" LIKE ?')
                    params.append(f'%{value}%')
                elif operator == 'equals':
                    where_clauses.append(f'json_data->>"$.{db_field}" = ?')
                    params.append(value)
                elif operator == 'not_contains':
                    where_clauses.append(f'json_data->>"$.{db_field}" NOT LIKE ?')
                    params.append(f'%{value}%')
                elif operator == 'not_equals':
                    where_clauses.append(f'json_data->>"$.{db_field}" != ?')
                    params.append(value)
                elif operator == 'is_empty':
                    where_clauses.append(f'(json_data->>"$.{db_field}" IS NULL OR json_data->>"$.{db_field}" = "")')
                elif operator == 'not_empty':
                    where_clauses.append(f'(json_data->>"$.{db_field}" IS NOT NULL AND json_data->>"$.{db_field}" != "")')
            else:
                # 新格式：直接使用列名
                if operator == 'contains':
                    where_clauses.append(f'"{db_field}" LIKE ?')
                    params.append(f'%{value}%')
                elif operator == 'equals':
                    where_clauses.append(f'"{db_field}" = ?')
                    params.append(value)
                elif operator == 'not_contains':
                    where_clauses.append(f'"{db_field}" NOT LIKE ?')
                    params.append(f'%{value}%')
                elif operator == 'not_equals':
                    where_clauses.append(f'"{db_field}" != ?')
                    params.append(value)
                elif operator == 'is_empty':
                    where_clauses.append(f'("{db_field}" IS NULL OR "{db_field}" = "")')
                elif operator == 'not_empty':
                    where_clauses.append(f'("{db_field}" IS NOT NULL AND "{db_field}" != "")')

        # 构建完整SQL
        base_sql = 'SELECT * FROM merge_results'
        count_sql = 'SELECT COUNT(*) as total FROM merge_results'

        if where_clauses:
            where_sql = ' AND '.join(where_clauses)
            base_sql += f' WHERE {where_sql}'
            count_sql += f' WHERE {where_sql}'

        # 获取总数
        cursor.execute(count_sql, params)
        total = cursor.fetchone()['total']

        # 获取分页数据
        offset = (page - 1) * page_size
        base_sql += ' ORDER BY id DESC LIMIT ? OFFSET ?'
        cursor.execute(base_sql, params + [page_size, offset])
        rows = cursor.fetchall()

        # 转换为字典列表
        data_list = []
        for row in rows:
            row_dict = dict(row)
            if use_json_format:
                # 旧格式：解析json_data
                if row_dict.get('json_data'):
                    try:
                        json_data = json.loads(row_dict['json_data'])
                        row_dict.update(json_data)
                    except:
                        pass
            # 新格式：直接使用列数据，无需额外处理
            data_list.append(row_dict)

        return jsonify({
            'success': True,
            'data': data_list,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    except ValueError as e:
        return jsonify({'success': False, 'error': str(e)}), 400
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({'success': False, 'error': str(e)}), 500
    finally:
        if conn:
            conn.close()


@common_bp.route('/api/assets/columns', methods=['GET'])
def get_assets_columns():
    """获取数据概览可搜索字段"""
    try:
        columns = [
            {'field': 'business_system', 'name': '业务系统'},
            {'field': 'data_asset_name', 'name': '数据资产名称'},
            {'field': 'ip_address', 'name': 'IP地址'},
            {'field': 'data_type', 'name': '数据类型'},
            {'field': 'data_level', 'name': '数据分级'},
            {'field': 'status', 'name': '状态'},
            {'field': 'responsible_person', 'name': '负责人'}
        ]
        return jsonify({'success': True, 'columns': columns})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500


@common_bp.route('/api/merge/search-columns', methods=['GET'])
def get_merge_search_columns():
    """获取合并结果可搜索字段"""
    try:
        columns = [
            {'field': 'business_system', 'name': '业务系统'},
            {'field': 'data_asset_name', 'name': '数据资产名称'},
            {'field': 'ip_address', 'name': 'IP地址'},
            {'field': 'data_type', 'name': '数据类型'},
            {'field': 'data_level', 'name': '数据分级'},
            {'field': 'status', 'name': '状态'},
            {'field': 'responsible_person', 'name': '负责人'}
        ]
        return jsonify({'success': True, 'columns': columns})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500
