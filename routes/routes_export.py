#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
导出功能路由模块
处理各类数据导出请求
"""
from flask import Blueprint, request, jsonify
import sqlite3
import json
import re
import os
import threading
from datetime import datetime
from urllib.parse import urlparse
from openpyxl import load_workbook, Workbook

from core.utils import get_db_connection
from core.data_mapper import get_mapper
from core.sensitive_data_handler import get_sensitive_handler
from core.export_engine import get_export_engine
from config.config import Config, PathConfig, DatabaseConfig, ExportConfig

# 创建Blueprint
export_bp = Blueprint('export', __name__)

def _new_export_progress(task_id='legacy', status='idle'):
    return {
        'task_id': task_id,
        'status': status,  # idle, running, canceling, completed, error, canceled
        'current': 0,
        'total': 0,
        'message': '',
        'filename': '',
        'total_rows': 0,
        'output_path': '',
        'cancel_requested': False
    }


# 导出进度追踪
export_progress = _new_export_progress()
export_progress_tasks = {}
export_progress_lock = threading.RLock()
export_task_threads = {}
MAPPING_CONFIG_FILE = os.path.join(Config.BASE_DIR, 'config', 'mapping_config.json')


class ExportCanceledError(Exception):
    """导出任务被用户取消。"""


def get_log_operation():
    """延迟导入日志记录函数"""
    from routes.routes_logs import log_operation
    return log_operation


def _get_export_task_id():
    task_id = (request.args.get('task_id') or '').strip()
    return task_id or 'legacy'


def _should_skip_export_log(operation_desc=''):
    skip_log = (request.args.get('skip_log') or '').strip().lower()
    if skip_log in {'1', 'true', 'yes', 'on'}:
        return True

    return operation_desc in {
        '导出合并结果数据全量',
        '导出数据概览数据全量',
    }


def _ensure_export_progress(task_id=None):
    global export_progress

    resolved_task_id = task_id or 'legacy'
    with export_progress_lock:
        progress = export_progress_tasks.get(resolved_task_id)
        if progress is None:
            progress = _new_export_progress(resolved_task_id)
            export_progress_tasks[resolved_task_id] = progress
        export_progress = progress
        return progress


def _snapshot_export_progress(task_id=None):
    progress = _ensure_export_progress(task_id)
    with export_progress_lock:
        return dict(progress)


def reject_cross_site_delete_request():
    """拦截明显来自外部站点的删除请求。"""
    origin = str(request.headers.get('Origin') or '').strip()
    sec_fetch_site = str(request.headers.get('Sec-Fetch-Site') or '').strip().lower()

    if sec_fetch_site == 'cross-site':
        return jsonify({'error': '跨站删除请求已被拒绝'}), 403

    if not origin:
        return None

    origin_netloc = urlparse(origin).netloc.lower()
    request_netloc = urlparse(request.host_url).netloc.lower()
    if origin_netloc and origin_netloc != request_netloc:
        return jsonify({'error': '跨站删除请求已被拒绝'}), 403

    return None


def _start_export_progress(task_id):
    global export_progress

    with export_progress_lock:
        progress = _new_export_progress(task_id, status='running')
        progress['message'] = '正在初始化...'
        export_progress_tasks[task_id] = progress
        export_progress = progress
        return progress


def _set_export_progress(task_id, **updates):
    global export_progress

    with export_progress_lock:
        progress = _ensure_export_progress(task_id)
        progress.update(updates)
        export_progress = progress
        return progress


def _delete_export_output(output_path):
    normalized_path = str(output_path or '').strip()
    if not normalized_path or not os.path.exists(normalized_path):
        return False

    try:
        os.remove(normalized_path)
        return True
    except FileNotFoundError:
        return False
    except Exception as remove_err:
        print(f'[WARNING] 删除导出半成品文件失败: {remove_err}')
        return False


def _mark_export_canceled(task_id, message='导出已取消', output_path=''):
    removed = _delete_export_output(output_path)
    final_message = message
    if removed:
        final_message = '导出已取消，已删除半成品文件'
    elif output_path:
        final_message = '导出已取消，半成品文件将在任务结束后清理'

    progress = _set_export_progress(
        task_id,
        status='canceled',
        message=final_message,
        filename='',
        output_path='',
        cancel_requested=True
    )
    return dict(progress)


def _request_export_cancel(task_id):
    with export_progress_lock:
        progress = _ensure_export_progress(task_id)
        current_status = progress.get('status') or 'idle'
        if current_status in {'completed', 'error', 'canceled'}:
            return dict(progress)

        progress['cancel_requested'] = True
        progress['status'] = 'canceling'
        progress['message'] = '正在取消导出，请稍候...'
        return dict(progress)


def _build_progress_callback(task_id):
    def progress_callback(key, value):
        current_progress = _snapshot_export_progress(task_id)
        if current_progress.get('cancel_requested'):
            raise ExportCanceledError('导出已取消')
        _set_export_progress(task_id, **{key: value})
        current_progress = _snapshot_export_progress(task_id)
        if current_progress.get('cancel_requested'):
            raise ExportCanceledError('导出已取消')

    return progress_callback


def _execute_config_export(task_id, category, report_key, operation_desc, limit, skip_export_log):
    try:
        engine = get_export_engine()
        result = engine.export_by_config(
            category,
            report_key,
            _build_progress_callback(task_id),
            limit=limit
        )

        result = dict(result)
        result['task_id'] = task_id

        current_progress = _snapshot_export_progress(task_id)
        if current_progress.get('cancel_requested'):
            _mark_export_canceled(
                task_id,
                output_path=result.get('output_path', '') or current_progress.get('output_path', '')
            )
            return result, 200

        if result.get('success'):
            _set_export_progress(
                task_id,
                status='completed',
                message=result.get('message', ''),
                filename=result.get('filename', ''),
                total_rows=result.get('total_rows', 0),
                output_path=result.get('output_path', '')
            )

            if limit == 0 and not skip_export_log:
                try:
                    log_operation = get_log_operation()
                    log_operation(
                        operation_type='data_export',
                        page_type='reporting',
                        operation_desc=operation_desc,
                        file_name=result.get('filename', ''),
                        record_count=result.get('total_rows', 0)
                    )
                except Exception as log_err:
                    print(f'[日志记录失败] {log_err}')
        else:
            _set_export_progress(
                task_id,
                status='error',
                message=f'导出失败: {result.get("error", "未知错误")}'
            )

        return result, 200

    except ExportCanceledError:
        progress = _snapshot_export_progress(task_id)
        _mark_export_canceled(task_id, output_path=progress.get('output_path', ''))
        return {
            'success': False,
            'canceled': True,
            'message': '导出已取消',
            'task_id': task_id
        }, 200

    except Exception as e:
        print(f'[ERROR] 导出失败: {str(e)}')
        import traceback
        traceback.print_exc()

        _set_export_progress(
            task_id,
            status='error',
            message=f'导出失败: {str(e)}'
        )

        return {
            'success': False,
            'error': str(e),
            'task_id': task_id
        }, 500

    finally:
        with export_progress_lock:
            export_task_threads.pop(task_id, None)


def _start_async_export_task(task_id, category, report_key, operation_desc, limit, skip_export_log):
    with export_progress_lock:
        worker = export_task_threads.get(task_id)
        if worker and worker.is_alive():
            return False

        worker = threading.Thread(
            target=_execute_config_export,
            args=(task_id, category, report_key, operation_desc, limit, skip_export_log),
            name=f'export-task-{task_id}',
            daemon=True
        )
        export_task_threads[task_id] = worker

    worker.start()
    return True


def _run_config_export(category, report_key, operation_desc):
    task_id = _get_export_task_id()
    limit = request.args.get('limit', 0, type=int)
    skip_export_log = _should_skip_export_log(operation_desc)

    _start_export_progress(task_id)

    if task_id != 'legacy':
        started = _start_async_export_task(
            task_id,
            category,
            report_key,
            operation_desc,
            limit,
            skip_export_log
        )
        if not started:
            return jsonify({
                'success': False,
                'error': '导出任务已在执行中',
                'task_id': task_id
            }), 409

        return jsonify({
            'success': True,
            'started': True,
            'task_id': task_id,
            'message': '导出任务已启动'
        })

    result, status_code = _execute_config_export(
        task_id,
        category,
        report_key,
        operation_desc,
        limit,
        skip_export_log
    )
    return jsonify(result), status_code


def _candidate_export_report_keys(category, report_key):
    rc = str(report_key or '').strip()
    if not rc:
        return []

    rc_short = rc.replace('i_10600_', '', 1) if rc.startswith('i_10600_') else rc

    if category == 'yeji':
        preferred_keys = [f'i_10600_{rc_short}', rc]
    elif category == 'smc':
        zh_to_en = {'附件三': 'attachment_3', '附件五': 'attachment_5'}
        en_to_zh = {'attachment_3': '附件三', 'attachment_5': '附件五'}
        canonical = zh_to_en.get(rc, rc)
        preferred_keys = [canonical, rc, en_to_zh.get(canonical, canonical)]
    elif category == 'xinan':
        preferred_keys = [rc_short, rc, f'i_10600_{rc_short}']
    else:
        preferred_keys = [rc]

    common_keys = [f'i_10600_{rc_short}', f'i_{rc_short}', f'smc_{rc_short}', rc]
    return list(dict.fromkeys([key for key in preferred_keys + common_keys if key]))


def _resolve_export_report_key(category, report_key):
    rc = str(report_key or '').strip()
    if not rc or not os.path.exists(MAPPING_CONFIG_FILE):
        return rc

    try:
        with open(MAPPING_CONFIG_FILE, 'r', encoding='utf-8') as f:
            mapping_config = json.load(f)
    except Exception:
        return rc

    category_reports = (mapping_config.get('reports') or {}).get(category) or {}
    if not category_reports:
        return rc

    rc_short = rc.replace('i_10600_', '', 1) if rc.startswith('i_10600_') else rc

    for key in _candidate_export_report_keys(category, rc):
        if key in category_reports:
            return key

    for key in category_reports:
        if key.endswith(f'_{rc_short}') or key == rc:
            return key

    return rc


@export_bp.route('/api/export/yeji/i_10600_10001', methods=['GET'])
def export_yeji_i_10600_10001():
    """导出业支上报-i_10600_10001上报表（数据库资产信息表）- 使用通用导出引擎

    支持参数:
        limit: 限制导出行数（用于测试），0表示全部
    """
    return _run_config_export('yeji', 'i_10600_10001', '导出业支上报-i_10600_10001')


@export_bp.route('/api/export/yeji/i_10600_10002', methods=['GET'])
def export_yeji_i_10600_10002():
    """导出业支上报-i_10600_10002上报表（数据库资产字段信息表）- 使用通用导出引擎

    支持参数:
        limit: 限制导出行数（用于测试），0表示全部
    """
    return _run_config_export('yeji', 'i_10600_10002', '导出业支上报-i_10600_10002')


@export_bp.route('/api/export/reports/<category>/<code>', methods=['GET'])
def export_report_by_code(category, code):
    """按类别和代码动态导出报表。"""
    actual_report_key = _resolve_export_report_key(category, code)
    return _run_config_export(category, actual_report_key, f'导出报表 {category}/{actual_report_key}')


@export_bp.route('/api/export/progress', methods=['GET'])
def get_export_progress():
    """获取导出进度"""
    task_id = (request.args.get('task_id') or '').strip()
    if task_id:
        return jsonify(_snapshot_export_progress(task_id))
    with export_progress_lock:
        return jsonify(dict(export_progress))


@export_bp.route('/api/export/cancel', methods=['POST'])
def cancel_export():
    """取消导出任务"""
    reject_result = reject_cross_site_delete_request()
    if reject_result:
        return reject_result

    payload = request.get_json(silent=True) or {}
    task_id = str(payload.get('task_id') or request.args.get('task_id') or '').strip()
    if not task_id:
        return jsonify({
            'success': False,
            'error': '缺少 task_id'
        }), 400

    progress = _snapshot_export_progress(task_id)
    if progress.get('status') == 'completed':
        return jsonify({
            'success': False,
            'error': '导出任务已完成，无法取消',
            'task_id': task_id,
            'status': 'completed'
        }), 409

    if progress.get('status') in {'error', 'canceled'}:
        return jsonify({
            'success': True,
            'task_id': task_id,
            'status': progress.get('status'),
            'message': progress.get('message') or '任务已结束'
        })

    progress = _request_export_cancel(task_id)

    with export_progress_lock:
        worker = export_task_threads.get(task_id)
        worker_alive = bool(worker and worker.is_alive())

    if not worker_alive:
        progress = _mark_export_canceled(task_id, output_path=progress.get('output_path', ''))

    return jsonify({
        'success': True,
        'task_id': task_id,
        'status': progress.get('status'),
        'message': progress.get('message') or '导出取消请求已提交'
    })


@export_bp.route('/api/export/merge-results/all', methods=['GET'])
def export_merge_results_all():
    """导出所有合并结果数据到Excel"""
    try:
        print('[START] 开始导出合并结果数据...')

        # 步骤1: 读取合并结果数据
        print('[STEP 1] 正在读取合并结果表数据...')
        merge_conn = get_db_connection(DatabaseConfig.MERGE_RESULTS_DB)
        try:
            merge_cursor = merge_conn.cursor()

            # 兼容两种格式：
            # 1) 旧格式：id + json_data
            # 2) 新格式：大文件导入后的列式结构
            merge_cursor.execute('PRAGMA table_info(merge_results)')
            columns_info = merge_cursor.fetchall()
            column_names = []
            for col in columns_info:
                if isinstance(col, dict):
                    column_names.append(col.get('name'))
                elif hasattr(col, 'keys'):
                    column_names.append(col['name'])
                else:
                    column_names.append(col[1])

            use_json_format = 'json_data' in column_names and len(column_names) <= 3

            if use_json_format:
                merge_cursor.execute('SELECT id, json_data FROM merge_results ORDER BY id')
            else:
                columns_str = ', '.join([f'"{col}"' for col in column_names])
                merge_cursor.execute(f'SELECT {columns_str} FROM merge_results ORDER BY id')

            records = merge_cursor.fetchall()
        finally:
            merge_conn.close()

        print(f'[INFO] 读取到 {len(records)} 条记录')

        if len(records) == 0:
            return jsonify({
                'success': False,
                'error': '没有可导出的数据'
            })

        # 步骤2: 解析JSON数据并转换
        print('[STEP 2] 正在解析JSON数据...')
        output_rows = []
        parse_failed_count = 0

        def pick_value(record_data, *keys, default=''):
            """从多候选字段中取第一个非空值"""
            for key in keys:
                value = record_data.get(key)
                if value is not None and value != '':
                    return value
            return default

        for idx, record in enumerate(records, 1):
            try:
                if use_json_format:
                    if isinstance(record, dict):
                        json_text = record.get('json_data', '')
                    elif hasattr(record, 'keys'):
                        json_text = record['json_data']
                    else:
                        json_text = record[1]
                    data = json.loads(json_text) if json_text else {}
                else:
                    data = dict(record) if hasattr(record, 'keys') else {}

                # 提取关键字段
                row = [
                    idx,  # 序号
                    pick_value(data, '业务系统名称', '业务系统'),  # 业务系统名称
                    pick_value(data, '主机IP', 'IP地址', '数据源IP'),  # 主机IP
                    pick_value(data, '数据类型'),  # 数据类型
                    pick_value(data, '实例名/数据库名', '实例名', '数据库名', '实例名/数据库名(schema)'),  # 实例名/数据库名
                    pick_value(data, '表名'),  # 表名
                    pick_value(data, '字段名称', '字段名'),  # 字段名称
                    pick_value(data, '字段数据分类', '数据分类', 'AI字段分类'),  # 字段数据分类
                    pick_value(data, '字段数据分级', '数据分级'),  # 字段数据分级
                    pick_value(data, '数据名称', '字段名称', '字段名'),  # 数据名称
                    pick_value(data, '数据样例', '样本'),  # 数据样例
                    pick_value(data, '数据条数', '记录数', '表记录数'),  # 数据条数
                    pick_value(data, '数据大小', '表空间大小', '数据总量'),  # 数据大小
                    pick_value(data, '保存期限', default='永久保存'),  # 保存期限
                    pick_value(data, '数据存储状态'),  # 数据存储状态
                    pick_value(data, '对外提供情况', '数据对外提供情况', default='不涉及对外提供'),  # 对外提供情况
                    pick_value(data, '处理方式', default='数据收集|数据传输|数据存储|数据使用加工|数据销毁'),  # 处理方式
                    pick_value(data, '处理目的', '数据处理目的'),  # 处理目的
                    pick_value(data, '流转路径'),  # 流转路径
                    pick_value(data, '业务场景'),  # 业务场景
                    pick_value(data, '保障措施', '提供数据生命周期各环节安全措施配套情况'),  # 保障措施
                    pick_value(data, '业务系统负责人', '责任人'),  # 责任人
                    pick_value(data, '业务系统负责人联系方式', '联系电话'),  # 联系方式
                    pick_value(data, '数据来源', default='生产运营中产生'),  # 数据来源
                    pick_value(data, '补充信息'),  # 补充信息
                    pick_value(data, '存储位置', '数据存储位置', default='物理机房')  # 存储位置
                ]
                output_rows.append(row)
            except Exception as e:
                parse_failed_count += 1
                print(f'[WARNING] 第{idx}条记录解析失败: {e}')
                continue

        print(f'[INFO] 成功解析 {len(output_rows)} 条记录')
        if parse_failed_count:
            print(f'[WARNING] 解析失败记录数: {parse_failed_count}')

        if not output_rows:
            return jsonify({
                'success': False,
                'error': '合并结果数据解析后未生成有效导出记录'
            })

        # 步骤3: 创建输出文件
        print('[STEP 3] 正在创建输出文件...')
        from datetime import datetime

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'合并结果数据全量_{timestamp}.xlsx'
        output_dir = PathConfig.EXPORT_FOLDER
        os.makedirs(output_dir, exist_ok=True)
        output_path = os.path.join(output_dir, filename)

        # 创建新的Excel文件
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "合并结果数据"

        # 写入表头
        headers = ['序号', '业务系统名称', '主机IP', '数据类型', '实例名/数据库名', '表名', '字段名称',
                    '字段数据分类', '字段数据分级', '数据名称', '数据样例', '数据条数', '数据大小',
                    '保存期限', '数据存储状态', '对外提供情况', '处理方式', '处理目的', '流转路径',
                    '业务场景', '保障措施', '责任人', '联系电话', '数据来源', '补充信息', '存储位置']
        ws.append(headers)

        # 写入数据（从第2行开始）
        for row in output_rows:
            ws.append(row)

        # 保存文件
        wb.save(output_path)

        print(f'[SUCCESS] 导出成功！')
        print(f'文件名: {filename}')
        print(f'保存路径: {output_path}')
        print(f'总行数: {len(output_rows)}')

        # 记录操作日志
        if not _should_skip_export_log('导出合并结果数据全量'):
            try:
                log_operation = get_log_operation()
                log_operation(
                    operation_type='data_export',
                    page_type='merge',
                    operation_desc='导出合并结果数据全量',
                    file_name=filename,
                    record_count=len(output_rows)
                )
            except Exception as log_err:
                print(f'[日志记录失败] {log_err}')

        return jsonify({
            'success': True,
            'filename': filename,
            'total_rows': len(output_rows),
            'output_path': output_path,
            'message': f'成功导出{len(output_rows)}条记录到{filename}'
        })

    except Exception as e:
        print(f'[ERROR] 导出失败: {str(e)}')
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@export_bp.route('/api/export/assets/all', methods=['GET'])
def export_assets_all():
    """导出所有数据概览数据到Excel"""
    try:
        print('[START] 开始导出数据概览数据...')

        # 步骤1: 读取assets表数据
        print('[STEP 1] 正在读取数据概览表数据...')
        conn = get_db_connection(DatabaseConfig.ASSETS_DB)
        cursor = conn.cursor()
        cursor.execute('SELECT * FROM assets')
        assets = cursor.fetchall()

        # 获取列名
        columns = [desc[0] for desc in cursor.description]
        print(f'读取到 {len(assets)} 条记录')
        print(f'列名: {columns}')

        conn.close()

        if len(assets) == 0:
            return jsonify({
                'success': False,
                'error': '没有可导出的数据'
            })

        # 步骤2: 创建输出文件
        print('[STEP 2] 正在创建输出文件...')
        from datetime import datetime

        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'数据概览全量_{timestamp}.xlsx'
        output_dir = PathConfig.EXPORT_FOLDER
        os.makedirs(output_dir, exist_ok=True)
        output_path = os.path.join(output_dir, filename)

        # 创建新的Excel文件
        wb = Workbook()
        ws = wb.active
        ws.title = "数据概览"

        # 写入表头
        ws.append(columns)

        # 写入数据
        for asset in assets:
            ws.append(list(asset))

        # 保存文件
        wb.save(output_path)

        print(f'[SUCCESS] 导出成功！')
        print(f'文件名: {filename}')
        print(f'保存路径: {output_path}')
        print(f'总行数: {len(assets)}')

        # 记录操作日志
        if not _should_skip_export_log('导出数据概览数据全量'):
            try:
                log_operation = get_log_operation()
                log_operation(
                    operation_type='data_export',
                    page_type='device',
                    operation_desc='导出数据概览数据全量',
                    file_name=filename,
                    record_count=len(assets)
                )
            except Exception as log_err:
                print(f'[日志记录失败] {log_err}')

        return jsonify({
            'success': True,
            'filename': filename,
            'total_rows': len(assets),
            'output_path': output_path,
            'message': f'成功导出{len(assets)}条记录到{filename}'
        })

    except Exception as e:
        print(f'[ERROR] 导出失败: {str(e)}')
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@export_bp.route('/api/export/reports/<category>/<code>', methods=['DELETE'])
def delete_report(category, code):
    """删除报表关联导出文件和映射配置"""
    try:
        blocked_response = reject_cross_site_delete_request()
        if blocked_response:
            return blocked_response

        print(f'[DELETE] 开始删除报表 {category}/{code}')

        normalized_code = str(code).strip()
        if category == 'smc':
            normalized_code = {'附件三': 'attachment_3', '附件五': 'attachment_5'}.get(
                normalized_code, normalized_code
            )
        elif category in ('yeji', 'xinan') and normalized_code.startswith('i_10600_'):
            normalized_code = normalized_code.replace('i_10600_', '', 1)

        export_dir = {
            'yeji': PathConfig.YEZHI_EXPORT_DIR,
            'smc': PathConfig.SMC_EXPORT_DIR,
            'xinan': PathConfig.XINAN_EXPORT_DIR
        }.get(category, PathConfig.YEZHI_EXPORT_DIR)

        deleted_files = []
        report_prefix = f'i_10600_{normalized_code}'
        if os.path.exists(export_dir):
            for filename in os.listdir(export_dir):
                if not filename.lower().endswith(('.xlsx', '.xls')):
                    continue

                matched = False
                if category in ('yeji', 'xinan'):
                    matched = filename.startswith(f'{report_prefix}_')
                elif category == 'smc':
                    keys = {
                        'attachment_3': ['附件三', '附件3', 'attachment_3'],
                        'attachment_5': ['附件五', '附件5', 'attachment_5']
                    }.get(normalized_code, [normalized_code])
                    matched = any(k in filename for k in keys)

                if matched:
                    try:
                        os.remove(os.path.join(export_dir, filename))
                        deleted_files.append(filename)
                    except Exception as file_err:
                        print(f'[WARNING] 删除导出文件失败 {filename}: {file_err}')

        mapping_config_file = os.path.join(
            os.path.dirname(os.path.dirname(__file__)),
            'config',
            'mapping_config.json'
        )
        mapping_deleted = False
        if os.path.exists(mapping_config_file):
            try:
                with open(mapping_config_file, 'r', encoding='utf-8') as f:
                    config = json.load(f)

                if category in config.get('reports', {}):
                    category_reports = config['reports'][category]
                    candidate_keys = [code, normalized_code]
                    if category in ('yeji', 'xinan'):
                        candidate_keys.append(f'i_10600_{normalized_code}')
                    elif category == 'smc':
                        candidate_keys.append({'attachment_3': '附件三', 'attachment_5': '附件五'}.get(
                            normalized_code, normalized_code
                        ))

                    for key in list(dict.fromkeys(candidate_keys)):
                        if key in category_reports:
                            del category_reports[key]
                            mapping_deleted = True

                    with open(mapping_config_file, 'w', encoding='utf-8') as f:
                        json.dump(config, f, ensure_ascii=False, indent=2)
            except Exception as map_err:
                print(f'[WARNING] 删除映射配置失败: {map_err}')

        try:
            log_operation = get_log_operation()
            log_operation(
                operation_type='data_delete',
                page_type='reporting',
                operation_desc=f'删除报表: {category}/{code}',
                file_name=', '.join(deleted_files) if deleted_files else '无导出文件',
                details=f'映射配置已删除: {mapping_deleted}'
            )
        except Exception as log_err:
            print(f'[WARNING] 记录日志失败: {log_err}')

        messages = []
        messages.append(f'已删除{len(deleted_files)}个导出文件' if deleted_files else '未找到导出文件')
        messages.append('已删除映射配置' if mapping_deleted else '未找到映射配置')
        result_message = '; '.join(messages)

        return jsonify({
            'success': True,
            'message': result_message,
            'deleted_files': deleted_files,
            'mapping_deleted': mapping_deleted
        })
    except Exception as e:
        print(f'[ERROR] 删除报表失败: {str(e)}')
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'message': f'删除失败: {str(e)}'
        }), 500

@export_bp.route('/api/templates/list', methods=['GET'])
def list_template_files():
    """获取所有模板文件列表"""
    try:
        import os
        from config.config import PathConfig
        
        category = request.args.get('category', "")  # 'yeji', 'smc', 'xinan' or empty (all)

        # 类别映射
        category_dirs = {
            'yeji': os.path.join(PathConfig.TEMPLATES_DIR, '业支上报'),
            'smc': os.path.join(PathConfig.TEMPLATES_DIR, 'SMC上报'),
            'xinan': os.path.join(PathConfig.TEMPLATES_DIR, '信安上报')
        }

        templates = []

        # 如果指定了类别，只返回该类别的模板
        if category and category in category_dirs:
            dirs_to_scan = [category_dirs[category]]
        else:
            dirs_to_scan = list(category_dirs.values())

        for dir_path in dirs_to_scan:
            if not os.path.exists(dir_path):
                continue

            for filename in os.listdir(dir_path):
                if filename.endswith('.xlsx') and not filename.startswith('~'):
                    file_path = os.path.join(dir_path, filename)

                    # 提取报表代码（去除_template.xlsx后缀）
                    base_name = filename.replace('_template.xlsx', '').replace('.xlsx', '')

                    # 确定类别
                    if '业支上报' in dir_path:
                        file_category = 'yeji'
                    elif 'SMC上报' in dir_path:
                        file_category = 'smc'
                    elif '信安上报' in dir_path:
                        file_category = 'xinan'
                    else:
                        file_category = 'yeji'  # 默认

                    templates.append({
                        'filename': filename,
                        'base_name': base_name,
                        'category': file_category,
                        'size': os.path.getsize(file_path),
                        'full_path': file_path
                    })

        # 按类别和文件名排序
        templates.sort(key=lambda x: (x['category'], x['filename']))

        return jsonify({
            'success': True,
            'data': templates
        })
    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500


@export_bp.route('/api/export/xinan/i_10600_00000', methods=['GET'])
def export_xinan_i_10600_00000():
    """导出信安上报-i_10600_00000上报表（数据资产汇总表）- 使用通用导出引擎

    支持参数:
        limit: 限制导出行数（用于测试），0表示全部
    """
    return _run_config_export('xinan', '00000', '导出信安上报-i_10600_00000')


@export_bp.route('/api/export/yeji/i_10600_10004', methods=['GET'])
def export_yeji_i_10600_10004():
    """导出业支上报-i_10600_10004上报表（数据策略上报清单）- 使用通用导出引擎

    支持参数:
        limit: 限制导出行数（用于测试），0表示全部
    """
    return _run_config_export('yeji', 'i_10600_10004', '导出业支上报-i_10600_10004')


@export_bp.route('/api/export/xinan/i_10600_10001', methods=['GET'])
def export_xinan_i_10600_10001():
    """导出信安上报-i_10600_10001上报表（数据资产字段信息表）- 使用通用导出引擎

    支持参数:
        limit: 限制导出行数（用于测试），0表示全部
    """
    return _run_config_export('xinan', '10001', '导出信安上报-i_10600_10001')


@export_bp.route('/api/export/smc/attachment_3', methods=['GET'])
def export_smc_attachment_3():
    """导出SMC上报-附件三上报表（数据资产清单）- 使用通用导出引擎

    支持参数:
        limit: 限制导出行数（用于测试），0表示全部
    """
    return _run_config_export('smc', 'attachment_3', '导出SMC上报-附件三')


@export_bp.route('/api/export/smc/attachment_5', methods=['GET'])
def export_smc_attachment_5():
    """导出SMC上报-附件五上报表（涉敏资产梳理汇总表）- 使用通用导出引擎

    支持参数:
        limit: 限制导出行数（用于测试），0表示全部
    """
    return _run_config_export('smc', 'attachment_5', '导出SMC上报-附件五')

